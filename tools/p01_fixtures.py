"""Synthetic P01 formats and scale datasets; never uses customer data."""
import hashlib
import json
from pathlib import Path
import shutil
import zipfile

import openpyxl

from collect_p01_samples import HEADERS


def write_json(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def workbook(path, count=2, variant="current", summary=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "利润明细"
    headers = list(HEADERS)
    if variant == "legacy":
        headers[1] = "用户"
    elif variant == "remark":
        headers.append("备注")
    elif variant == "reordered":
        headers.reverse()
    elif variant == "duplicate_headers":
        headers[2] = headers[1]
    ws.append(["合成基线，不是真实成交账本"])
    ws.append([])
    ws.append([])
    if summary:
        ws.append(["总收入", count * 10, "总成本", count * 4, "总利润", count * 6, "综合利润率", 0.6 if count else 0])
    else:
        ws.append([])
    ws.append([])
    ws.append(headers)
    for index in range(count):
        record = dict(zip(HEADERS, ["2026-10-10 12:00:00", "合成客户0001", f"合成订单{index:06d}", "合成微信", "合成商品", "合成商品", "是", None, None, 1, 10, 4, 1, None, None, None, 10, 4, 6, "60.0%", "", "", "1", "10", "4", "1"]))
        record["用户"] = record["淘宝ID"]
        record["备注"] = "合成行备注" if index == 0 else None
        if variant in {"formula", "formula_cached"} and index == 0:
            record["销售额"] = "=1+1"
        ws.append([record.get(header) for header in headers])
    if variant == "no_summary":
        ws.append(["--- 本次总计 ---"] + [None] * 15 + [count * 10, count * 4, count * 6])
    if variant == "latest_note":
        note = wb.create_sheet("备注")
        note.append(["最近备注"])
        note.cell(5, 1, "备注")
        note.cell(5, 2, "合成文件备注，归属未知")
    if variant == "empty_evidence_shape":
        issue = wb.create_sheet("输入问题")
        issue.append(["合成输入诊断"])
        issue.append([])
        issue.append([])
        issue.append(["记录时间", "级别", "位置", "字段", "输入内容", "提示"])
        issue.append(["2026-10-10 12:00:00", "提醒", "行1", "数量", "", "合成诊断"])
        price = wb.create_sheet("价格设定")
        price.append(["合成价格快照"])
        price.append([])
        price.append([])
        price.append(["项目名称", "销售单价原文", "成本单价原文", "K系数原文", "是否单件"])
        for i in range(9):
            price.append([f"合成商品{i}", "10", "4", "1", "是" if i >= 6 else "否"])
    wb.save(path)
    wb.close()
    if variant == "formula_cached":
        # Synthetic OOXML only: supply a cache without invoking Excel/formulas.
        with zipfile.ZipFile(path) as archive:
            parts = [(info, archive.read(info.filename)) for info in archive.infolist()]
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for info, data in parts:
                if info.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b"<f>1+1</f><v></v>", b"<f>1+1</f><v>2</v>")
                archive.writestr(info, data)


def formats(folder):
    folder.mkdir()
    entries = []
    configurations = [
        ("items", {"items": [{"name": "合成商品", "sale": "10", "cost": "4", "k": "1", "single": True}]}, "json_items", []),
        ("legacy", {"合成商品": ["10", "4", "1"]}, "json_legacy_price_dict", []),
        ("false_string", {"items": [{"name": "合成商品", "sale": "10", "cost": "4", "k": "1", "single": "false"}]}, "json_items", ["SINGLE_NOT_BOOLEAN"]),
        ("duplicate_name", {"items": [{"name": "合成商品", "sale": "10", "cost": "4", "k": "1"}] * 2}, "json_items", ["DUPLICATE_PRODUCT_NAME"]),
        ("invalid_member", {"items": [None]}, "json_items", ["ITEM_NOT_OBJECT"]),
        ("unknown", {"other": "合成未知字段"}, "unknown_json", ["LEGACY_PRICE_SHAPE_UNRECOGNIZED"]),
    ]
    for name, data, schema, issues in configurations:
        path = folder / f"config_{name}.json"
        write_json(path, data)
        entries.append({"path": path.name, "kind": "config", "schema": schema, "issues": issues})
    for variant, schema in [("current", "excel_current_26"), ("empty_evidence_shape", "excel_current_26"), ("reordered", "excel_current_26_reordered"), ("remark", "excel_current_26_with_remark"), ("legacy", "excel_legacy_user_header"), ("no_summary", "excel_current_26"), ("latest_note", "excel_current_26"), ("formula", "excel_current_26"), ("formula_cached", "excel_current_26"), ("duplicate_headers", "unknown_duplicate_headers")]:
        path = folder / f"ledger_{variant}.xlsx"
        workbook(path, 0 if variant == "empty_evidence_shape" else 2, variant, summary=variant != "no_summary")
        entries.append({"path": path.name, "kind": "excel", "first_sheet_schema": schema})
    duplicate = folder / "ledger_duplicate_file.xlsx"
    shutil.copyfile(folder / "ledger_current.xlsx", duplicate)
    entries.append({"path": duplicate.name, "kind": "excel", "first_sheet_schema": "excel_current_26", "same_bytes_as": "ledger_current.xlsx"})
    path = folder / "ledger_corrupt.xlsx"
    path.write_bytes(b"synthetic truncated xlsx, not a ZIP")
    entries.append({"path": path.name, "kind": "excel", "expected_read_error": "BadZipFile"})
    for entry in entries:
        entry["sha256"] = hashlib.sha256((folder / entry["path"]).read_bytes()).hexdigest()
    write_json(folder / "manifest.json", {"origin": "synthetic", "actual_legacy_existence": "unknown", "files": entries})
    return entries


def scales(folder):
    folder.mkdir()
    result = {"origin": "synthetic", "row_sets": [], "file_sets": [], "hash_policy": "generated OOXML bytes vary with library/timestamp; row contents and counts fixed"}
    for count in (100, 1000, 10000):
        path = folder / f"rows_{count}.xlsx"
        workbook(path, count)
        result["row_sets"].append({"path": path.name, "detail_rows": count, "income": str(count * 10), "cost": str(count * 4), "profit": str(count * 6), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
    template = folder / "file_template.xlsx"
    workbook(template, 1)
    for count in (10, 100, 1000):
        directory = folder / f"files_{count}"
        directory.mkdir()
        for i in range(count):
            shutil.copyfile(template, directory / f"synthetic_{i:04d}.xlsx")
        result["file_sets"].append({"path": directory.name, "files": count, "detail_rows_per_file": 1, "same_bytes_intentionally": True})
    result["checks"] = check_scales(folder, result)
    write_json(folder / "manifest.json", result)
    return result


def check_scales(folder, manifest):
    checks = []
    for entry in manifest["row_sets"]:
        wb = openpyxl.load_workbook(folder / entry["path"], read_only=True, data_only=False)
        try:
            income = cost = profit = count = 0
            orders = set()
            for row in wb.active.iter_rows(min_row=7, values_only=True):
                count += 1
                orders.add(row[2])
                income += row[16]
                cost += row[17]
                profit += row[18]
            passed = count == entry["detail_rows"] == len(orders) and (income, cost, profit) == tuple(int(entry[key]) for key in ("income", "cost", "profit"))
            checks.append({"path": entry["path"], "detail_rows": count, "distinct_order_refs": len(orders), "totals": [income, cost, profit], "passed": passed})
        finally:
            wb.close()
    template_hash = hashlib.sha256((folder / "file_template.xlsx").read_bytes()).hexdigest()
    for entry in manifest["file_sets"]:
        paths = list((folder / entry["path"]).glob("*.xlsx"))
        passed = len(paths) == entry["files"] and all(hashlib.sha256(path.read_bytes()).hexdigest() == template_hash for path in paths)
        checks.append({"path": entry["path"], "files": len(paths), "passed": passed})
    return checks
