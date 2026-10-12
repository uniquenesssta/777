"""Performance-v1 measurements, independent of Tk and financial computation."""
from collections import Counter
from contextlib import contextmanager, ExitStack
import math
from pathlib import Path
import statistics
from threading import Lock
from unittest.mock import patch
from zipfile import ZipFile

import openpyxl
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet.worksheet import Worksheet


def percentile(values, fraction=.95):
    values = sorted(values)
    if not values:
        return None
    return values[max(0, math.ceil(len(values) * fraction) - 1)]


def distribution(values):
    return {'samples': len(values), 'p50': statistics.median(values) if values else None,
            'p95': percentile(values), 'max': max(values) if values else None}


def healthy_heartbeat(data, p95_limit=100, max_limit=500):
    gaps = data.get('gaps_ms', [])
    if len(gaps) < 20 or data.get('samples') != len(gaps) or any(not math.isfinite(g) or g < 0 for g in gaps):
        return False
    stats = distribution(gaps)
    return data.get('status') == 'passed' and data.get('p95_gap_ms') == stats['p95'] and data.get('max_gap_ms') == stats['max'] and stats['p95'] <= p95_limit and stats['max'] <= max_limit


class Heartbeat:
    """Record every gap inside an explicit phase, including its boundaries."""
    def __init__(self, p95_limit=100, max_limit=500):
        self.p95_limit, self.max_limit = p95_limit, max_limit
        self.last, self.gaps = None, []

    def start(self, now):
        self.started = now
        self.last, self.gaps = now, []

    def ready(self, now):
        # Timer delivery can be slower than the requested 50ms (e.g. Windows
        # clock rounding). Elapsed time alone cannot establish sample count.
        return self.last is not None and now - self.started >= 1.1 and len(self.gaps) >= 20

    def pulse(self, now):
        if self.last is not None:
            gap = (now - self.last) * 1000
            if not math.isfinite(gap) or gap < 0:
                raise ValueError('INVALID_HEARTBEAT_CLOCK')
            self.gaps.append(gap)
            self.last = now

    def finish(self, now):
        self.pulse(now)
        gaps, self.last = self.gaps, None
        stats = distribution(gaps)
        enough = len(gaps) >= 20
        passed = enough and stats['p95'] <= self.p95_limit and stats['max'] <= self.max_limit
        return {'interval_ms': 50, 'samples': len(gaps), 'p95_gap_ms': stats['p95'],
                'max_gap_ms': stats['max'], 'gaps_ms': gaps,
                'over_100_ms': sum(g > self.p95_limit for g in gaps),
                'over_500_ms': sum(g > self.max_limit for g in gaps),
                'status': 'passed' if passed else 'failed' if enough else 'insufficient_samples',
                'observation_seconds': now - self.started,
                'scope': '50ms actual Tk heartbeat during operation, at least 1.1s observation and 20 timer callbacks; includes every phase boundary gap and idle settling. Preparation and post-measurement validation are separate.'}


@contextmanager
def instrumentation(namespace=None):
    """Count matching library boundaries for original and current implementations."""
    counts, lock = Counter(), Lock()
    def increment(key):
        with lock:
            counts[key] += 1
    opened = openpyxl.load_workbook
    xml_open = ZipFile.open
    source = ReadOnlyWorksheet._get_source
    regular_rows, readonly_rows = Worksheet.iter_rows, ReadOnlyWorksheet.iter_rows
    def load(*args, **kwargs):
        increment('workbook_opens')
        return opened(*args, **kwargs)
    def xml(package, name, *args, **kwargs):
        part = getattr(name, 'filename', name)
        if isinstance(part, str) and part.startswith('xl/worksheets/') and part.endswith('.xml'):
            increment('worksheet_xml_source_opens')
        return xml_open(package, name, *args, **kwargs)
    def read_source(sheet):
        increment('readonly_get_source_calls')
        return source(sheet)
    def rows(original):
        def counted(sheet, *args, **kwargs):
            increment('iter_rows_calls')
            for row in original(sheet, *args, **kwargs):
                increment('yielded_rows')
                yield row
        return counted
    with ExitStack() as stack:
        stack.enter_context(patch.object(openpyxl, 'load_workbook', load))
        stack.enter_context(patch.object(ZipFile, 'open', xml))
        stack.enter_context(patch.object(ReadOnlyWorksheet, '_get_source', read_source))
        stack.enter_context(patch.object(Worksheet, 'iter_rows', rows(regular_rows)))
        stack.enter_context(patch.object(ReadOnlyWorksheet, 'iter_rows', rows(readonly_rows)))
        if namespace is not None:
            globals_ = namespace['PriceCalculator'].refresh_excel_files.__globals__
            stack.enter_context(patch.dict(globals_, {'load_workbook': load}))
        yield counts


def input_hashes(folder):
    import hashlib
    return {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted(Path(folder).glob('*.xlsx'))}
