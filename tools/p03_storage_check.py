"""P03-03 temporary storage/price/screenshot checks and optional Windows Tk pair."""
import argparse
import hashlib
import io
from pathlib import Path
import platform
import subprocess
import sys
import time
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / "tools"), str(ROOT / "tests")]
from p03_format_check import case_ids
import p02_gui_check as gui
BASE_SHARED_SMOKE = gui.shared_directory_smoke


def storage_gui_smoke(manager, sales, sandbox):
    import tkinter as tk
    from PIL import Image
    from openpyxl import load_workbook
    original_prices = (sandbox / "config.json").read_bytes()
    checks = BASE_SHARED_SMOKE(manager, sales, sandbox)
    # Existing helper owns close assertions; postpone them until all checks end.
    for app in (manager, sales):
        for job in tuple(app._after_jobs):
            app.root.after_cancel(job)
        app._after_jobs.clear()
    assert (sandbox / "config.json").read_bytes() == original_prices
    checks["directory_save_preserves_price_bytes"] = "passed"
    clone = type(manager)(tk.Toplevel(manager.root))
    try:
        assert clone.data_root == manager.data_root
        assert clone.config_file == manager.config_file
        assert clone.report_dir == manager.report_dir and clone.image_dir == manager.image_dir
        assert clone._storage_valid
    finally:
        clone.on_close()
    checks["actual_new_manager_window_restores_local_connection"] = "passed"
    class FirstRun(type(manager)):
        local_settings_file = str(sandbox / "first-run-settings.json")
    first_run = FirstRun(tk.Toplevel(manager.root))
    try:
        assert not first_run.data_root
        assert first_run.notebook.select() == str(first_run.excel_page)
        first_run.data_source_var.set(str(sandbox))
        assert first_run.apply_connection()
        assert first_run.report_dir == manager.report_dir
    finally:
        first_run.on_close()
    checks["first_launch_visible_connection_then_apply"] = "passed"
    path = sales.get_daily_excel_path()
    book = load_workbook(path)
    try:
        before = list(book["订单价格快照"].values)
    finally:
        book.close()
    next(row for row in manager.setting_rows if row["name"].get() == "P02面积")["sale"].set("20")
    assert manager.save_settings()
    manager.add_calc_row("新价格订单", "P02面积", False)
    manager.calculate(notify=False)
    manager.confirm_current_prices()
    assert manager.export_to_excel() == path
    book = load_workbook(path)
    try:
        after = list(book["订单价格快照"].values)
        assert after[:len(before)] == before
        assert len(after) == len(before) + 2
        assert after[-1][7] == "20"
    finally:
        book.close()
    checks["old_order_price_snapshot_preserved_after_new_price_export"] = "passed"
    with patch("PIL.ImageGrab.grab", return_value=Image.new("RGB", (1, 1))), patch("quote_app.integration.time.time_ns", return_value=1):
        first = manager._capture_current_quote()
        second = manager._capture_current_quote()
        assert first and second and first != second
    checks["same_timestamp_png_outputs_unique"] = "passed"
    for app in (manager, sales):
        app.schedule_quote_callback(lambda: (_ for _ in ()).throw(AssertionError("closed callback ran")), delay=1)
    return checks


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--gui-smoke", action="store_true", help="actual mother/sales Tk pair in temporary directories, Windows only")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args(argv)
    paths = sorted([*ROOT.glob("quote_files/*.py"), *ROOT.glob("quote_app/*.py"),
                    *ROOT.glob("报价*.py"), ROOT / "tests/test_storage_configuration.py", Path(__file__),
                    ROOT / "tools/p02_gui_check.py"])
    suite = unittest.defaultTestLoader.loadTestsFromName("test_storage_configuration")
    ids = list(case_ids(suite))
    log = io.StringIO()
    started = time.monotonic()
    result = unittest.TextTestRunner(stream=log, verbosity=2).run(suite)
    report = {
        "task": "P03-03",
        "commit": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(),
        "environment": {"os": platform.system(), "python": platform.python_version()},
        "implementation_sha256_LF": {p.relative_to(ROOT).as_posix(): hashlib.sha256(p.read_bytes().replace(b"\r\n", b"\n")).hexdigest() for p in paths},
        "scope": "Synthetic temporary settings/prices/workbooks/images; no production paths, installation or backup scheduling",
        "production_paths_accessed": False,
        "windows_gui": "not_requested",
        "checks": {"status": "passed" if result.wasSuccessful() else "failed", "tests_run": result.testsRun,
                   "failures": len(result.failures), "errors": len(result.errors), "skipped": result.skipped,
                   "elapsed_seconds": round(time.monotonic() - started, 3), "test_ids": ids, "log": log.getvalue()},
        "status": "passed" if result.wasSuccessful() else "failed",
    }
    if args.gui_smoke:
        if platform.system() != "Windows":
            report["windows_gui"] = "blocked: Windows GUI only"
            report["status"] = "blocked" if result.wasSuccessful() else "failed"
        elif result.wasSuccessful():
            import p02_gui_check as gui
            try:
                with patch.object(gui, "shared_directory_smoke", side_effect=storage_gui_smoke):
                    report["windows_gui"] = gui.run_pair()
                report["status"] = report["windows_gui"]["status"]
            except Exception as exc:
                report["windows_gui"] = {"status": "failed", "error": f"{type(exc).__name__}: {exc}"}
                report["status"] = "failed"
    from baseline_check import emit
    emit(report, args.output)
    return 2 if report["status"] == "blocked" else int(report["status"] == "failed")


if __name__ == "__main__":
    raise SystemExit(main())
