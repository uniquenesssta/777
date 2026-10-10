"""Read explicitly selected P01-03 files; write only a new reports subdirectory.

Private: read-only original snapshots and source-path provenance, never upload.
Public: pseudonymized JSON/schema observations, review before sharing.
No application import/startup, recursive search, formula evaluation, or SMB probe.
"""
import argparse
from datetime import date, datetime
import hashlib
import json
import math
from pathlib import Path
import platform
import re
import stat
import sys
import zipfile

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
MAX_BYTES = 20 * 1024 * 1024
MAX_ZIP_BYTES = 64 * 1024 * 1024
MAX_ROWS, MAX_COLS = 1000, 64
SAFE_CODES = {"JSON_DUPLICATE_KEY", "JSON_NON_FINITE", "ZIP_CAPACITY_EXCEEDED", "SOURCE_NOT_REGULAR_FILE", "SOURCE_CAPACITY_EXCEEDED", "SNAPSHOT_NOT_READONLY", "SOURCE_CHANGED_DURING_COLLECTION", "OUTPUT_LINK_DISALLOWED", "OUTPUT_MUST_BE_NEW_REPORTS_SUBDIRECTORY", "OUTPUT_ALREADY_EXISTS", "SOURCE_INSIDE_OUTPUT"}
HEADERS = ["记录时间", "淘宝ID", "订单号", "微信号", "项目名称", "基础类型", "是否单件", "高度(m)", "宽度(m)", "数量", "销售单价", "成本单价", "K系数", "折扣", "折扣系数", "折前销售额", "销售额", "总成本", "利润", "利润率", "高度原文", "宽度原文", "数量原文", "销售单价原文", "成本单价原文", "K系数原文"]
NUMERIC = set(HEADERS[7:]) - {"备注"}
LABELS = set(HEADERS) | {"用户", "备注", "文件备注", "总收入", "总利润", "综合利润率", "最后导出时间", "最近导出淘宝ID", "保存时间", "淘宝ID/用户", "级别", "位置", "字段", "输入内容", "提示"}
SHEETS = {"利润明细", "价格设定", "备注", "输入问题"}
MARKERS = {"备注", "文件备注", "折扣", "--- 本次总计 ---"}
NUMERIC_TEXT = re.compile(r"[\d\s.,+*/%()eE\-￥¥元块人民币米mM个件张套倍×÷xX（）％，]+\Z")
TIME_TEXT = re.compile(r"\d{4}-\d{2}-\d{2}(?: \d{2}:\d{2}:\d{2})?\Z")


class Redactor:
    def __init__(self):
        self.maps = {}

    def token(self, category, value):
        mapping = self.maps.setdefault(category, {})
        key = str(value)
        if key not in mapping:
            mapping[key] = f"{category}{len(mapping) + 1:04d}"
        return mapping[key]

    def value(self, value, field=None):
        if value is None:
            return None
        if field in {"淘宝ID", "用户", "淘宝ID/用户", "最近导出淘宝ID"}:
            return self.token("客户", value)
        if field in {"订单号", "微信号"}:
            return self.token("订单" if field == "订单号" else "微信", value)
        if field in {"项目名称", "基础类型", "name"}:
            return value if isinstance(value, str) and value in MARKERS else self.token("商品", value)
        if field in NUMERIC | {"sale", "cost", "k", "总收入", "总利润", "综合利润率"}:
            if isinstance(value, (int, float, bool)):
                return value if not isinstance(value, float) or math.isfinite(value) else "<NON_FINITE>"
            if isinstance(value, str) and len(value) <= 80 and (value == "" or NUMERIC_TEXT.fullmatch(value)):
                return value
        if field in {"是否单件", "single", "qty_only"}:
            if isinstance(value, bool) or value in ("是", "否", "折扣行", "true", "false"):
                return value
        if field in {"记录时间", "保存时间", "最后导出时间"}:
            if isinstance(value, (date, datetime)):
                return value.isoformat()
            if isinstance(value, str) and TIME_TEXT.fullmatch(value):
                return value
        if field is None and isinstance(value, str) and value in LABELS:
            return value
        return self.token("文本", value)


def json_pairs(pairs):
    obj = {}
    for key, value in pairs:
        if key in obj:
            raise ValueError("JSON_DUPLICATE_KEY")
        obj[key] = value
    return obj


def reject_constant(_value):
    raise ValueError("JSON_NON_FINITE")


def inspect_config(data, redactor):
    obj = json.loads(data.decode("utf-8-sig"), object_pairs_hook=json_pairs, parse_constant=reject_constant)
    issues = []
    if not isinstance(obj, dict):
        return {"schema": "unknown_json", "issues": ["TOP_LEVEL_NOT_OBJECT"]}
    if isinstance(obj.get("items"), list):
        schema = "json_items"
        items = []
        names = []
        for item in obj["items"]:
            if not isinstance(item, dict):
                issues.append("ITEM_NOT_OBJECT")
                items.append({"invalid_type": type(item).__name__})
                continue
            name = item.get("name")
            names.append(name.strip() if isinstance(name, str) else "")
            if not isinstance(name, str) or not name.strip():
                issues.append("NAME_INVALID")
            out = {key: redactor.value(item[key], key) for key in ("name", "sale", "cost", "k", "single", "qty_only") if key in item}
            out["unknown_field_count"] = len(set(item) - {"name", "sale", "cost", "k", "single", "qty_only"})
            for flag in ("single", "qty_only"):
                if flag in item and not isinstance(item[flag], bool):
                    issues.append("SINGLE_NOT_BOOLEAN")
            if not {"sale", "cost", "k"}.issubset(item):
                issues.append("PRICE_FIELDS_MISSING")
            items.append(out)
        if len(names) != len(set(names)):
            issues.append("DUPLICATE_PRODUCT_NAME")
        unknown = len(set(obj) - {"items", "export_dir"})
    else:
        entries = [(key, value) for key, value in obj.items() if key != "export_dir"]
        if not entries or not all(isinstance(value, list) and len(value) >= 3 for _, value in entries):
            return {"schema": "unknown_json", "issues": ["LEGACY_PRICE_SHAPE_UNRECOGNIZED"]}
        schema, unknown = "json_legacy_price_dict", 0
        items = [{"name": redactor.value(name, "name"), "prices": [redactor.value(value[i], field) for i, field in enumerate(("sale", "cost", "k"))], "extra_value_count": len(value) - 3} for name, value in entries]
    if not items:
        issues.append("NO_PRICE_ITEMS")
    return {"schema": schema, "issues": sorted(set(issues)), "items": items, "export_dir_present": "export_dir" in obj, "unknown_top_field_count": unknown}


def classify_header(rows):
    for row_index, values in rows[:30]:
        texts = [value if isinstance(value, str) else "" for value in values]
        fields = {v for v in texts if v}
        if {"记录时间", "项目名称"}.issubset(fields) and fields.intersection({"淘宝ID", "用户"}):
            if len([v for v in texts if v]) != len(fields):
                return row_index, texts, "unknown_duplicate_headers"
            if fields == set(HEADERS) and texts[:len(HEADERS)] == HEADERS:
                schema = "excel_current_26"
            elif fields == set(HEADERS):
                schema = "excel_current_26_reordered"
            elif fields == set(HEADERS) | {"备注"}:
                schema = "excel_current_26_with_remark"
            elif "用户" in fields and {"销售额", "总成本", "利润"}.issubset(fields):
                schema = "excel_legacy_user_header"
            else:
                schema = "unknown_partial_headers"
            return row_index, texts, schema
        if {"项目名称", "销售单价原文", "成本单价原文", "K系数原文", "是否单件"}.issubset(fields):
            return row_index, texts, "excel_price_snapshot"
        if texts[:6] == ["记录时间", "级别", "位置", "字段", "输入内容", "提示"]:
            return row_index, texts, "excel_input_issues"
    return None, [], "unknown_sheet"


def inspect_workbook(path, redactor):
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > 2000 or sum(info.file_size for info in infos) > MAX_ZIP_BYTES:
            raise ValueError("ZIP_CAPACITY_EXCEEDED")
        external_links = sum(info.filename.startswith("xl/externalLinks/externalLink") and info.filename.endswith(".xml") for info in infos)
        macro_parts = sum(info.filename == "xl/vbaProject.bin" for info in infos)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = None
    try:
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        sheets = []
        for index, ws in enumerate(wb.worksheets, 1):
            rows = [(i, list(row)) for i, row in enumerate(ws.iter_rows(max_row=min(ws.max_row or MAX_ROWS, MAX_ROWS), max_col=min(ws.max_column or MAX_COLS, MAX_COLS)), 1)]
            raw_rows = [(i, [cell.value for cell in row]) for i, row in rows]
            header_row, headers, schema = classify_header(raw_rows)
            if ws.title == "备注" and raw_rows and raw_rows[0][1][0] == "最近备注":
                schema = "excel_latest_note"
            known = ws.title in SHEETS and not schema.startswith("unknown")
            out = {"sheet": ws.title if ws.title in SHEETS else f"未知工作表{index}", "schema": schema if ws.title in SHEETS else "unknown_sheet", "reported_rows": ws.max_row, "reported_columns": ws.max_column, "truncated": bool((ws.max_row or 0) > MAX_ROWS or (ws.max_column or 0) > MAX_COLS), "header_row": header_row, "sampled_formula_count": sum(cell.data_type == "f" for _, row in rows for cell in row), "sampled_formula_without_cache": None, "cells": []}
            # Unknown schema: metadata only, never release unclassified values.
            if known:
                out["sampled_formula_without_cache"] = 0
                cache_rows = iter(cached.worksheets[index - 1].iter_rows(max_row=len(rows), max_col=len(rows[0][1]) if rows else 1, values_only=True))
                for i, row in rows:
                    cache_row = next(cache_rows, ())
                    for j, cell in enumerate(row, 1):
                        value = cell.value
                        if value is None:
                            continue
                        field = headers[j - 1] if header_row and i > header_row and j <= len(headers) else None
                        if schema == "excel_latest_note" and j == 2:
                            field = raw_rows[i - 1][1][0]
                        if (not header_row or i < header_row) and j % 2 == 0 and raw_rows[i - 1][1][j - 2] in LABELS:
                            field = raw_rows[i - 1][1][j - 2]
                        if cell.data_type == "f":
                            cache = cache_row[j - 1] if j <= len(cache_row) else None
                            public = {"formula_redacted": True, "has_cached_value": cache is not None}
                            out["sampled_formula_without_cache"] += cache is None
                            if field in NUMERIC and isinstance(cache, (int, float)) and not isinstance(cache, bool) and math.isfinite(cache):
                                public["cached_numeric_value"] = cache
                        else:
                            public = redactor.value(value, field)
                        out["cells"].append({"row": i, "column": j, "type": cell.data_type, "value": public})
            sheets.append(out)
        return {"schema": "excel_workbook", "external_link_parts": external_links, "macro_parts": macro_parts, "sheets": sheets}
    finally:
        wb.close()
        if cached is not None:
            cached.close()


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def collect(selected, output, root=ROOT):
    requested = Path(output).absolute()
    for ancestor in (requested, *requested.parents):
        if ancestor.is_symlink() or getattr(ancestor, "is_junction", lambda: False)():
            raise ValueError("OUTPUT_LINK_DISALLOWED")
    output = Path(output).resolve()
    if not output.is_relative_to((root / "reports").resolve()) or output == (root / "reports").resolve():
        raise ValueError("OUTPUT_MUST_BE_NEW_REPORTS_SUBDIRECTORY")
    if output.exists():
        raise ValueError("OUTPUT_ALREADY_EXISTS")
    sources = [(role, kind, Path(path).resolve()) for role, kind, path in selected]
    if any(path.is_relative_to(output) for _, _, path in sources):
        raise ValueError("SOURCE_INSIDE_OUTPUT")
    output.mkdir(parents=True, exist_ok=False)
    (output / "private").mkdir()
    (output / "public").mkdir()
    redactor, files, provenance, review = Redactor(), [], [], []
    for index, (role, kind, path) in enumerate(sources, 1):
        file_id = f"S{index:03d}"
        entry = {"id": file_id, "role": role, "kind": kind}
        private = {"id": file_id, "source_path": str(path)}
        try:
            if not path.is_file():
                raise ValueError("SOURCE_NOT_REGULAR_FILE")
            if path.stat().st_size > MAX_BYTES:
                raise ValueError("SOURCE_CAPACITY_EXCEEDED")
            # Cap the actual read too: the source may grow after stat().
            with path.open("rb") as stream:
                data = stream.read(MAX_BYTES + 1)
            if len(data) > MAX_BYTES:
                raise ValueError("SOURCE_CAPACITY_EXCEEDED")
            digest = hashlib.sha256(data).hexdigest()
            private.update(source_sha256=digest, source_mtime_ns=path.stat().st_mtime_ns)
            snapshot = output / "private" / (file_id + (".json" if kind == "config" else ".xlsx"))
            snapshot.write_bytes(data)
            snapshot.chmod(stat.S_IRUSR | stat.S_IRGRP | stat.S_IROTH)
            if snapshot.stat().st_mode & stat.S_IWUSR:
                raise ValueError("SNAPSHOT_NOT_READONLY")
            entry.update(source_sha256=digest, size_bytes=len(data), readonly_snapshot=True)
            private["snapshot"] = snapshot.name
            with path.open("rb") as stream:
                second = stream.read(MAX_BYTES + 1)
            if hashlib.sha256(second).hexdigest() != digest:
                raise ValueError("SOURCE_CHANGED_DURING_COLLECTION")
            observed = inspect_config(data, redactor) if kind == "config" else inspect_workbook(snapshot, redactor)
            if observed["schema"].startswith("unknown") or observed.get("issues"):
                review.append({"id": file_id, "schema": observed["schema"], "reasons": observed.get("issues", ["UNKNOWN_SCHEMA"])})
            for sheet in observed.get("sheets", []):
                reasons = []
                if sheet["schema"].startswith("unknown"):
                    reasons.append("UNKNOWN_SCHEMA")
                if sheet["truncated"]:
                    reasons.append("SAMPLE_WINDOW_TRUNCATED")
                if sheet["sampled_formula_without_cache"]:
                    reasons.append("FORMULA_CACHE_MISSING")
                if reasons:
                    review.append({"id": file_id, "sheet": sheet["sheet"], "schema": sheet["schema"], "reasons": reasons})
            if observed.get("external_link_parts") or observed.get("macro_parts"):
                review.append({"id": file_id, "reasons": ["EXTERNAL_LINKS_OR_MACROS_PRESENT"]})
            public_path = output / "public" / f"{file_id}.json"
            write_json(public_path, observed)
            entry.update(status="observed", schema=observed["schema"], public_snapshot=public_path.name)
        except Exception as exc:
            # Paths, raw JSON decoder excerpts and customer values never enter
            # the public failure report. The private snapshot remains local.
            code = str(exc) if isinstance(exc, ValueError) and str(exc) in SAFE_CODES else type(exc).__name__
            entry.update(status="blocked", error_code=code)
        files.append(entry)
        provenance.append(private)
    config_hashes = {e["role"]: e["source_sha256"] for e in files if e["status"] == "observed" and e["kind"] == "config" and e["role"] in {"manager", "sales"}}
    report = {"task": "P01-03", "status": "evidence_only_acceptance_pending", "environment": {"os": platform.system(), "python": platform.python_version(), "openpyxl": openpyxl.__version__}, "scope": "显式选中文件的只读采集，不启动应用、不评估公式、不执行网络探测或部署", "originals_modified": False, "limits": {"source_bytes": MAX_BYTES, "zip_expanded_bytes": MAX_ZIP_BYTES, "rows_per_sheet": MAX_ROWS, "columns_per_sheet": MAX_COLS}, "files": files, "review_required": review, "manager_sales_config_bytes_equal": (config_hashes["manager"] == config_hashes["sales"]) if len(config_hashes) == 2 else None, "same_bytes_do_not_prove_same_physical_share": True, "unknown_formats_require_review": True, "private_directory_must_not_be_shared": True}
    write_json(output / "private" / "provenance.json", provenance)
    write_json(output / "public" / "manifest.json", report)
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manager-config", type=Path)
    parser.add_argument("--sales-config", type=Path)
    parser.add_argument("--config", action="append", type=Path, default=[])
    parser.add_argument("--excel", action="append", type=Path, default=[])
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args(argv)
    selected = []
    for role, path in (("manager", args.manager_config), ("sales", args.sales_config)):
        if path:
            selected.append((role, "config", path))
    selected.extend(("unspecified", "config", path) for path in args.config)
    selected.extend(("unspecified", "excel", path) for path in args.excel)
    if not selected:
        parser.error("至少显式选择一份config或Excel；不自动访问源码中的生产目录")
    try:
        report = collect(selected, args.output)
    except (OSError, ValueError) as exc:
        print(json.dumps({"task": "P01-03", "status": "blocked", "error_code": str(exc) if isinstance(exc, ValueError) and str(exc) in SAFE_CODES else type(exc).__name__}, ensure_ascii=False))
        return 1
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if any(e["status"] == "blocked" for e in report["files"]) else 0


if __name__ == "__main__":
    sys.exit(main())
