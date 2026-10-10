"""P01-04 isolated targets/formats/datasets and optional baseline measurements.

Run from repo root. All output must be a NEW reports subdirectory. Business
sources are fingerprint-gated before definitions load. No production path IO.
"""
import argparse
from contextlib import contextmanager, redirect_stdout
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import platform
import runpy
import shutil
import subprocess
import sys
import tempfile
import time
import tracemalloc

import openpyxl
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

import baseline_check
import collect_p01_samples as collector
import p01_fixtures as fixtures
import p01_oracle as oracle

ROOT = Path(__file__).resolve().parents[1]


def module(path):
    spec = importlib.util.spec_from_file_location(path.stem, path)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


def gate():
    manifest = json.loads((ROOT / "docs/baseline/manifest.json").read_text(encoding="utf-8"))
    sources = [baseline_check.inspect_source(ROOT, entry) for entry in manifest["files"]]
    if any(source["status"] != "ok" for source in sources):
        raise ValueError("SOURCE_BASELINE_CHANGED")
    return manifest, sources


def new_output(output, root=ROOT):
    requested = Path(output).absolute()
    for path in (requested, *requested.parents):
        if path.is_symlink() or getattr(path, "is_junction", lambda: False)():
            raise ValueError("OUTPUT_LINK_DISALLOWED")
    output = requested.resolve()
    if not output.is_relative_to((root / "reports").resolve()) or output == (root / "reports").resolve():
        raise ValueError("OUTPUT_MUST_BE_NEW_REPORTS_SUBDIRECTORY")
    output.mkdir(parents=True, exist_ok=False)
    return output


def golden(manifest):
    suite = json.loads((ROOT / "docs/baseline/p01-04-golden.json").read_text(encoding="utf-8"))
    examples = module(ROOT / "docs/baseline/p01_02_examples.py")
    namespaces = {entry["role"]: runpy.run_path(str(ROOT / entry["path"]), run_name="p01_validation_only") for entry in manifest["files"]}
    results = []
    with tempfile.TemporaryDirectory(prefix="777-p01-golden-") as temp:
        for case in suite["cases"]:
            actual = oracle.quote(case)
            expected = case["expected"]
            passed = {k: actual[k] for k in expected} == expected
            legacy_lines = []
            for line in case["input"]:
                if line["kind"] == "discount":
                    legacy_lines.append(line)
                else:
                    legacy_lines.append({**line, **case["price_snapshot"]["items"][line["product_id"]]})
            original = {role: examples.original(ns, Path(temp), legacy_lines) for role, ns in namespaces.items()}
            differs = {}
            for role, observation in original.items():
                if expected["status"] == "ok":
                    differs[role] = observation["displayed_totals"] != expected["totals"] or observation["displayed_line_income"] != [r["income"] for r in expected["lines"]]
                else:
                    differs[role] = not (observation["errors"] and observation["retained_records"] == 0)
            results.append({"id": case["id"], "target_check_passed": passed, "oracle": actual, "original": original, "legacy_differs_from_target": differs})
    return results


def check_formats(folder, entries):
    results = []
    for entry in entries:
        path = folder / entry["path"]
        before = path.read_bytes()
        try:
            observed = collector.inspect_config(before, collector.Redactor()) if entry["kind"] == "config" else collector.inspect_workbook(path, collector.Redactor())
            if entry["kind"] == "config":
                passed = observed["schema"] == entry["schema"] and observed["issues"] == entry["issues"]
            else:
                first = observed["sheets"][0]
                passed = first["schema"] == entry.get("first_sheet_schema")
                if "formula" in entry["path"]:
                    passed &= first["sampled_formula_count"] == 1 and first["sampled_formula_without_cache"] == (0 if "cached" in entry["path"] else 1)
                    if "cached" in entry["path"]:
                        formula_cells = [cell for cell in first["cells"] if cell["type"] == "f"]
                        passed &= len(formula_cells) == 1 and formula_cells[0]["value"].get("cached_numeric_value") == 2
                if "same_bytes_as" in entry:
                    passed &= before == (folder / entry["same_bytes_as"]).read_bytes()
            results.append({"path": entry["path"], "passed": bool(passed), "observed": observed})
        except Exception as exc:
            results.append({"path": entry["path"], "passed": type(exc).__name__ == entry.get("expected_read_error"), "error": type(exc).__name__})
        results[-1]["source_unchanged"] = before == path.read_bytes()
    return results


def format_summary(results):
    failures = []
    for result in results:
        if not result["passed"] or not result["source_unchanged"]:
            observed = result.get("observed", {})
            failures.append({"path": result["path"], "error": result.get("error"), "source_unchanged": result["source_unchanged"], "schema": observed.get("schema"), "issues": observed.get("issues", []), "sheets": [{key: sheet[key] for key in ("schema", "sampled_formula_count", "sampled_formula_without_cache")} for sheet in observed.get("sheets", [])]})
    return {"count": len(results), "passed": sum(result["passed"] for result in results), "originals_unchanged": all(result["source_unchanged"] for result in results), "failures": failures}


@contextmanager
def instrumentation(namespace):
    globals_ = namespace["PriceCalculator"].refresh_excel_files.__globals__
    load = globals_["load_workbook"]
    source = ReadOnlyWorksheet._get_source
    iterate = ReadOnlyWorksheet.iter_rows
    counts = {"workbook_opens": 0, "worksheet_xml_source_opens": 0, "iter_rows_calls": 0, "yielded_rows": 0}
    def counted_load(*args, **kwargs):
        counts["workbook_opens"] += 1
        return load(*args, **kwargs)
    def counted_source(self):
        counts["worksheet_xml_source_opens"] += 1
        return source(self)
    def counted_rows(self, *args, **kwargs):
        counts["iter_rows_calls"] += 1
        for row in iterate(self, *args, **kwargs):
            counts["yielded_rows"] += 1
            yield row
    globals_["load_workbook"] = counted_load
    ReadOnlyWorksheet._get_source = counted_source
    ReadOnlyWorksheet.iter_rows = counted_rows
    try:
        yield counts
    finally:
        globals_["load_workbook"] = load
        ReadOnlyWorksheet._get_source = source
        ReadOnlyWorksheet.iter_rows = iterate


def refresh(app, role):
    app.refresh_excel_files()
    if role == "manager":
        app.aggregate_all_excels(silent=True)  # actual repeated refresh chain


class CaptureTree:
    def __init__(self):
        self.values = {}

    def get_children(self):
        return tuple(self.values)

    def delete(self, item):
        self.values.pop(item)

    def insert(self, *args, **kwargs):
        self.values[kwargs["iid"]] = kwargs["values"]

    def item(self, item, option):
        return self.values[item]


def measure(namespace, app, role, condition):
    with instrumentation(namespace) as counts:
        tracemalloc.start()
        started = time.perf_counter()
        try:
            refresh(app, role)
            elapsed = time.perf_counter() - started
            _, peak = tracemalloc.get_traced_memory()
        finally:
            tracemalloc.stop()
    children = app.excel_file_tree.get_children()
    values = [app.excel_file_tree.item(item, "values") for item in children]
    observed = {"displayed_files": len(children)}
    if role == "sales":
        observed["visible_detail_rows"] = sum(int(value[2]) for value in values)
    else:
        observed["totals"] = {field: getattr(app, "excel_total_" + field).get() for field in ("income", "cost", "profit")}
    return {"condition": condition, "seconds": elapsed, "python_allocation_peak_bytes": peak, "observed": observed, **counts}


def worker(role, folder, gui=False, size="small"):
    manifest, _ = gate()
    entry = next(entry for entry in manifest["files"] if entry["role"] == role)
    folder = Path(folder).resolve()
    if not folder.is_relative_to((ROOT / "reports").resolve()):
        raise ValueError("WORKER_INPUT_OUTSIDE_REPORTS")
    row_count, file_count = {"small": (100, 10), "medium": (1000, 100), "large": (10000, 1000)}[size]
    paths = sorted((folder / f"scales/files_{file_count}").glob("*.xlsx")) if role == "manager" else [folder / f"scales/rows_{row_count}.xlsx"]
    if not gui:
        examples = module(ROOT / "docs/baseline/p01_02_examples.py")
        namespace = runpy.run_path(str(ROOT / entry["path"]), run_name="p01_performance_only")
        app = examples.audit.instance(namespace, folder)
        app.excel_file_tree = CaptureTree()
        app.get_report_excel_paths = lambda: [str(path) for path in paths]
        for field in ("count", "income", "cost", "profit", "margin"):
            setattr(app, "excel_total_" + field, examples.audit.Var())
        return {"role": role, "size": size, "input_files": len(paths), "detail_rows_per_file": 1 if role == "manager" else row_count, "scope": "original methods with control stubs; synthetic local files", "ui_heartbeat": "not_measured", "passes": [measure(namespace, app, role, condition) for condition in ("first_pass_os_cache_uncontrolled", "same_process_warm_pass")]}
    if platform.system() != "Windows":
        return {"role": role, "status": "blocked_non_windows", "ui_heartbeat": "not_measured"}
    with tempfile.TemporaryDirectory(prefix=f"gui-{role}-", dir=folder) as temp:
        sandbox = Path(temp)
        namespace = baseline_check.sandbox_namespace(ROOT, entry, sandbox)
        window = namespace["tk"].Tk()
        try:
            app = baseline_check.sandbox_application_type(namespace, sandbox)(window)
            for directory in (app.report_dir, app.image_dir):
                if not Path(directory).resolve().is_relative_to(sandbox.resolve()):
                    raise ValueError("GUI_SANDBOX_ESCAPE")
            for path in paths:
                shutil.copyfile(path, Path(app.report_dir) / path.name)
            gaps, result, callbacks = [], {}, []
            last = time.perf_counter()
            def pulse():
                nonlocal last
                now = time.perf_counter()
                gaps.append((now - last) * 1000)
                last = now
                window.after(50, pulse)
            def work():
                try:
                    result.update(measure(namespace, app, role, "gui_synthetic_local_refresh"))
                except Exception as exc:
                    callbacks.append(type(exc).__name__ + ": " + str(exc))
                window.after(500, app.on_close)
            window.report_callback_exception = lambda error, value, trace: callbacks.append(error.__name__ + ": " + str(value))
            window.after(50, pulse)
            window.after(1000, work)
            window.mainloop()
            gaps.sort()
            return {"role": role, "status": "measured" if result and not callbacks else "failed", "scope": "original GUI in temporary directories, 50ms after heartbeat; local synthetic refresh", "ui_heartbeat": {"samples": len(gaps), "p95_gap_ms": gaps[max(0, (95 * len(gaps) + 99) // 100 - 1)] if gaps else None, "max_gap_ms": max(gaps, default=None)}, "measurement": result, "callback_errors": callbacks}
        finally:
            try:
                window.destroy()
            except namespace["tk"].TclError:
                pass


def benchmarks(output, gui=False, timeout=30, size="small"):
    results = []
    for role in ("manager", "sales"):
        started = time.perf_counter()
        command = [sys.executable, str(Path(__file__).resolve()), "--worker", role, "--worker-root", str(output), "--size", size]
        if gui:
            command.append("--gui-heartbeat")
        try:
            run = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", timeout=timeout, cwd=ROOT, env={**os.environ, "PYTHONIOENCODING": "utf-8"})
            results.append(json.loads(run.stdout) if run.returncode == 0 else {"role": role, "status": "worker_failed", "returncode": run.returncode, "stderr": run.stderr[-2000:]})
        except subprocess.TimeoutExpired:
            results.append({"role": role, "status": "timeout", "timeout_seconds": timeout, "elapsed_seconds": time.perf_counter() - started, "metrics": "unavailable_after_process_termination; not a successful baseline"})
        finally:
            if gui:
                # The parent owns cleanup even if a blocked GUI worker is killed.
                for sandbox in output.glob(f"gui-{role}-*"):
                    shutil.rmtree(sandbox)
    return results


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--formats-only", action="store_true", help="recheck 18 formats only; skips already-measured calculations, scales and GUI")
    parser.add_argument("--datasets", action="store_true", help="generate 100/1000/10000 rows and 10/100/1000 files")
    parser.add_argument("--benchmark", action="store_true", help="bounded small baseline: manager 10 files, sales 100 rows; implies datasets")
    parser.add_argument("--size", choices=("small", "medium", "large"), default="small", help="performance scales: rows/files 100/10, 1000/100, 10000/1000; each worker bounded to 30 seconds")
    parser.add_argument("--gui-heartbeat", action="store_true", help="optional Windows original GUI heartbeat; implies datasets")
    parser.add_argument("--worker", choices=("manager", "sales"), help=argparse.SUPPRESS)
    parser.add_argument("--worker-root", type=Path, help=argparse.SUPPRESS)
    args = parser.parse_args()
    if args.formats_only and (args.datasets or args.benchmark or args.gui_heartbeat or args.worker):
        parser.error("--formats-only cannot be combined with dataset/measurement modes")
    if args.worker:
        print(json.dumps(worker(args.worker, args.worker_root, args.gui_heartbeat, args.size), ensure_ascii=False, allow_nan=False))
        return 0
    if args.output is None:
        parser.error("--output reports/<new-directory> is required")
    manifest, sources = gate()
    output = new_output(args.output)
    cases = [] if args.formats_only else golden(manifest)
    fixtures.write_json(output / "golden-results.json", cases)
    formats = fixtures.formats(output / "formats")
    format_results = check_formats(output / "formats", formats)
    fixtures.write_json(output / "format-results.json", format_results)
    if not args.formats_only:
        audit = module(ROOT / "docs/audit/reproduce.py")
        audit_output = io.StringIO()
        with redirect_stdout(audit_output):
            audit.main()
        audit_result = json.loads(audit_output.getvalue())
        manager = next(entry for entry in manifest["files"] if entry["role"] == "manager")
        namespace = runpy.run_path(str(ROOT / manager["path"]), run_name="p01_summary_observation_only")
        app = audit.instance(namespace, output / "formats")
        audit_result["evidence"]["A"]["summary_header_collision"] = {"fixture": "ledger_current.xlsx", "expected": {"income": 20, "cost": 8, "profit": 12}, "observed": app.parse_excel_summary(str(output / "formats/ledger_current.xlsx"))}
        fixtures.write_json(output / "legacy-audit.json", audit_result)
    dataset = fixtures.scales(output / "scales") if args.datasets or args.benchmark or args.gui_heartbeat else None
    perf = benchmarks(output, size=args.size) if args.benchmark else []
    gui = benchmarks(output, gui=True, size=args.size) if args.gui_heartbeat and platform.system() == "Windows" else []
    if args.gui_heartbeat and platform.system() != "Windows":
        gui = [{"status": "blocked_non_windows", "ui_heartbeat": "not_measured"}]
    hashes = {path.name: hashlib.sha256(path.read_bytes().replace(b"\r\n", b"\n")).hexdigest() for path in [ROOT / "docs/baseline/p01-04-golden.json", ROOT / "docs/baseline/p01-04-numeric-profile.json", Path(__file__).resolve(), ROOT / "tools/p01_oracle.py", ROOT / "tools/p01_fixtures.py"]}
    passed = all(case["target_check_passed"] for case in cases) and all(case["passed"] and case["source_unchanged"] for case in format_results)
    if dataset:
        passed &= all(check["passed"] for check in dataset["checks"])
    report = {"task": "P01-04", "status": "diagnostic_checks_passed" if passed else "diagnostic_checks_failed", "environment": {"os": platform.system(), "python": platform.python_version(), "openpyxl": openpyxl.__version__}, "scope": "synthetic targets/formats, original methods, optional bounded measurement; not production correctness acceptance", "production_paths_accessed": False, "business_sources_modified": False, "numeric_profile": oracle.PROFILE["profile_id"], "numeric_profile_status": oracle.PROFILE["status"], "normalized_artifact_sha256": hashes, "source_commit": manifest["source_commit"], "sources": sources, "golden": {"count": len(cases), "passed": sum(case["target_check_passed"] for case in cases), "legacy_differences": {role: sum(case["legacy_differs_from_target"][role] for case in cases) for role in ("manager", "sales")}}, "formats": format_summary(format_results), "datasets": dataset, "performance": perf, "gui": gui, "acceptance": "technical_recommendations_and_P01_stage_user_review_pending; Windows_SMB_full_matrix_future"}
    report["environment"]["xml_backend"] = "lxml" if openpyxl.LXML else "stdlib"
    if args.formats_only:
        report["scope"] = "synthetic formats only; calculations, scale datasets, legacy methods and GUI not run"
        report["golden"]["status"] = "not_run_formats_only"
    fixtures.write_json(output / "summary.json", report)
    print(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False))
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
