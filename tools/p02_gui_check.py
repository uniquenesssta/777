"""P02-04 live Windows GUI checks, with synthetic files in temporary directories.

P01 tools continue to load the immutable archived original scripts. This tool
loads the current entrypoints and redirects all configured storage before init.
"""
import argparse
import hashlib
import json
from pathlib import Path
import platform
import subprocess
import sys
import tempfile
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))
import baseline_check

FIXTURE = {"items": [
    {"name": "P02单件", "sale": "0.005", "cost": "0.005", "k": "1", "single": True},
    {"name": "P02面积", "sale": "10", "cost": "4", "k": "1.2", "single": False},
]}


def load_namespace(root, role, sandbox):
    manifest = json.loads((root / "docs/baseline/manifest.json").read_text(encoding="utf-8"))
    entry = next(e for e in manifest["files"] if e["role"] == role)
    path = root / entry["path"]  # Deliberately live, NOT baseline.source_path().
    ns = {"__name__": "p02_current_sandbox", "__file__": str(path)}
    exec(compile(path.read_bytes(), str(path), "exec"), ns)
    reports, images, config = sandbox / "reports", sandbox / "screenshots", sandbox / "config.json"
    reports.mkdir(parents=True, exist_ok=True)
    images.mkdir(parents=True, exist_ok=True)
    config.write_text(json.dumps(FIXTURE, ensure_ascii=False), encoding="utf-8")
    ns.update(APP_SAVE_ROOT=str(sandbox), DEFAULT_REPORT_DIR=str(reports),
              DEFAULT_CONFIG_FILE=str(config), DEFAULT_SCREENSHOT_DIR=str(images),
              LAN_DATA_DIR=str(reports), LAN_PRICE_CONFIG_FILE=str(config),
              LAN_SCREENSHOT_DIR=str(images))
    return ns, entry


def live_sandbox_type(namespace, sandbox):
    parent = baseline_check.sandbox_application_type(namespace, sandbox)
    class SandboxedCurrent(parent):
        def _price_path(self):
            path = super()._price_path()
            if not Path(path).resolve().is_relative_to(sandbox.resolve()):
                raise ValueError("P02隔离检查只允许临时价格文件")
            return path
    return SandboxedCurrent


def files(folder):
    return {str(p.relative_to(folder)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in folder.rglob("*") if p.is_file()}


def smoke(app, sandbox):
    """Actual Tk controls and current class methods; no real screen capture."""
    result = {}
    app.add_calc_row("半分甲", "P02单件", True)
    app.add_calc_row("半分乙", "P02单件", True)
    assert app.calculate(notify=False).totals.income_cents == 2
    assert app.summary_income.get() == "0.02"
    result["two_half_cent_lines"] = "0.02"
    app.user_name.set("000123")
    assert app.quote_session.state.result is None
    app.calculate(notify=False)
    path = app.export_to_excel()
    assert path and Path(path).is_relative_to(sandbox)
    from openpyxl import load_workbook
    book = load_workbook(path)
    try:
        detail = list(book["利润明细"].iter_rows(min_row=7, values_only=True))
        assert detail[0][1] == "000123" and detail[0][16] == "0.01"
        assert detail[1][16] == "0.01" and detail[2][16] == "0.02"
    finally:
        book.close()
    result["frozen_excel_and_text_identity"] = "passed"
    from PIL import Image
    with patch("PIL.ImageGrab.grab", return_value=Image.new("RGB", (1, 1), "white")):
        image_path = app._capture_current_quote()
        assert image_path and Path(image_path).is_relative_to(sandbox)
        with Image.open(image_path) as image:
            assert image.format == "PNG"
    result["valid_screenshot_writer_with_synthetic_image"] = "passed"
    before = files(sandbox)
    app.calc_rows[1]["q"].set("bad")
    assert app.quote_session.state.result is None
    assert app.export_to_excel() is False
    assert app._capture_current_quote() is False
    app.calculate(notify=False)
    assert app.quote_session.state.result.totals is None
    assert app.export_to_excel() is False
    assert app._capture_current_quote() is False
    assert files(sandbox) == before
    result["stale_and_invalid_outputs_zero_writes"] = "passed"
    app.calc_rows[1]["q"].set("1")
    app.calculate(notify=False)
    config = sandbox / "config.json"
    changed = json.loads(config.read_text(encoding="utf-8"))
    changed["items"][0]["sale"] = "20"
    config.write_text(json.dumps(changed, ensure_ascii=False), encoding="utf-8")
    before = files(sandbox)
    assert app.export_to_excel() is False
    assert app.quote_session.state.prices.items[0].sale == "0.005"
    assert files(sandbox) == before
    app.reload_price_settings_from_json(silent=True)
    assert app.price_changes and app.quote_session.state.price_confirmation_required
    app.calculate(notify=False)
    assert app.export_to_excel() is False
    app.confirm_current_prices()
    assert app.quote_session.prepare_output("excel").allowed
    result["price_change_refresh_recalculate_confirm"] = "passed"
    app.add_discount_calc_row("精确折扣")
    discount = app.calc_rows[-1]
    discount["discount_choice"].set("自定义")
    discount["discount_custom"].set("89.123456789012")
    calculated = app.calculate(notify=False)
    assert app.quote_session.state.result is calculated
    assert discount["discount_choice"].get() == "自定义"
    assert calculated.lines[-1].normalized == (("percent", "89.123456789012"),)
    assert app.quote_session.prepare_output("excel").allowed
    result["actual_discount_controls_full_precision"] = "passed"
    retained = app.active_price_config
    config.write_text('{"items":[{"single":"false"}]}', encoding="utf-8")
    before = files(sandbox)
    assert app.reload_price_settings_from_json(silent=True) is False
    assert app.active_price_config is retained
    assert app.export_to_excel() is False and app._capture_current_quote() is False
    assert files(sandbox) == before
    result["bad_config_retained_zero_writes"] = "passed"
    return result


def run_role(role, manual=False):
    import tkinter as tk
    result = {"role": role, "mode": "manual" if manual else "automatic"}
    with tempfile.TemporaryDirectory(prefix="777-p02-current-") as folder:
        sandbox = Path(folder)
        namespace, entry = load_namespace(ROOT, role, sandbox)
        result["source_sha256_LF"] = hashlib.sha256((ROOT / entry["path"]).read_bytes().replace(b"\r\n", b"\n")).hexdigest()
        window = tk.Tk()
        try:
            app = live_sandbox_type(namespace, sandbox)(window)
            result["startup"] = "ok"
            result["tabs"] = [app.notebook.tab(tab, "text").strip() for tab in app.notebook.tabs()]
            if manual:
                print("当前P02窗口：配置、Excel和截图仅在临时目录；关闭后自动清理。", file=sys.stderr)
                window.mainloop()
                result["status"] = "manual_session_finished; user_observation_required"
            else:
                def check():
                    try:
                        with patch("tkinter.messagebox.showinfo"), patch("tkinter.messagebox.showerror"), patch("tkinter.messagebox.showwarning"), patch("PIL.ImageGrab.grab", side_effect=AssertionError("Blocked screenshot attempted capture")):
                            result["checks"] = smoke(app, sandbox)
                        for tab in app.notebook.tabs():
                            app.notebook.select(tab)
                        result["status"] = "passed"
                    except Exception as exc:
                        result["status"] = "failed"
                        result["error"] = f"{type(exc).__name__}: {exc}"
                    finally:
                        app.on_close()
                window.after(100, check)
                window.mainloop()
        finally:
            try:
                window.destroy()
            except tk.TclError:
                pass
        result["close"] = "ok"
    result["sandbox_cleaned"] = True
    return result


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--gui-smoke", action="store_true")
    mode.add_argument("--launch", choices=("manager", "sales"))
    parser.add_argument("--output", type=Path)
    args = parser.parse_args(argv)
    report = {"task": "P02-04", "environment": {"os": platform.system(), "python": platform.python_version()},
              "production_paths_accessed": False, "scope": "Live Tk entrypoints and synthetic outputs; no SMB or real screen capture"}
    report["commit"] = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip()
    report["implementation_sha256_LF"] = {
        str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes().replace(b"\r\n", b"\n")).hexdigest()
        for p in sorted([*ROOT.glob("quote_app/*.py"), *ROOT.glob("quote_core/*.py"), Path(__file__)])
    }
    if platform.system() != "Windows":
        report.update(status="blocked", reason="Windows GUI acceptance only")
        baseline_check.emit(report, args.output)
        return 2
    report["applications"] = []
    for role in ([args.launch] if args.launch else ["manager", "sales"]):
        try:
            report["applications"].append(run_role(role, bool(args.launch)))
        except Exception as exc:
            report["applications"].append({"role": role, "status": "failed", "error": f"{type(exc).__name__}: {exc}"})
    report["status"] = "passed" if all(r["status"] == "passed" for r in report["applications"]) else "manual_observation_pending" if args.launch else "failed"
    baseline_check.emit(report, args.output)
    return int(report["status"] == "failed")


if __name__ == "__main__":
    raise SystemExit(main())
