"""P06-03 native Excel IO and legacy worker compatibility, isolated only."""
import argparse
import base64
import hashlib
import io
import json
import os
from pathlib import Path
import platform
import shutil
import subprocess
import sys
import time
import traceback
import unittest
import uuid
ROOT=Path(__file__).resolve().parents[1];sys.path[:0]=[str(ROOT),str(ROOT/'tools'),str(ROOT/'tests')]
from p01_validation import new_output
from p03_format_check import case_ids
from p06_core_check import run,compare
from p06_ledger_check import fingerprints as ledger_fingerprints
MODULES=['test_native_excel','test_p06_excel_check','test_rust_ledger','test_rust_service','test_p06_ledger_check','test_service_outputs','test_legacy_import','test_p05_import_entry']
def fingerprints():
    result=ledger_fingerprints()
    for name in ['tools/p06_excel_check.py','tools/p05_import.py','tests/test_native_excel.py','tests/legacy_worker_fixture.py','tests/test_p06_excel_check.py','tests/test_p05_import_entry.py','docs/verification/P06-03-protocol.json','docs/verification/P06-03-candidates.json']:
        result[name]=hashlib.sha256((ROOT/name).read_bytes().replace(b'\r\n',b'\n')).hexdigest()
    return dict(sorted(result.items()))
def run_checks():
    suite=unittest.defaultTestLoader.loadTestsFromNames(MODULES);ids=list(case_ids(suite));log=io.StringIO();result=unittest.TextTestRunner(stream=log,verbosity=2).run(suite)
    return dict(status='passed' if result.wasSuccessful() else 'failed',tests_run=result.testsRun,failures=len(result.failures),errors=len(result.errors),skipped=result.skipped,test_ids=ids,log=log.getvalue(),failure_details=[dict(test_id=c.id(),error=t.strip().splitlines()[-1]) for c,t in result.failures+result.errors])
def windows_smoke():
    if os.name!='nt' or platform.system()!='Windows':raise RuntimeError('ACTUAL_WINDOWS_REQUIRED')
    return protocol_smoke(windows=True)

def protocol_smoke(*, windows):
    if windows and (os.name!='nt' or platform.system()!='Windows'):raise RuntimeError('ACTUAL_WINDOWS_REQUIRED')
    from rust_service_fixture import Hosted
    from import_fixture import current_book,append_note_and_void
    from quote_import import preview_bytes,apply_source
    from quote_service.journal import Journal,MemoryJournal
    from quote_service.client import CommitUnknown
    from quote_service.protocol import APIError
    from quote_service.recovery import restore
    from quote_service.hosting import load
    from quote_core import QuoteSnapshot,NormalLine,DraftDetails
    from openpyxl import load_workbook
    host=Hosted()
    try:
        assert host.service.ledger.format_backend=='rust-excel-v1'
        sales=host.client();manager=host.client('manager');owner=host.service.auth.authenticate(sales.token).subject
        source=current_book(host.root/'old-source.xlsx');append_note_and_void(source);raw=source.read_bytes();mtime=source.stat().st_mtime_ns
        plan=preview_bytes(raw,'xlsx',native=host.service.ledger);reference=preview_bytes(raw,'xlsx');assert plan.payload()==reference.payload()
        imported=apply_source(host.service.ledger,host.admin,owner,raw,'xlsx',expected_review_sha256=plan.review_sha256)
        assert apply_source(host.service.ledger,host.admin,owner,raw,'xlsx',expected_review_sha256=plan.review_sha256)==imported
        assert source.read_bytes()==raw and source.stat().st_mtime_ns==mtime
        host.service.ledger.test_fault='after_commit';journal=Journal(host.root/'pending.dat') if windows else MemoryJournal()
        try:
            try:sales.submit(QuoteSnapshot('fixture-v1',(NormalLine('单件','4'),)),DraftDetails('=中文','00003','微信'),journal)
            except CommitUnknown:pass
            else:raise AssertionError('unknown commit not reproduced')
        finally:host.service.ledger.test_fault=None
        assert journal.read() is not None
        if windows:assert b'00003' not in journal.path.read_bytes()
        order=sales.check_pending(journal)['result'];assert journal.read() is None
        export=host.root/'public-excel';export.mkdir()
        with host.service.auth.connection() as c:c.execute('UPDATE preferences SET export_dir=?,revision=revision+1 WHERE id=1',(str(export),))
        tasks=[];contents=[]
        for api,kind in [(sales,'excel_quote'),(manager,'excel_admin')]:
            task=api.mutate('/v1/outputs',dict(order_id=order['order_id'],expected_revision=1,kind=kind),Journal(host.root/(kind+'.pending')) if windows else MemoryJournal())
            deadline=time.monotonic()+20
            while True:
                state=api.request('GET','/v1/outputs/'+task['job_id'])['state']
                if state=='done':break
                if state=='failed' or time.monotonic()>deadline:raise RuntimeError('NATIVE_OUTPUT_NOT_DONE')
                time.sleep(.02)
            content=api.request('GET','/v1/outputs/'+task['job_id']+'/content');data=base64.b64decode(content['data_base64'],validate=True)
            assert hashlib.sha256(data).hexdigest()==content['sha256']
            book=load_workbook(io.BytesIO(data))
            try:
                fields=dict(list(book['Order'].values)[1:]);assert fields['order_ref']=='00003' and fields['income_cents']=='4000'
                assert ('cost_cents' in fields)==(kind=='excel_admin')
                assert all(cell.data_type!='f' and cell.number_format=='@' for ws in book for row in ws for cell in row)
                assert all(ws.freeze_panes=='A2' and all(c.font.bold for c in ws[1]) for ws in book)
            finally:book.close()
            tasks.append(task);contents.append(content)
        try:sales.request('GET','/v1/outputs/'+tasks[1]['job_id']+'/content')
        except APIError as e:assert e.status in (403,404)
        else:raise AssertionError('private output leaked')
        published='777-'+order['order_id']+'-'+tasks[0]['job_id']+'.xlsx'
        private_name='777-'+order['order_id']+'-'+tasks[1]['job_id']+'.xlsx'
        assert (export/published).read_bytes()==base64.b64decode(contents[0]['data_base64'])
        assert not (export/private_name).exists()
        before=host.service.ledger.inventory();assert before['counts']['orders']==3 and before['counts']['lines']==4 and before['counts']['notes']==1
        assert [before[k] for k in ['active_income_cents','active_cost_cents','active_profit_cents']]==['11500','4600','6900']
        saved=host.service.backup(kind='maintenance');target=host.root/'restored';ledger,auth=restore(saved['service_bundle'],target,'manager','new-smoke-password-123',auxiliary=not windows)
        _,_,reopened=load(target,allow_maintenance=True,auxiliary=not windows);assert ledger.inventory()==before and reopened.ledger.format_backend=='rust-excel-v1'
        for task,content in zip(tasks,contents):assert reopened.outputs.content(host.admin,task['job_id'])==content
        assert ledger.query_request(imported['request_id'],host.admin)==imported
        try:auth.authenticate(sales.token)
        except APIError:pass
        else:raise AssertionError('old session survived restore')
        result=dict(status='passed',backend='rust',excel_backend='rust-excel-v1',source_hash_and_mtime_unchanged=True,complete_native_and_reference_import_plan_equal=True,native_repeat_batch_single_commit=True,DPAPI_native_unknown_commit_and_original_receipt_recovered=windows,actual_HTTPS_original_clients_native_excel=True,independent_reader_all_text_money_front_zero_and_styles=True,private_cost_output_restricted=True,public_directory_only_quote_excel=True,v3_native_excel_files_receipts_and_ledger_restore_equal=True,old_sessions_removed_on_restore=True,inventory=before,host_private_directory_native_security=host.config['private_security'],scope='Windows temp host/Rust workbook processes/old HTTPS clients/DPAPI/import/new XLSX/v3 restore; no GUI, installed service, real E/share or production switch')
    finally:host.close()
    result['sandbox_cleaned']=True
    if not windows:result['scope']='Auxiliary native workbook/real HTTPS/import/output/backup protocol rehearsal with memory journal; no Windows DPAPI/ACL, GUI, installed service, real E/share or production claims'
    return result

def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__);mode=parser.add_mutually_exclusive_group(required=True);mode.add_argument('--windows-smoke',action='store_true');mode.add_argument('--checks-only',action='store_true');parser.add_argument('--output',required=True)
    args=parser.parse_args(argv);output=new_output(args.output)
    report=dict(task='P06-03',status='running',environment=dict(os=platform.system(),python=platform.python_version()),source_commit=None,implementation_sha256_LF=None,commands=[],checks=None,shadow_comparison=None,native_smoke=None,production_paths_accessed=False,production_writer_replaced=False,Windows_acceptance='pending')
    try:
        if args.windows_smoke and (os.name!='nt' or platform.system()!='Windows'):raise RuntimeError('ACTUAL_WINDOWS_REQUIRED')
        report['source_commit']=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True,encoding='utf-8').strip();report['worktree_dirty']=bool(subprocess.check_output(['git','status','--porcelain'],cwd=ROOT,text=True,encoding='utf-8').strip());report['implementation_sha256_LF']=fingerprints()
        cargo=shutil.which('cargo');rustc=shutil.which('rustc')
        if not cargo or not rustc:raise RuntimeError('PINNED_RUST_AND_MSVC_REQUIRED')
        v=run([rustc,'+1.97.1','-vV']);report['commands'].append(v)
        if v['exit_code'] or not v['stdout'].startswith('rustc 1.97.1 '):raise RuntimeError('PINNED_TOOLCHAIN_REQUIRED')
        if args.windows_smoke and 'host: x86_64-pc-windows-msvc' not in v['stdout']:raise RuntimeError('NATIVE_MSVC_REQUIRED')
        report['environment']['rustc']=v['stdout'].strip()
        for tail in [['fmt','--all','--','--check'],['clippy','--locked','--workspace','--all-targets','--','-D','warnings'],['test','--locked','--workspace'],['build','--locked','--release','--workspace']]:
            record=run([cargo,'+1.97.1',*tail]);report['commands'].append(record)
            if record['exit_code']:raise RuntimeError('RUST_CHECK_FAILED')
        meta=run([cargo,'+1.97.1','metadata','--locked','--format-version','1']);packages=json.loads(meta['stdout'])['packages']
        external=[dict(name=v['name'],version=v['version'],license=v['license']) for v in packages if v['source'] is not None];expected=json.loads((ROOT/'docs/verification/P06-03-dependencies.json').read_text(encoding='utf-8'))['packages']
        if {(v['name'],v['version'],v['license']) for v in external}!={(v['name'],v['version'],v['license']) for v in expected}:raise RuntimeError('DEPENDENCY_LICENSE_MISMATCH')
        report['dependencies']=external
        from p05_service_check import dependencies
        report['service_dependencies']=dependencies()
        if any(v['status']!='available' for v in report['service_dependencies']):raise RuntimeError('EXISTING_PINNED_SERVICE_DEPENDENCIES_REQUIRED')
        report['checks']=run_checks()
        import legacy_worker_fixture
        report['legacy_worker_build']=legacy_worker_fixture.build_record
        report['shadow_comparison']=compare(ROOT/'target/release'/('quote-shadow.exe' if os.name=='nt' else 'quote-shadow'),output)
        worker=ROOT/'target/release'/('quote-ledger.exe' if os.name=='nt' else 'quote-ledger');report['worker_sha256']=hashlib.sha256(worker.read_bytes()).hexdigest()
        if report['checks']['status']!='passed' or report['shadow_comparison']['status']!='passed':raise RuntimeError('FUNCTIONAL_CHECK_FAILED')
        if args.windows_smoke:report['native_smoke']=windows_smoke();report['Windows_acceptance']='isolated_native_excel_import_output_restore_passed_user_acceptance_pending'
        report['status']='Windows_native_excel_passed_acceptance_pending' if args.windows_smoke else 'auxiliary_native_excel_passed_Windows_pending'
    except Exception as e:report['status']='failed';report['failure']=dict(reason=str(e),traceback=traceback.format_exc())
    (output/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    summary={k:report[k] for k in ['task','status','environment','Windows_acceptance','native_smoke','production_paths_accessed','production_writer_replaced']}
    for key in ['checks','shadow_comparison']:
        v=report[key];summary[key]=None if v is None else {k:x for k,x in v.items() if k not in ['log','test_ids']}
    if 'failure' in report:summary['failure']=report['failure']['reason']
    summary['report_path']=str(output/'report.json');print(json.dumps(summary,ensure_ascii=False,indent=2));return 0 if report['status'].endswith(('acceptance_pending','Windows_pending')) else 1
if __name__=='__main__':raise SystemExit(main())
