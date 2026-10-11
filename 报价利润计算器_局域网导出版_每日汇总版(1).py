# -*- coding: utf-8 -*-
"""
六六派对报价 / 利润计算器 - 局域网导出版

本版重点：
1. 价格设定不强制填写；空白项目行会被忽略，真正计算时才集中校验。
2. 价格设定只负责基础项目：项目名称、销售单价、成本单价、K系数、是否单件。
3. “折扣”改为计算主页里的独立项目类型，不再放在价格设定中。
4. 添加“折扣”行后，可选择 95%、90%、85%、80%、75%、70%、50% 或自定义百分比。
5. 折扣行不参与高度、宽度、数量、销售单价、成本单价和 K 系数计算，只按它前面的当前小计折减销售额。
6. 发现输入错误时，弹窗提醒并自动切换到“输入反馈”页面。
7. Excel 查看页按“淘宝ID + 记录时间”分组删除，备注保存到所选分组底部。
8. 计算主页顶部改为淘宝ID，并新增订单号、微信号。
9. Excel 导出目录支持本地文件夹、映射盘和局域网共享路径。

依赖：
- tkinter：Python 标准库自带
- openpyxl：用于导出 Excel，安装命令：pip install openpyxl
- pillow：用于截图，安装命令：pip install pillow
"""

import os
import re
import time
import tkinter as tk
from quote_app import QuoteApplicationMixin
from tkinter import filedialog, messagebox, ttk

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None
    load_workbook = None


# ======================== 固定保存位置 ========================
# Excel 每天共用一个文件，统一保存到：E:\14T共享盘\葛宁\PyInstaller\利润表
# 价格设定 config.json 保存到：E:\14T共享盘\葛宁\PyInstaller\config.json，子版本会读取这个文件。
APP_SAVE_ROOT = r"E:\14T共享盘\葛宁\PyInstaller"
DEFAULT_REPORT_DIR = os.path.join(APP_SAVE_ROOT, "利润表")
DEFAULT_CONFIG_FILE = os.path.join(APP_SAVE_ROOT, "config.json")
DEFAULT_SCREENSHOT_DIR = os.path.join(APP_SAVE_ROOT, "截图")
DAILY_EXCEL_PREFIX = "利润明细"
# =============================================================


class ScrollableFrame(ttk.Frame):
    def __init__(self, parent, *args, **kwargs):
        super().__init__(parent, *args, **kwargs)
        self.canvas = tk.Canvas(self, highlightthickness=0)
        self.scrollbar = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.inner = ttk.Frame(self.canvas)

        self.inner.bind(
            "<Configure>",
            lambda event: self.canvas.configure(scrollregion=self.canvas.bbox("all")),
        )

        self.window_id = self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.canvas.configure(yscrollcommand=self.scrollbar.set)
        self.canvas.pack(side="left", fill="both", expand=True)
        self.scrollbar.pack(side="right", fill="y")

        self.canvas.bind("<Configure>", self._on_canvas_configure)
        self.canvas.bind_all("<MouseWheel>", self._on_mousewheel)

    def _on_canvas_configure(self, event):
        self.canvas.itemconfigure(self.window_id, width=event.width)

    def _on_mousewheel(self, event):
        try:
            widget = self.winfo_containing(event.x_root, event.y_root)
            if widget is not None and str(widget).startswith(str(self)):
                self.canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")
        except Exception:
            pass


class TableRowHandle:
    """计算列表表格行句柄。

    表头和所有数据行都直接放在同一个 grid 父容器里，
    普通项目行、折扣行、表头会共享同一组列宽，避免列线错位。
    """

    def __init__(self, parent, row_index):
        self.parent = parent
        self.row_index = row_index
        self.cells = []

    def add_cell(self, cell):
        self.cells.append(cell)

    def destroy(self):
        for cell in list(self.cells):
            try:
                cell.destroy()
            except Exception:
                pass
        self.cells = []
        try:
            self.parent.grid_rowconfigure(self.row_index, minsize=0)
        except Exception:
            pass


class PriceCalculator(QuoteApplicationMixin):
    quote_role = "manager"
    SPECIAL_DISCOUNT_TYPE = "折扣"
    DISCOUNT_CHOICES = ["95%", "90%", "85%", "80%", "75%", "70%", "50%", "自定义"]
    CALC_COLUMNS = [
        # 标题、最小像素宽、扩展权重。所有列在同一个 grid 中统一分配宽度。
        ("项目名称", 250, 3),
        ("高(m)", 110, 1),
        ("宽(m)", 110, 1),
        ("数量", 90, 1),
        ("折扣/说明", 180, 2),
        ("销售额小计", 130, 1),
        ("状态", 95, 1),
        ("操作", 90, 1),
    ]

    def __init__(self, root):
        self.root = root
        self.root.title("六六派对报价 / 利润计算器")
        self.root.geometry("1180x840")
        self.root.minsize(1040, 720)

        self.config_file = DEFAULT_CONFIG_FILE
        self.default_report_dir = DEFAULT_REPORT_DIR
        self.report_dir = self.default_report_dir
        self.image_dir = DEFAULT_SCREENSHOT_DIR
        self.ensure_directories()

        self.default_items = [
            {"name": "背景板", "sale": "11.5", "cost": "5.0", "k": "1.0", "single": False},
            {"name": "喷绘", "sale": "15.0", "cost": "7.0", "k": "1.0", "single": False},
            {"name": "KT板", "sale": "12.5", "cost": "11.5", "k": "2.2", "single": False},
            {"name": "支架", "sale": "50.0", "cost": "20.0", "k": "1.0", "single": True},
            {"name": "设计费", "sale": "100.0", "cost": "0.0", "k": "1.0", "single": True},
            {"name": "配套气球", "sale": "30.0", "cost": "15.0", "k": "1.0", "single": True},
            {"name": "其他", "sale": "1.0", "cost": "1.0", "k": "1.0", "single": False},
        ]

        self.setting_rows = []
        self.calc_rows = []
        self.calc_table_next_row = 1
        self.last_export_records = []
        self.validation_issues = []
        self.last_config_snapshot = {}
        self.current_excel_path = None
        self.excel_view_row_map = {}

        self.user_name = tk.StringVar(value="默认淘宝ID")  # 兼容旧代码命名：这里实际表示“淘宝ID”
        self.order_no = tk.StringVar(value="")
        self.wechat_no = tk.StringVar(value="")
        self.feedback_var = tk.StringVar(value="提示：价格设定可以留空；如需整单折扣，请在计算主页选择项目类型“折扣”单独添加。")
        self.issue_summary_var = tk.StringVar(value="暂无输入问题。")
        self.total_val = tk.StringVar(value="总计：0.00 元")
        self.summary_income = tk.StringVar(value="0.00")
        self.summary_cost = tk.StringVar(value="0.00")
        self.summary_profit = tk.StringVar(value="0.00")
        self.summary_margin = tk.StringVar(value="0%")
        self.excel_total_income = tk.StringVar(value="0.00")
        self.excel_total_cost = tk.StringVar(value="0.00")
        self.excel_total_profit = tk.StringVar(value="0.00")
        self.excel_total_margin = tk.StringVar(value="0%")
        self.excel_total_count = tk.StringVar(value="0 个文件")
        self.excel_selected_file = tk.StringVar(value="未选择 Excel 文件")
        self.excel_sheet_var = tk.StringVar(value="")
        self.export_dir_var = tk.StringVar(value=self.report_dir)

        self.initialize_quote_state()
        self._build_style()
        self._build_layout()
        self.load_settings_to_ui()
        self.refresh_item_selector()

        self.install_quote_controls()
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

    # ---------- 基础工具 ----------

    def _build_style(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except Exception:
            pass
        style.configure("Title.TLabel", font=("Arial", 15, "bold"))
        style.configure("Section.TLabelframe.Label", font=("Arial", 10, "bold"))
        style.configure("Feedback.TLabel", font=("Arial", 10))
        style.configure("Total.TLabel", font=("Arial", 14, "bold"))
        style.configure("Accent.TButton", font=("Arial", 10, "bold"))


    def normalize_path_text(self, text):
        path = str(text or "").strip().strip('"').strip("'")
        path = os.path.expandvars(os.path.expanduser(path))
        return path or getattr(self, "default_report_dir", "利润表")

    def ensure_directories(self):
        for folder in [self.report_dir, self.image_dir]:
            if not os.path.exists(folder):
                os.makedirs(folder)

    def format_display(self, value):
        return f"{round(float(value), 2):.2f}"

    def clean_filename(self, text):
        return re.sub(r"[\\/:*?\"<>|]", "_", str(text).strip() or "默认淘宝ID")

    def set_feedback(self, text, is_error=False):
        self.feedback_var.set(text)
        if hasattr(self, "feedback_label"):
            self.feedback_label.configure(foreground="#c0392b" if is_error else "#2c3e50")

    def add_issue(self, issues, level, location, field, raw_value, message):
        issues.append(
            {
                "级别": level,
                "位置": location,
                "字段": field,
                "输入内容": str(raw_value),
                "提示": message,
            }
        )


    # ---------- 界面 ----------

    def _build_layout(self):
        top = ttk.Frame(self.root, padding=(14, 10, 14, 6))
        top.pack(fill=tk.X)

        ttk.Label(top, text="六六派对报价 / 利润计算器", style="Title.TLabel").pack(side=tk.LEFT)

        user_box = ttk.Frame(top)
        user_box.pack(side=tk.RIGHT)
        ttk.Label(user_box, text="淘宝ID：").grid(row=0, column=0, sticky=tk.E, padx=(0, 4))
        ttk.Entry(user_box, textvariable=self.user_name, width=14).grid(row=0, column=1, sticky=tk.W, padx=(0, 8))
        ttk.Label(user_box, text="订单号：").grid(row=0, column=2, sticky=tk.E, padx=(0, 4))
        ttk.Entry(user_box, textvariable=self.order_no, width=15).grid(row=0, column=3, sticky=tk.W, padx=(0, 8))
        ttk.Label(user_box, text="微信号：").grid(row=0, column=4, sticky=tk.E, padx=(0, 4))
        ttk.Entry(user_box, textvariable=self.wechat_no, width=14).grid(row=0, column=5, sticky=tk.W)

        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill=tk.BOTH, expand=True, padx=12, pady=(0, 8))

        self.calc_page = ttk.Frame(self.notebook)
        self.feedback_page = ttk.Frame(self.notebook)
        self.profit_page = ttk.Frame(self.notebook)
        self.excel_page = ttk.Frame(self.notebook)
        self.settings_page = ttk.Frame(self.notebook)

        self.notebook.add(self.calc_page, text=" 计算主页 ")
        self.notebook.add(self.feedback_page, text=" 输入反馈 ")
        self.notebook.add(self.profit_page, text=" 利润分析 ")
        self.notebook.add(self.excel_page, text=" Excel记录 ")
        self.notebook.add(self.settings_page, text=" 价格设定 ")

        self.setup_calc_page()
        self.setup_feedback_page()
        self.setup_profit_page()
        self.setup_excel_page()
        self.setup_settings_page()

    def setup_calc_page(self):
        add_frame = ttk.LabelFrame(self.calc_page, text="添加项目", padding=10, style="Section.TLabelframe")
        add_frame.pack(fill=tk.X, padx=10, pady=(10, 6))

        ttk.Label(add_frame, text="品名备注：").grid(row=0, column=0, sticky=tk.W, padx=(0, 6))
        self.custom_item_entry = ttk.Entry(add_frame, width=22)
        self.custom_item_entry.grid(row=0, column=1, sticky=tk.W, padx=(0, 14))

        ttk.Label(add_frame, text="项目类型：").grid(row=0, column=2, sticky=tk.W, padx=(0, 6))
        self.item_selector = ttk.Combobox(add_frame, width=24, state="readonly")
        self.item_selector.grid(row=0, column=3, sticky=tk.W, padx=(0, 14))

        ttk.Button(add_frame, text="+ 添加到列表", command=self.handle_add_item, style="Accent.TButton").grid(row=0, column=4, sticky=tk.W)
        ttk.Button(add_frame, text="应用价格设定", command=self.apply_settings_to_calc).grid(row=0, column=5, sticky=tk.W, padx=(10, 0))

        tip = "价格设定留空不会马上报错；计算时会检查被使用项目。折扣请作为“折扣”项目单独添加，按它前面的当前小计折减。"
        ttk.Label(add_frame, text=tip, foreground="#7f8c8d").grid(row=1, column=0, columnspan=6, sticky=tk.W, pady=(8, 0))

        list_frame = ttk.LabelFrame(self.calc_page, text="计算列表", padding=8, style="Section.TLabelframe")
        list_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=6)

        self.calc_scroll = ScrollableFrame(list_frame)
        self.calc_scroll.pack(fill=tk.BOTH, expand=True)
        self.calc_table_next_row = 1
        self.build_calc_header(self.calc_scroll.inner)

        bottom = ttk.Frame(self.calc_page, padding=(10, 4, 10, 10))
        bottom.pack(fill=tk.X)

        buttons = ttk.Frame(bottom)
        buttons.pack(side=tk.LEFT)
        ttk.Button(buttons, text="开始计算", command=self.calculate, style="Accent.TButton").pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(buttons, text="查看输入反馈", command=self.go_to_feedback_page).pack(side=tk.LEFT, padx=6)
        ttk.Button(buttons, text="清空所有行", command=self.reset).pack(side=tk.LEFT, padx=6)
        ttk.Button(buttons, text="保存计算截图", command=self.save_screenshot).pack(side=tk.LEFT, padx=6)

        ttk.Label(bottom, textvariable=self.total_val, style="Total.TLabel", foreground="#2c3e50").pack(side=tk.RIGHT)

        feedback_box = ttk.Frame(self.calc_page, padding=(12, 0, 12, 12))
        feedback_box.pack(fill=tk.X)
        self.feedback_label = ttk.Label(
            feedback_box,
            textvariable=self.feedback_var,
            style="Feedback.TLabel",
            wraplength=1100,
            foreground="#2c3e50",
        )
        self.feedback_label.pack(fill=tk.X)

    def setup_feedback_page(self):
        container = ttk.Frame(self.feedback_page, padding=14)
        container.pack(fill=tk.BOTH, expand=True)

        top = ttk.Frame(container)
        top.pack(fill=tk.X, pady=(0, 8))
        ttk.Label(top, text="输入校验反馈", style="Title.TLabel").pack(side=tk.LEFT)
        ttk.Button(top, text="返回计算主页", command=lambda: self.notebook.select(self.calc_page)).pack(side=tk.RIGHT, padx=(6, 0))
        ttk.Button(top, text="重新计算", command=self.calculate).pack(side=tk.RIGHT, padx=(6, 0))

        ttk.Label(
            container,
            text="这里集中显示计算时发现的问题。出现“错误”时，相关行不会进入利润明细和 Excel 明细；“提醒”表示系统仍会计算，但建议人工确认。",
            foreground="#7f8c8d",
            wraplength=1080,
        ).pack(fill=tk.X, pady=(0, 8))

        summary_box = ttk.LabelFrame(container, text="反馈概览", padding=10, style="Section.TLabelframe")
        summary_box.pack(fill=tk.X, pady=(0, 10))
        ttk.Label(summary_box, textvariable=self.issue_summary_var, font=("Arial", 11, "bold"), foreground="#2c3e50").pack(anchor=tk.W)

        tree_frame = ttk.Frame(container)
        tree_frame.pack(fill=tk.BOTH, expand=True)

        issue_columns = ("level", "location", "field", "value", "message")
        self.issue_tree = ttk.Treeview(tree_frame, columns=issue_columns, show="headings", height=18)
        issue_headings = {
            "level": "级别",
            "location": "位置",
            "field": "字段",
            "value": "输入内容",
            "message": "提示",
        }
        issue_widths = {"level": 80, "location": 190, "field": 140, "value": 240, "message": 540}
        for col in issue_columns:
            self.issue_tree.heading(col, text=issue_headings[col])
            self.issue_tree.column(col, width=issue_widths[col], anchor=tk.W)

        y_scroll = ttk.Scrollbar(tree_frame, orient="vertical", command=self.issue_tree.yview)
        x_scroll = ttk.Scrollbar(tree_frame, orient="horizontal", command=self.issue_tree.xview)
        self.issue_tree.configure(yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)
        self.issue_tree.grid(row=0, column=0, sticky="nsew")
        y_scroll.grid(row=0, column=1, sticky="ns")
        x_scroll.grid(row=1, column=0, sticky="ew")
        tree_frame.rowconfigure(0, weight=1)
        tree_frame.columnconfigure(0, weight=1)

        self.issue_tree.tag_configure("错误", foreground="#c0392b")
        self.issue_tree.tag_configure("提醒", foreground="#8e44ad")

    def setup_profit_page(self):
        container = ttk.Frame(self.profit_page, padding=14)
        container.pack(fill=tk.BOTH, expand=True)

        ttk.Label(container, text="项目利润明细表", style="Title.TLabel").pack(anchor=tk.W, pady=(0, 8))

        tree_frame = ttk.Frame(container)
        tree_frame.pack(fill=tk.BOTH, expand=True)

        columns = ("item", "discount", "income", "cost", "profit", "margin")
        self.profit_tree = ttk.Treeview(tree_frame, columns=columns, show="headings", height=14)
        headings = {"item": "项目", "discount": "折扣/说明", "income": "销售额", "cost": "成本", "profit": "利润", "margin": "利润率"}
        widths = {"item": 320, "discount": 100, "income": 130, "cost": 130, "profit": 130, "margin": 110}
        for col in columns:
            self.profit_tree.heading(col, text=headings[col])
            self.profit_tree.column(col, width=widths[col], anchor=tk.CENTER)
        self.profit_tree.column("item", anchor=tk.W)

        y_scroll = ttk.Scrollbar(tree_frame, orient="vertical", command=self.profit_tree.yview)
        self.profit_tree.configure(yscrollcommand=y_scroll.set)
        self.profit_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        y_scroll.pack(side=tk.RIGHT, fill=tk.Y)

        summary = ttk.LabelFrame(container, text="全局汇总", padding=10, style="Section.TLabelframe")
        summary.pack(fill=tk.X, pady=12)

        stats = [("总收入", self.summary_income), ("总成本", self.summary_cost), ("总利润", self.summary_profit), ("综合利润率", self.summary_margin)]
        for idx, (name, var) in enumerate(stats):
            card = ttk.Frame(summary)
            card.grid(row=0, column=idx, sticky="ew", padx=8)
            summary.columnconfigure(idx, weight=1)
            ttk.Label(card, text=name).pack()
            ttk.Label(card, textvariable=var, font=("Arial", 14, "bold")).pack()

        ttk.Button(container, text="导出当前全部明细到 Excel", command=self.export_to_excel, style="Accent.TButton").pack(anchor=tk.E)

    def setup_settings_page(self):
        wrapper = ttk.Frame(self.settings_page, padding=10)
        wrapper.pack(fill=tk.BOTH, expand=True)

        top = ttk.Frame(wrapper)
        top.pack(fill=tk.X, pady=(0, 8))
        ttk.Label(top, text="价格设定", style="Title.TLabel").pack(side=tk.LEFT)
        ttk.Button(top, text="+ 新增自定义项目", command=self.add_setting_row).pack(side=tk.RIGHT, padx=(6, 0))
        ttk.Button(top, text="保存并应用", command=self.save_and_apply_settings, style="Accent.TButton").pack(side=tk.RIGHT)

        desc = "价格设定不强制填写：空白项目行会自动忽略。折扣不在这里设置，请在计算主页选择项目类型“折扣”单独添加。"
        ttk.Label(wrapper, text=desc, foreground="#7f8c8d", wraplength=1080).pack(fill=tk.X, pady=(0, 8))

        header = ttk.Frame(wrapper)
        header.pack(fill=tk.X)
        header_cols = [
            ("项目名称", 22),
            ("销售单价", 16),
            ("成本单价", 16),
            ("K系数", 12),
            ("是否单件", 10),
            ("操作", 8),
        ]
        for text, width in header_cols:
            ttk.Label(header, text=text, width=width, font=("Arial", 10, "bold")).pack(side=tk.LEFT, padx=4)

        self.settings_scroll = ScrollableFrame(wrapper)
        self.settings_scroll.pack(fill=tk.BOTH, expand=True, pady=(4, 0))


    def get_calc_column(self, col_index):
        column = self.CALC_COLUMNS[col_index]
        if len(column) == 2:
            title, width = column
            weight = 1
        else:
            title, width, weight = column
        return title, width, weight

    def build_calc_header(self, parent):
        """表头和数据行共用同一个 grid 父容器，确保列线严格对齐。"""
        for col_index, _column in enumerate(self.CALC_COLUMNS):
            title, width, weight = self.get_calc_column(col_index)
            parent.grid_columnconfigure(col_index, minsize=width, weight=weight, uniform="calc_table")
            cell = tk.Frame(parent, height=32, bd=1, relief="groove", bg="#eef3f7")
            cell.grid(row=0, column=col_index, sticky="nsew", padx=0, pady=(0, 2))
            cell.grid_propagate(False)
            label = tk.Label(cell, text=title, bg="#eef3f7", fg="#2c3e50", font=("Arial", 10, "bold"))
            label.pack(fill=tk.BOTH, expand=True)
        parent.grid_rowconfigure(0, minsize=32)

    def create_calc_row_frame(self, kind="normal"):
        bg = "#ffffff" if kind == "normal" else "#fffaf2"
        row_handle = TableRowHandle(self.calc_scroll.inner, self.calc_table_next_row)
        self.calc_table_next_row += 1
        self.calc_scroll.inner.grid_rowconfigure(row_handle.row_index, minsize=44)
        return row_handle, bg

    def create_calc_cell(self, frame, col_index, bg="#ffffff"):
        _title, _width, _weight = self.get_calc_column(col_index)
        # 直接把每一个单元格 grid 到 calc_scroll.inner，所有行共享一套列宽。
        cell = tk.Frame(self.calc_scroll.inner, height=42, bd=1, relief="sunken", bg=bg)
        cell.grid(row=frame.row_index, column=col_index, sticky="nsew", padx=0, pady=1)
        cell.grid_propagate(False)
        frame.add_cell(cell)
        return cell

    def pack_in_cell(self, parent, widget, padx=6, pady=5):
        widget.pack(fill=tk.BOTH, expand=True, padx=padx, pady=pady)
        return widget

    def go_to_feedback_page(self):
        if hasattr(self, "feedback_page"):
            self.notebook.select(self.feedback_page)

    def notify_validation_errors(self, error_count, warn_count):
        self.go_to_feedback_page()
        messagebox.showwarning(
            "发现输入错误",
            f"本次计算发现 {error_count} 个错误、{warn_count} 个提醒。\n\n"
            "已切换到“输入反馈”页面，请根据提示修改后重新计算。",
        )

    # ---------- 设置数据 ----------


    def add_setting_row(self, name="", sale="", cost="", k="", single=False):
        frame = ttk.Frame(self.settings_scroll.inner)
        frame.pack(fill=tk.X, pady=3)

        name_var = tk.StringVar(value=name)
        sale_var = tk.StringVar(value=sale)
        cost_var = tk.StringVar(value=cost)
        k_var = tk.StringVar(value=k)
        single_var = tk.BooleanVar(value=single)

        ttk.Entry(frame, textvariable=name_var, width=24).pack(side=tk.LEFT, padx=4)
        ttk.Entry(frame, textvariable=sale_var, width=18).pack(side=tk.LEFT, padx=4)
        ttk.Entry(frame, textvariable=cost_var, width=18).pack(side=tk.LEFT, padx=4)
        ttk.Entry(frame, textvariable=k_var, width=14).pack(side=tk.LEFT, padx=4)
        ttk.Checkbutton(frame, variable=single_var).pack(side=tk.LEFT, padx=(28, 26))
        ttk.Button(frame, text="删除", width=6, command=lambda: self.remove_setting_row(frame)).pack(side=tk.LEFT, padx=4)

        row = {
            "frame": frame,
            "name": name_var,
            "sale": sale_var,
            "cost": cost_var,
            "k": k_var,
            "single": single_var,
        }
        self.setting_rows.append(row)
        self.watch_quote_row(row, settings=True)

    def update_discount_editor_state(self, row):
        choice = row["discount_choice"].get()
        if choice == "自定义":
            row["custom_entry"].configure(state="normal")
            if not row["discount_custom"].get().strip():
                row["discount_custom"].set("90")
        else:
            row["custom_entry"].configure(state="disabled")

    def remove_setting_row(self, frame):
        for idx, row in enumerate(self.setting_rows):
            if row["frame"] == frame:
                row["frame"].destroy()
                self.setting_rows.pop(idx)
                self.quote_input_changed()
                break
        self.refresh_item_selector()


    def refresh_item_selector(self):
        config, _issues = self.get_settings_snapshot()
        values = list(config.keys()) + [self.SPECIAL_DISCOUNT_TYPE]
        self.item_selector.configure(values=values)
        current = self.item_selector.get().strip()
        if values and current not in values:
            self.item_selector.set(values[0])
        elif not values:
            self.item_selector.set("")


    # ---------- 计算主页 ----------

    def handle_add_item(self):
        config, issues = self.get_settings_snapshot()
        self.refresh_item_selector()
        base_type = self.item_selector.get().strip()
        prefix = self.custom_item_entry.get().strip()

        if base_type == self.SPECIAL_DISCOUNT_TYPE:
            display_name = f"{prefix}（折扣）" if prefix else "整单折扣"
            self.add_discount_calc_row(display_name)
            self.custom_item_entry.delete(0, tk.END)
            self.render_issues(issues)
            self.set_feedback("已添加折扣行：折扣行按它前面的当前小计计算，建议放在需要折扣的项目后面。", is_error=False)
            return

        if not config:
            self.render_issues(issues)
            self.set_feedback("请先在价格设定中新增至少一个有效项目名称，或直接选择“折扣”添加折扣行。", is_error=True)
            return

        if base_type not in config:
            self.add_issue(issues, "错误", "添加项目", "项目类型", base_type, "请选择一个有效的项目类型。")
            self.render_issues(issues)
            self.set_feedback("请选择一个有效的项目类型。", is_error=True)
            return

        display_name = f"{prefix}（{base_type}）" if prefix else base_type
        self.add_calc_row(display_name, base_type, config[base_type].get("single", False))
        self.custom_item_entry.delete(0, tk.END)

        if issues:
            self.render_issues(issues)
            self.set_feedback(f"项目已添加，但价格设定有 {len(issues)} 条提醒。", is_error=True)
        else:
            self.set_feedback(f"已添加：{display_name}", is_error=False)

    def add_calc_row(self, display_name, base_type, single=False):
        frame, bg = self.create_calc_row_frame("normal")

        h_var = tk.StringVar(value="" if single else "1.0")
        w_var = tk.StringVar(value="" if single else "1.0")
        q_var = tk.StringVar(value="1")
        note_var = tk.StringVar(value="—")
        result_var = tk.StringVar(value="0.00")
        status_var = tk.StringVar(value="未计算")

        name_cell = self.create_calc_cell(frame, 0, bg)
        tk.Label(name_cell, text=display_name, anchor="w", bg=bg, fg="#2c3e50").pack(fill=tk.BOTH, expand=True, padx=8)

        h_cell = self.create_calc_cell(frame, 1, bg)
        w_cell = self.create_calc_cell(frame, 2, bg)
        q_cell = self.create_calc_cell(frame, 3, bg)
        h_entry = ttk.Entry(h_cell, textvariable=h_var, justify="center")
        w_entry = ttk.Entry(w_cell, textvariable=w_var, justify="center")
        q_entry = ttk.Entry(q_cell, textvariable=q_var, justify="center")
        self.pack_in_cell(h_cell, h_entry)
        self.pack_in_cell(w_cell, w_entry)
        self.pack_in_cell(q_cell, q_entry)

        note_cell = self.create_calc_cell(frame, 4, bg)
        tk.Label(note_cell, textvariable=note_var, anchor="center", bg=bg, fg="#7f8c8d").pack(fill=tk.BOTH, expand=True)

        result_cell = self.create_calc_cell(frame, 5, bg)
        tk.Label(result_cell, textvariable=result_var, anchor="center", bg=bg, fg="#e67e22", font=("Arial", 10, "bold")).pack(fill=tk.BOTH, expand=True)

        status_cell = self.create_calc_cell(frame, 6, bg)
        tk.Label(status_cell, textvariable=status_var, anchor="center", bg=bg, fg="#2c3e50").pack(fill=tk.BOTH, expand=True)

        action_cell = self.create_calc_cell(frame, 7, bg)
        ttk.Button(action_cell, text="删除", width=6, command=lambda: self.remove_calc_row(frame)).pack(expand=True, padx=6, pady=4)

        row = {
            "row_kind": "normal",
            "frame": frame,
            "display_name": display_name,
            "base_type": base_type,
            "h": h_var,
            "w": w_var,
            "q": q_var,
            "discount": note_var,
            "res": result_var,
            "status": status_var,
            "h_entry": h_entry,
            "w_entry": w_entry,
        }
        self.calc_rows.append(row)
        self.watch_quote_row(row)
        self.set_row_dimension_state(row, single)

    def add_discount_calc_row(self, display_name):
        frame, bg = self.create_calc_row_frame("discount")

        h_var = tk.StringVar(value="")
        w_var = tk.StringVar(value="")
        q_var = tk.StringVar(value="")
        discount_choice_var = tk.StringVar(value="90%")
        discount_custom_var = tk.StringVar(value="90")
        result_var = tk.StringVar(value="0.00")
        status_var = tk.StringVar(value="未计算")

        name_cell = self.create_calc_cell(frame, 0, bg)
        tk.Label(name_cell, text=display_name, anchor="w", bg=bg, fg="#8a5a00", font=("Arial", 10, "bold")).pack(fill=tk.BOTH, expand=True, padx=8)

        h_cell = self.create_calc_cell(frame, 1, bg)
        w_cell = self.create_calc_cell(frame, 2, bg)
        q_cell = self.create_calc_cell(frame, 3, bg)
        h_entry = ttk.Entry(h_cell, textvariable=h_var, justify="center", state="disabled")
        w_entry = ttk.Entry(w_cell, textvariable=w_var, justify="center", state="disabled")
        q_entry = ttk.Entry(q_cell, textvariable=q_var, justify="center", state="disabled")
        self.pack_in_cell(h_cell, h_entry)
        self.pack_in_cell(w_cell, w_entry)
        self.pack_in_cell(q_cell, q_entry)

        discount_cell = self.create_calc_cell(frame, 4, bg)
        discount_box = tk.Frame(discount_cell, bg=bg)
        discount_box.pack(fill=tk.BOTH, expand=True, padx=5, pady=4)
        discount_combo = ttk.Combobox(
            discount_box,
            textvariable=discount_choice_var,
            values=self.DISCOUNT_CHOICES,
            width=7,
            state="readonly",
        )
        discount_combo.pack(side=tk.LEFT, fill=tk.X, expand=True)
        vcmd = (self.root.register(self.validate_discount_key), "%P")
        custom_entry = ttk.Entry(discount_box, textvariable=discount_custom_var, width=6, justify="center", validate="key", validatecommand=vcmd)
        custom_entry.pack(side=tk.LEFT, padx=(5, 2))
        tk.Label(discount_box, text="%", bg=bg, fg="#7f8c8d").pack(side=tk.LEFT)

        result_cell = self.create_calc_cell(frame, 5, bg)
        tk.Label(result_cell, textvariable=result_var, anchor="center", bg=bg, fg="#c0392b", font=("Arial", 10, "bold")).pack(fill=tk.BOTH, expand=True)

        status_cell = self.create_calc_cell(frame, 6, bg)
        tk.Label(status_cell, textvariable=status_var, anchor="center", bg=bg, fg="#2c3e50").pack(fill=tk.BOTH, expand=True)

        action_cell = self.create_calc_cell(frame, 7, bg)
        ttk.Button(action_cell, text="删除", width=6, command=lambda: self.remove_calc_row(frame)).pack(expand=True, padx=6, pady=4)

        row = {
            "row_kind": "discount",
            "frame": frame,
            "display_name": display_name,
            "base_type": self.SPECIAL_DISCOUNT_TYPE,
            "h": h_var,
            "w": w_var,
            "q": q_var,
            "discount": discount_choice_var,
            "discount_choice": discount_choice_var,
            "discount_custom": discount_custom_var,
            "custom_entry": custom_entry,
            "res": result_var,
            "status": status_var,
            "h_entry": h_entry,
            "w_entry": w_entry,
        }
        self.calc_rows.append(row)
        self.watch_quote_row(row)
        discount_combo.bind("<<ComboboxSelected>>", lambda _event, r=row: self.update_discount_editor_state(r))
        custom_entry.bind("<FocusOut>", lambda _event, v=discount_custom_var: self.normalize_discount_entry(v))
        self.update_discount_editor_state(row)

    def set_row_dimension_state(self, row, single):
        if row.get("row_kind") == "discount":
            return
        state = "disabled" if single else "normal"
        row["h_entry"].configure(state=state)
        row["w_entry"].configure(state=state)
        if single:
            row["h"].set("")
            row["w"].set("")

    def update_calc_row_states(self, config):
        for row in self.calc_rows:
            if row.get("row_kind") == "discount":
                continue
            item = config.get(row["base_type"])
            if item is not None:
                self.set_row_dimension_state(row, item.get("single", False))
                row["discount"].set("—")
            else:
                row["discount"].set("需检查")

    def remove_calc_row(self, frame):
        for idx, row in enumerate(self.calc_rows):
            if row["frame"] == frame:
                row["frame"].destroy()
                self.calc_rows.pop(idx)
                break
        self.quote_input_changed()


    def render_issues(self, issues):
        if not hasattr(self, "issue_tree"):
            return

        for item in self.issue_tree.get_children():
            self.issue_tree.delete(item)

        for issue in issues:
            level = issue["级别"]
            self.issue_tree.insert(
                "",
                tk.END,
                values=(level, issue["位置"], issue["字段"], issue["输入内容"], issue["提示"]),
                tags=(level,),
            )

        error_count = sum(1 for issue in issues if issue["级别"] == "错误")
        warn_count = len(issues) - error_count
        if error_count:
            self.issue_summary_var.set(f"发现 {error_count} 个错误、{warn_count} 个提醒。请先处理错误；整单没有有效金额，禁止正式导出和截图。")
            self.notebook.tab(self.feedback_page, text=f" 输入反馈 ({error_count}) ")
        elif warn_count:
            self.issue_summary_var.set(f"发现 {warn_count} 个提醒。系统仍会计算，但建议人工确认。")
            self.notebook.tab(self.feedback_page, text=f" 输入反馈 ({warn_count}) ")
        else:
            self.issue_summary_var.set("暂无输入问题。")
            self.notebook.tab(self.feedback_page, text=" 输入反馈 ")


    def reset(self):
        for row in self.calc_rows:
            row["frame"].destroy()
        self.calc_rows = []
        self.quote_input_changed()
        self.calc_table_next_row = 1
        self.last_export_records = []
        self.validation_issues = []
        for item in self.profit_tree.get_children():
            self.profit_tree.delete(item)
        self.render_issues([])
        self.total_val.set("总计：0.00 元")
        self.summary_income.set("0.00")
        self.summary_cost.set("0.00")
        self.summary_profit.set("0.00")
        self.summary_margin.set("0%")
        self.set_feedback("已清空计算列表。", is_error=False)


    # ---------- Excel 记录 / 读取 / 统算 ----------

    def setup_excel_page(self):
        wrapper = ttk.Frame(self.excel_page, padding=12)
        wrapper.pack(fill=tk.BOTH, expand=True)

        self.excel_list_frame = ttk.Frame(wrapper)
        self.excel_detail_frame = ttk.Frame(wrapper)

        # 上一层：Excel 文件列表 + 全部统算
        list_top = ttk.Frame(self.excel_list_frame)
        list_top.pack(fill=tk.X, pady=(0, 8))
        ttk.Label(list_top, text="Excel 记录中心", style="Title.TLabel").pack(side=tk.LEFT)
        ttk.Button(list_top, text="刷新列表", command=self.refresh_excel_files).pack(side=tk.RIGHT, padx=(6, 0))
        ttk.Button(list_top, text="选择 Excel 查看", command=self.browse_excel_file).pack(side=tk.RIGHT, padx=(6, 0))
        ttk.Button(list_top, text="统算全部 Excel", command=self.aggregate_all_excels, style="Accent.TButton").pack(side=tk.RIGHT, padx=(6, 0))

        export_path_box = ttk.LabelFrame(self.excel_list_frame, text="Excel 导出位置（可填局域网共享路径）", padding=8, style="Section.TLabelframe")
        export_path_box.pack(fill=tk.X, pady=(0, 10))
        ttk.Label(export_path_box, text="导出目录：").pack(side=tk.LEFT)
        ttk.Entry(export_path_box, textvariable=self.export_dir_var).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(4, 8))
        ttk.Button(export_path_box, text="选择文件夹", command=self.choose_export_dir).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(export_path_box, text="应用导出位置", command=self.apply_export_dir, style="Accent.TButton").pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(export_path_box, text="恢复本地", command=self.reset_export_dir_to_local).pack(side=tk.LEFT)
        ttk.Label(
            export_path_box,
            text=r"示例：\\192.168.1.10\共享文件夹\利润表，或 Z:\利润表。电脑需提前有访问权限。",
            foreground="#7f8c8d",
            wraplength=520,
        ).pack(side=tk.LEFT, padx=(12, 0))

        stat_frame = ttk.LabelFrame(self.excel_list_frame, text="全部 Excel 利润统算", padding=10, style="Section.TLabelframe")
        stat_frame.pack(fill=tk.X, pady=(0, 10))
        stats = [
            ("文件数量", self.excel_total_count),
            ("总收入", self.excel_total_income),
            ("总成本", self.excel_total_cost),
            ("总利润", self.excel_total_profit),
            ("综合利润率", self.excel_total_margin),
        ]
        for idx, (label, var) in enumerate(stats):
            card = ttk.Frame(stat_frame)
            card.grid(row=0, column=idx, sticky="ew", padx=8)
            stat_frame.columnconfigure(idx, weight=1)
            ttk.Label(card, text=label).pack()
            ttk.Label(card, textvariable=var, font=("Arial", 13, "bold")).pack()

        list_body = ttk.Frame(self.excel_list_frame)
        list_body.pack(fill=tk.BOTH, expand=True)
        columns = ("file", "mtime", "income", "cost", "profit", "margin", "remark")
        self.excel_file_tree = ttk.Treeview(list_body, columns=columns, show="headings", height=16)
        headings = {
            "file": "文件名",
            "mtime": "修改时间",
            "income": "收入",
            "cost": "成本",
            "profit": "利润",
            "margin": "利润率",
            "remark": "备注",
        }
        widths = {"file": 320, "mtime": 150, "income": 110, "cost": 110, "profit": 110, "margin": 90, "remark": 360}
        for col in columns:
            self.excel_file_tree.heading(col, text=headings[col])
            self.excel_file_tree.column(col, width=widths[col], anchor=tk.W if col in ("file", "remark") else tk.CENTER)
        y_scroll = ttk.Scrollbar(list_body, orient="vertical", command=self.excel_file_tree.yview)
        x_scroll = ttk.Scrollbar(list_body, orient="horizontal", command=self.excel_file_tree.xview)
        self.excel_file_tree.configure(yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)
        self.excel_file_tree.grid(row=0, column=0, sticky="nsew")
        y_scroll.grid(row=0, column=1, sticky="ns")
        x_scroll.grid(row=1, column=0, sticky="ew")
        list_body.rowconfigure(0, weight=1)
        list_body.columnconfigure(0, weight=1)
        self.excel_file_tree.bind("<Double-1>", lambda _event: self.open_selected_excel_from_list())

        list_actions = ttk.Frame(self.excel_list_frame)
        list_actions.pack(fill=tk.X, pady=(8, 0))
        ttk.Button(list_actions, text="查看选中 Excel", command=self.open_selected_excel_from_list).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(list_actions, text="删除选中 Excel", command=self.delete_selected_excel).pack(side=tk.LEFT, padx=6)
        ttk.Label(list_actions, text="双击文件也可以查看。备注会写入选中行同一淘宝ID+记录时间分组的底部。", foreground="#7f8c8d").pack(side=tk.LEFT, padx=12)

        # 下一层：单个 Excel 查看 + 删除 + 备注
        detail_top = ttk.Frame(self.excel_detail_frame)
        detail_top.pack(fill=tk.X, pady=(0, 8))
        ttk.Label(detail_top, text="Excel 查看", style="Title.TLabel").pack(side=tk.LEFT)
        ttk.Button(detail_top, text="返回上一层", command=self.show_excel_list).pack(side=tk.RIGHT, padx=(6, 0))
        ttk.Button(detail_top, text="查看另一个 Excel", command=self.browse_excel_file).pack(side=tk.RIGHT, padx=(6, 0))

        selected_box = ttk.LabelFrame(self.excel_detail_frame, text="当前文件", padding=8, style="Section.TLabelframe")
        selected_box.pack(fill=tk.X, pady=(0, 8))
        ttk.Label(selected_box, textvariable=self.excel_selected_file, foreground="#2c3e50", wraplength=1080).pack(anchor=tk.W)

        sheet_bar = ttk.Frame(self.excel_detail_frame)
        sheet_bar.pack(fill=tk.X, pady=(0, 6))
        ttk.Label(sheet_bar, text="工作表：").pack(side=tk.LEFT)
        self.excel_sheet_combo = ttk.Combobox(sheet_bar, textvariable=self.excel_sheet_var, width=28, state="readonly")
        self.excel_sheet_combo.pack(side=tk.LEFT, padx=(4, 8))
        ttk.Button(sheet_bar, text="读取", command=self.load_current_excel_sheet).pack(side=tk.LEFT)
        ttk.Button(sheet_bar, text="删除同组行", command=self.delete_current_excel_row).pack(side=tk.LEFT, padx=(8, 0))
        ttk.Label(sheet_bar, text="提示：点选一行后，删除会同步删除同一淘宝ID+记录时间的所有行。", foreground="#7f8c8d").pack(side=tk.LEFT, padx=10)

        view_body = ttk.Frame(self.excel_detail_frame)
        view_body.pack(fill=tk.BOTH, expand=True)
        self.excel_view_tree = ttk.Treeview(view_body, show="headings", height=13, selectmode="browse")
        self.excel_view_tree.bind("<Delete>", lambda _event: self.delete_current_excel_row())
        view_y = ttk.Scrollbar(view_body, orient="vertical", command=self.excel_view_tree.yview)
        view_x = ttk.Scrollbar(view_body, orient="horizontal", command=self.excel_view_tree.xview)
        self.excel_view_tree.configure(yscrollcommand=view_y.set, xscrollcommand=view_x.set)
        self.excel_view_tree.grid(row=0, column=0, sticky="nsew")
        view_y.grid(row=0, column=1, sticky="ns")
        view_x.grid(row=1, column=0, sticky="ew")
        view_body.rowconfigure(0, weight=1)
        view_body.columnconfigure(0, weight=1)
        self.excel_sheet_combo.bind("<<ComboboxSelected>>", lambda _event: self.load_current_excel_sheet())

        remark_box = ttk.LabelFrame(self.excel_detail_frame, text="备注（先选择一行，再保存到同组底部）", padding=8, style="Section.TLabelframe")
        remark_box.pack(fill=tk.X, pady=(8, 0))
        self.excel_remark_text = tk.Text(remark_box, height=4, wrap="word")
        self.excel_remark_text.pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Button(remark_box, text="保存备注", command=self.save_current_excel_remark).pack(side=tk.RIGHT, padx=(8, 0))

        self.show_excel_list()
        self.root.after(300, self.refresh_excel_files)

    def choose_export_dir(self):
        initial_dir = self.normalize_path_text(self.export_dir_var.get() or self.report_dir)
        if not os.path.isdir(initial_dir):
            initial_dir = os.getcwd()
        path = filedialog.askdirectory(title="选择 Excel 导出目录", initialdir=initial_dir)
        if path:
            self.export_dir_var.set(path)
            self.apply_export_dir()

    def apply_export_dir(self):
        new_dir = self.normalize_path_text(self.export_dir_var.get())
        try:
            os.makedirs(new_dir, exist_ok=True)
            test_file = os.path.join(new_dir, ".write_test.tmp")
            with open(test_file, "w", encoding="utf-8") as f:
                f.write("ok")
            os.remove(test_file)
        except Exception as exc:
            messagebox.showerror(
                "导出位置不可用",
                f"无法使用这个导出目录：\n{new_dir}\n\n原因：{exc}\n\n请确认局域网共享已打开、账号有写入权限，或先在系统里映射为盘符。",
            )
            return False

        self.report_dir = new_dir
        self.export_dir_var.set(new_dir)
        self.save_settings()
        self.refresh_excel_files()
        self.aggregate_all_excels(silent=True)
        self.set_feedback(f"Excel 导出位置已切换为：{new_dir}", is_error=False)
        messagebox.showinfo("导出位置已应用", f"之后导出的 Excel 会保存到：\n{new_dir}")
        return True

    def reset_export_dir_to_local(self):
        self.export_dir_var.set(self.default_report_dir)
        self.apply_export_dir()

    def show_excel_list(self):
        self.excel_detail_frame.pack_forget()
        self.excel_list_frame.pack(fill=tk.BOTH, expand=True)
        self.refresh_excel_files()

    def show_excel_detail(self):
        self.excel_list_frame.pack_forget()
        self.excel_detail_frame.pack(fill=tk.BOTH, expand=True)

    def get_report_excel_paths(self):
        self.ensure_directories()
        paths = []
        for name in os.listdir(self.report_dir):
            if name.startswith("~$"):
                continue
            if name.lower().endswith(".xlsx"):
                paths.append(os.path.join(self.report_dir, name))
        paths.sort(key=lambda p: os.path.getmtime(p), reverse=True)
        return paths

    def _to_float(self, value, default=0.0):
        try:
            if value is None:
                return default
            if isinstance(value, str):
                text = value.strip().replace(",", "")
                if text.endswith("%"):
                    return float(text[:-1]) / 100.0
                return float(text)
            return float(value)
        except Exception:
            return default

    def parse_excel_summary(self, path):
        result = {"income": 0.0, "cost": 0.0, "profit": 0.0, "margin": 0.0}
        if load_workbook is None:
            return result
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
            if "利润明细" not in wb.sheetnames:
                wb.close()
                return result
            ws = wb["利润明细"]
            # 优先读取导出时的汇总行：A4:H4
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True):
                values = list(row)
                for idx, value in enumerate(values):
                    if value == "总收入" and idx + 1 < len(values):
                        result["income"] = self._to_float(values[idx + 1])
                    elif value == "总成本" and idx + 1 < len(values):
                        result["cost"] = self._to_float(values[idx + 1])
                    elif value == "总利润" and idx + 1 < len(values):
                        result["profit"] = self._to_float(values[idx + 1])
                    elif value == "综合利润率" and idx + 1 < len(values):
                        result["margin"] = self._to_float(values[idx + 1])
            # 如果旧文件没有摘要，则尝试找“本次总计”行。
            if result["income"] == 0 and result["cost"] == 0 and result["profit"] == 0:
                for row in ws.iter_rows(values_only=True):
                    values = list(row)
                    if "--- 本次总计 ---" in values:
                        result["income"] = self._to_float(values[14] if len(values) > 14 else 0)
                        result["cost"] = self._to_float(values[15] if len(values) > 15 else 0)
                        result["profit"] = self._to_float(values[16] if len(values) > 16 else 0)
                        break
            if result["income"]:
                result["margin"] = result["profit"] / result["income"]
            wb.close()
        except Exception:
            pass
        return result

    def read_excel_remark(self, path):
        if load_workbook is None:
            return ""
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
            if "备注" in wb.sheetnames:
                ws_note = wb["备注"]
                value = ws_note["B5"].value or ws_note["B4"].value
                wb.close()
                return str(value or "")
            if "利润明细" in wb.sheetnames:
                ws = wb["利润明细"]
                header_row, headers = self.find_excel_table_header(ws)
                remark_col = headers.get("备注") if headers else None
                if remark_col:
                    last_remark = ""
                    for r in range((ws.max_row or 1), header_row, -1):
                        value = ws.cell(row=r, column=remark_col).value
                        if value:
                            last_remark = str(value)
                            break
                    wb.close()
                    return last_remark
            wb.close()
        except Exception:
            pass
        return ""

    def refresh_excel_files(self):
        if not hasattr(self, "excel_file_tree"):
            return
        for item in self.excel_file_tree.get_children():
            self.excel_file_tree.delete(item)

        paths = self.get_report_excel_paths()
        for path in paths:
            summary = self.parse_excel_summary(path)
            remark = self.read_excel_remark(path)
            modified = time.strftime("%Y-%m-%d %H:%M", time.localtime(os.path.getmtime(path)))
            margin = f"{summary['margin'] * 100:.1f}%" if summary["income"] else "0%"
            self.excel_file_tree.insert(
                "",
                tk.END,
                iid=path,
                values=(
                    os.path.basename(path),
                    modified,
                    self.format_display(summary["income"]),
                    self.format_display(summary["cost"]),
                    self.format_display(summary["profit"]),
                    margin,
                    remark,
                ),
            )

    def aggregate_all_excels(self, silent=False):
        paths = self.get_report_excel_paths()
        total_income = 0.0
        total_cost = 0.0
        total_profit = 0.0
        valid_count = 0
        for path in paths:
            summary = self.parse_excel_summary(path)
            # 只统计能识别出利润明细的文件；纯备注或无关 Excel 会被跳过。
            if summary["income"] or summary["cost"] or summary["profit"]:
                valid_count += 1
                total_income += summary["income"]
                total_cost += summary["cost"]
                total_profit += summary["profit"]
        margin = (total_profit / total_income) if total_income else 0.0
        self.excel_total_count.set(f"{valid_count} 个文件")
        self.excel_total_income.set(self.format_display(total_income))
        self.excel_total_cost.set(self.format_display(total_cost))
        self.excel_total_profit.set(self.format_display(total_profit))
        self.excel_total_margin.set(f"{margin * 100:.1f}%")
        self.refresh_excel_files()
        if not silent:
            messagebox.showinfo("统算完成", f"已统计 {valid_count} 个可识别的导出 Excel。")

    def browse_excel_file(self):
        self.ensure_directories()
        path = filedialog.askopenfilename(
            title="选择要查看的 Excel",
            initialdir=os.path.abspath(self.report_dir),
            filetypes=[("Excel 文件", "*.xlsx"), ("所有文件", "*.*")],
        )
        if path:
            self.open_excel_path(path)

    def open_selected_excel_from_list(self):
        selected = self.excel_file_tree.selection()
        if not selected:
            messagebox.showwarning("未选择", "请先选择一个 Excel 文件。")
            return
        self.open_excel_path(selected[0])

    def open_excel_path(self, path):
        if load_workbook is None:
            messagebox.showerror("无法读取", "缺少 openpyxl，请先运行：pip install openpyxl")
            return
        if not os.path.exists(path):
            messagebox.showerror("文件不存在", f"找不到文件：{path}")
            return
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
            sheet_names = wb.sheetnames
            wb.close()
            if not sheet_names:
                messagebox.showwarning("无法查看", "这个 Excel 没有工作表。")
                return
            self.current_excel_path = path
            self.excel_selected_file.set(path)
            self.excel_sheet_combo.configure(values=sheet_names)
            preferred = "利润明细" if "利润明细" in sheet_names else sheet_names[0]
            self.excel_sheet_var.set(preferred)
            self.load_current_excel_sheet()
            self.load_current_excel_remark()
            self.show_excel_detail()
        except Exception as exc:
            messagebox.showerror("读取失败", f"无法打开 Excel：{exc}")

    def load_current_excel_sheet(self):
        path = self.current_excel_path
        if not path:
            return
        sheet_name = self.excel_sheet_var.get().strip()
        if not sheet_name:
            return
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
            ws = wb[sheet_name]
            max_cols = min(ws.max_column or 1, 36)
            max_rows = min(ws.max_row or 1, 300)
            data_columns = [f"c{i}" for i in range(1, max_cols + 1)]
            columns = ["row_no"] + data_columns
            self.excel_view_row_map = {}
            self.excel_view_tree.delete(*self.excel_view_tree.get_children())
            self.excel_view_tree.configure(columns=columns, show="headings")
            self.excel_view_tree.heading("row_no", text="行")
            self.excel_view_tree.column("row_no", width=55, anchor=tk.CENTER, stretch=False)
            for idx, col in enumerate(data_columns, start=1):
                letter = get_column_letter(idx)
                self.excel_view_tree.heading(col, text=letter)
                self.excel_view_tree.column(col, width=140, anchor=tk.W)
            for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols, values_only=True), start=1):
                values = [str(row_idx)]
                for value in row:
                    if value is None:
                        values.append("")
                    elif isinstance(value, float):
                        values.append(self.format_display(value))
                    else:
                        values.append(str(value))
                item_id = f"row_{row_idx}"
                self.excel_view_row_map[item_id] = row_idx
                self.excel_view_tree.insert("", tk.END, iid=item_id, values=values)
            wb.close()
        except Exception as exc:
            messagebox.showerror("读取失败", f"无法读取工作表：{exc}")


    def find_excel_table_header(self, ws):
        """返回利润明细表头行和列映射；兼容旧版“用户”和新版“淘宝ID”。"""
        max_scan = min(ws.max_row or 1, 40)
        for row_idx in range(1, max_scan + 1):
            headers = {}
            for cell in ws[row_idx]:
                if cell.value is None:
                    continue
                name = str(cell.value).strip()
                if name:
                    headers[name] = cell.column
            has_time = "记录时间" in headers
            has_id = "淘宝ID" in headers or "用户" in headers
            has_item = "项目名称" in headers
            if has_time and has_id and has_item:
                return row_idx, headers
        return None, {}

    def get_excel_id_column(self, headers):
        return headers.get("淘宝ID") or headers.get("用户")

    def get_selected_excel_row_index(self):
        selected = self.excel_view_tree.selection() if hasattr(self, "excel_view_tree") else []
        if not selected:
            return None
        return self.excel_view_row_map.get(selected[0])

    def get_excel_group_rows(self, ws, row_index, headers, header_row):
        """按淘宝ID/旧版用户 + 记录时间定位同一次导出的所有行。"""
        time_col = headers.get("记录时间")
        id_col = self.get_excel_id_column(headers)
        if not time_col or not id_col:
            return [], "", ""
        selected_time = str(ws.cell(row=row_index, column=time_col).value or "").strip()
        selected_id = str(ws.cell(row=row_index, column=id_col).value or "").strip()
        if not selected_time or not selected_id:
            return [], selected_id, selected_time
        matched = []
        for r in range(header_row + 1, (ws.max_row or 0) + 1):
            row_time = str(ws.cell(row=r, column=time_col).value or "").strip()
            row_id = str(ws.cell(row=r, column=id_col).value or "").strip()
            if row_id == selected_id and row_time == selected_time:
                matched.append(r)
        return matched, selected_id, selected_time

    def ensure_excel_remark_column(self, ws, headers, header_row):
        remark_col = headers.get("备注")
        if remark_col:
            return remark_col
        remark_col = (ws.max_column or 0) + 1
        ws.cell(row=header_row, column=remark_col, value="备注")
        headers["备注"] = remark_col
        try:
            ws.column_dimensions[get_column_letter(remark_col)].width = 45
            ws.cell(row=header_row, column=remark_col).font = Font(bold=True)
            ws.cell(row=header_row, column=remark_col).fill = PatternFill("solid", fgColor="D9EAF7")
        except Exception:
            pass
        return remark_col

    def recompute_excel_summary(self, ws, headers):
        """删除/备注后重算 A4:H4 摘要，统算全部Excel会以这里为准。"""
        income_col = headers.get("销售额")
        cost_col = headers.get("总成本")
        profit_col = headers.get("利润")
        item_col = headers.get("项目名称")
        header_row, _headers = self.find_excel_table_header(ws)
        if not header_row or not income_col or not cost_col or not profit_col:
            return
        total_income = total_cost = total_profit = 0.0
        for r in range(header_row + 1, (ws.max_row or 0) + 1):
            item_text = str(ws.cell(row=r, column=item_col).value or "").strip() if item_col else ""
            if not item_text:
                continue
            if item_text.startswith("---") or item_text in {"备注", "文件备注"}:
                continue
            total_income += self._to_float(ws.cell(row=r, column=income_col).value)
            total_cost += self._to_float(ws.cell(row=r, column=cost_col).value)
            total_profit += self._to_float(ws.cell(row=r, column=profit_col).value)
        margin = total_profit / total_income if total_income else 0.0
        ws.cell(row=4, column=1, value="总收入")
        ws.cell(row=4, column=2, value=total_income)
        ws.cell(row=4, column=3, value="总成本")
        ws.cell(row=4, column=4, value=total_cost)
        ws.cell(row=4, column=5, value="总利润")
        ws.cell(row=4, column=6, value=total_profit)
        ws.cell(row=4, column=7, value="综合利润率")
        ws.cell(row=4, column=8, value=f"{margin * 100:.1f}%")

    def delete_current_excel_row(self):
        path = self.current_excel_path
        if not path:
            messagebox.showwarning("未选择", "请先选择一个 Excel 文件。")
            return
        sheet_name = self.excel_sheet_var.get().strip()
        if not sheet_name:
            messagebox.showwarning("未选择", "请先选择一个工作表。")
            return
        row_index = self.get_selected_excel_row_index()
        if not row_index:
            messagebox.showwarning("未选择行", "请先在下方表格中点选要删除的行。")
            return
        try:
            wb = load_workbook(path)
            if sheet_name not in wb.sheetnames:
                wb.close()
                messagebox.showerror("删除失败", "当前工作表不存在，请重新打开 Excel。")
                return
            ws = wb[sheet_name]
            header_row, headers = self.find_excel_table_header(ws)
            if not header_row:
                wb.close()
                messagebox.showwarning("无法分组删除", "当前工作表没有识别到“记录时间 + 淘宝ID/用户 + 项目名称”表头。")
                return
            group_rows, taobao_id, record_time = self.get_excel_group_rows(ws, row_index, headers, header_row)
            if not group_rows:
                wb.close()
                messagebox.showwarning("无法分组删除", "当前选中行没有有效的记录时间或淘宝ID，无法判断同组数据。")
                return
            ok = messagebox.askyesno(
                "确认删除同组行",
                f"将删除当前工作表中同一组的 {len(group_rows)} 行：\n\n淘宝ID/用户：{taobao_id}\n记录时间：{record_time}\n文件：{os.path.basename(path)}\n工作表：{sheet_name}\n\n不会删除其他淘宝ID或其他记录时间的数据。删除后会直接保存到这个 Excel 文件。",
            )
            if not ok:
                wb.close()
                return
            for r in sorted(group_rows, reverse=True):
                ws.delete_rows(r, 1)
            self.recompute_excel_summary(ws, headers)
            wb.save(path)
            wb.close()
            self.load_current_excel_sheet()
            self.load_current_excel_remark()
            self.refresh_excel_files()
            self.aggregate_all_excels(silent=True)
            messagebox.showinfo("删除成功", f"已删除淘宝ID/用户“{taobao_id}”在 {record_time} 的 {len(group_rows)} 行。")
        except PermissionError:
            messagebox.showerror("删除失败", "无法保存 Excel。请先关闭正在打开的 Excel 文件后再试。")
        except Exception as exc:
            messagebox.showerror("删除失败", f"无法删除同组行：{exc}")

    def load_current_excel_remark(self):
        if not hasattr(self, "excel_remark_text"):
            return
        self.excel_remark_text.delete("1.0", tk.END)
        if self.current_excel_path:
            self.excel_remark_text.insert("1.0", self.read_excel_remark(self.current_excel_path))

    def save_current_excel_remark(self):
        path = self.current_excel_path
        if not path:
            messagebox.showwarning("未选择", "请先选择一个 Excel 文件。")
            return
        if load_workbook is None:
            messagebox.showerror("无法保存", "缺少 openpyxl，请先运行：pip install openpyxl")
            return
        sheet_name = self.excel_sheet_var.get().strip()
        if not sheet_name:
            messagebox.showwarning("未选择", "请先选择一个工作表。")
            return
        row_index = self.get_selected_excel_row_index()
        if not row_index:
            messagebox.showwarning("未选择行", "请先在下方表格中点选一行，再保存备注。")
            return
        remark = self.excel_remark_text.get("1.0", tk.END).strip()
        if not remark:
            messagebox.showwarning("备注为空", "请先输入备注内容。")
            return
        try:
            wb = load_workbook(path)
            if sheet_name not in wb.sheetnames:
                wb.close()
                messagebox.showerror("保存失败", "当前工作表不存在，请重新打开 Excel。")
                return
            ws = wb[sheet_name]
            header_row, headers = self.find_excel_table_header(ws)
            if not header_row:
                wb.close()
                messagebox.showwarning("无法保存备注", "当前工作表没有识别到“记录时间 + 淘宝ID/用户 + 项目名称”表头。")
                return
            group_rows, taobao_id, record_time = self.get_excel_group_rows(ws, row_index, headers, header_row)
            if not group_rows:
                wb.close()
                messagebox.showwarning("无法保存备注", "当前选中行没有有效的记录时间或淘宝ID，无法判断备注应该加入哪一组。")
                return

            remark_col = self.ensure_excel_remark_column(ws, headers, header_row)
            insert_at = max(group_rows) + 1
            ws.insert_rows(insert_at, 1)

            copy_cols = ["记录时间", "淘宝ID", "用户", "订单号", "微信号"]
            for name in copy_cols:
                col = headers.get(name)
                if col:
                    ws.cell(row=insert_at, column=col, value=ws.cell(row=row_index, column=col).value)
            item_col = headers.get("项目名称")
            if item_col:
                ws.cell(row=insert_at, column=item_col, value="备注")
            base_col = headers.get("基础类型")
            if base_col:
                ws.cell(row=insert_at, column=base_col, value="备注")
            ws.cell(row=insert_at, column=remark_col, value=remark)

            try:
                fill = PatternFill("solid", fgColor="FFF2CC")
                for col in range(1, ws.max_column + 1):
                    cell = ws.cell(row=insert_at, column=col)
                    cell.fill = fill
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            except Exception:
                pass

            if "备注" in wb.sheetnames:
                note_ws = wb["备注"]
            else:
                note_ws = wb.create_sheet("备注")
            note_ws["A1"] = "最近备注"
            note_ws["A2"] = "保存时间"
            note_ws["B2"] = time.strftime("%Y-%m-%d %H:%M:%S")
            note_ws["A3"] = "淘宝ID/用户"
            note_ws["B3"] = taobao_id
            note_ws["A4"] = "记录时间"
            note_ws["B4"] = record_time
            note_ws["A5"] = "备注"
            note_ws["B5"] = remark
            note_ws.column_dimensions["A"].width = 14
            note_ws.column_dimensions["B"].width = 80
            note_ws["B5"].alignment = Alignment(wrap_text=True, vertical="top")

            wb.save(path)
            wb.close()
            self.refresh_excel_files()
            self.open_excel_path(path)
            self.excel_sheet_var.set(sheet_name)
            self.load_current_excel_sheet()
            messagebox.showinfo("保存成功", f"备注已加入淘宝ID/用户“{taobao_id}”在 {record_time} 的分组底部。")
        except PermissionError:
            messagebox.showerror("保存失败", "无法保存 Excel。请先关闭正在打开的 Excel 文件后再试。")
        except Exception as exc:
            messagebox.showerror("保存失败", f"无法保存备注：{exc}")

    def delete_selected_excel(self):
        selected = self.excel_file_tree.selection()
        if not selected:
            messagebox.showwarning("未选择", "请先选择要删除的 Excel 文件。")
            return
        path = selected[0]
        self.delete_excel_path(path)

    def delete_current_excel(self):
        if not self.current_excel_path:
            messagebox.showwarning("未选择", "当前没有打开的 Excel 文件。")
            return
        self.delete_excel_path(self.current_excel_path)

    def delete_excel_path(self, path):
        if not os.path.exists(path):
            messagebox.showerror("文件不存在", f"找不到文件：{path}")
            return
        ok = messagebox.askyesno("确认删除", f"确定删除这个 Excel 文件吗？\n\n{path}\n\n删除后不可从程序内恢复。")
        if not ok:
            return
        try:
            os.remove(path)
            if self.current_excel_path == path:
                self.current_excel_path = None
                self.excel_selected_file.set("未选择 Excel 文件")
                self.show_excel_list()
            self.refresh_excel_files()
            self.aggregate_all_excels(silent=True)
        except Exception as exc:
            messagebox.showerror("删除失败", f"无法删除文件：{exc}")

    # ---------- Excel 导出 ----------

    def style_excel_sheet(self, ws, header_row=None, freeze=None, widths=None):
        title_fill = PatternFill("solid", fgColor="1F4E78")
        header_fill = PatternFill("solid", fgColor="D9EAF7")
        white_font = Font(color="FFFFFF", bold=True, size=14)
        header_font = Font(bold=True)
        thin = Side(style="thin", color="D9D9D9")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        if ws.max_row >= 1:
            ws["A1"].font = white_font
            ws["A1"].fill = title_fill
            ws["A1"].alignment = Alignment(horizontal="left")

        if header_row:
            for cell in ws[header_row]:
                cell.font = header_font
                cell.fill = header_fill
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)

        if freeze:
            ws.freeze_panes = freeze

        if header_row and ws.max_column:
            last_col = get_column_letter(ws.max_column)
            ws.auto_filter.ref = f"A{header_row}:{last_col}{ws.max_row}"

        if widths:
            for col, width in widths.items():
                ws.column_dimensions[col].width = width
        else:
            for col_idx in range(1, ws.max_column + 1):
                col_letter = get_column_letter(col_idx)
                max_len = 10
                for cell in ws[col_letter]:
                    value = "" if cell.value is None else str(cell.value)
                    max_len = max(max_len, min(len(value) + 2, 40))
                ws.column_dimensions[col_letter].width = max_len


    def get_excel_headers(self):
        """母版和子版共用同一套 Excel 表头，确保同一天追加到同一个 Excel。"""
        return [
            "记录时间",
            "淘宝ID",
            "订单号",
            "微信号",
            "项目名称",
            "基础类型",
            "是否单件",
            "高度(m)",
            "宽度(m)",
            "数量",
            "销售单价",
            "成本单价",
            "K系数",
            "折扣",
            "折扣系数",
            "折前销售额",
            "销售额",
            "总成本",
            "利润",
            "利润率",
            "高度原文",
            "宽度原文",
            "数量原文",
            "销售单价原文",
            "成本单价原文",
            "K系数原文",
        ]

    def get_daily_excel_path(self):
        """每天只创建一个 Excel：利润明细_YYYYMMDD.xlsx。"""
        date_text = time.strftime("%Y%m%d")
        return os.path.join(self.report_dir, f"{DAILY_EXCEL_PREFIX}_{date_text}.xlsx")

    def create_profit_sheet_template(self, ws, timestamp):
        """初始化利润明细工作表模板。"""
        if ws.max_row:
            ws.delete_rows(1, ws.max_row)
        ws.title = "利润明细"
        ws.append(["六六派对利润明细"])
        ws.append(["最后导出时间", timestamp])
        ws.append(["最近导出淘宝ID", self.user_name.get().strip(), "订单号", self.order_no.get().strip(), "微信号", self.wechat_no.get().strip()])
        ws.append(["总收入", 0.0, "总成本", 0.0, "总利润", 0.0, "综合利润率", "0.0%"])
        ws.append([])
        ws.append(self.get_excel_headers())
        return ws

    def open_daily_workbook_for_append(self, timestamp):
        """打开或创建当天 Excel。母版和子版都只写这个文件。"""
        xlsx_path = self.get_daily_excel_path()
        if os.path.exists(xlsx_path):
            wb = load_workbook(xlsx_path)
            if "利润明细" in wb.sheetnames:
                ws = wb["利润明细"]
            else:
                ws = wb.create_sheet("利润明细", 0)
            header_row, _headers = self.find_excel_table_header(ws)
            if not header_row:
                self.create_profit_sheet_template(ws, timestamp)
        else:
            wb = Workbook()
            ws = wb.active
            self.create_profit_sheet_template(ws, timestamp)
        return xlsx_path, wb, ws





    def style_daily_profit_sheet(self, ws):
        """设置当天利润明细表样式。"""
        header_row, _headers = self.find_excel_table_header(ws)
        if not header_row:
            header_row = 6
        for row in range(header_row + 1, ws.max_row + 1):
            for col in [8, 9, 10, 11, 12, 13, 16, 17, 18, 19]:
                ws.cell(row=row, column=col).number_format = "#,##0.00"
            ws.cell(row=row, column=15).number_format = "0.0%"
            ws.cell(row=row, column=20).number_format = "0.0%"
        for col in [2, 4, 6]:
            ws.cell(row=4, column=col).number_format = "#,##0.00"
        widths = {
            "A": 20, "B": 16, "C": 18, "D": 16, "E": 28, "F": 16,
            "G": 10, "H": 10, "I": 10, "J": 10, "K": 12, "L": 12,
            "M": 10, "N": 12, "O": 12, "P": 14, "Q": 14, "R": 14,
            "S": 14, "T": 10, "U": 14, "V": 14, "W": 14, "X": 16,
            "Y": 16, "Z": 14,
        }
        self.style_excel_sheet(ws, header_row=header_row, freeze=f"A{header_row + 1}", widths=widths)
        try:
            if "A1:Z1" not in [str(rng) for rng in ws.merged_cells.ranges]:
                ws.merge_cells("A1:Z1")
        except Exception:
            pass
        ws.row_dimensions[1].height = 24


    # ---------- 截图 ----------


if __name__ == "__main__":
    root = tk.Tk()
    app = PriceCalculator(root)
    root.mainloop()
