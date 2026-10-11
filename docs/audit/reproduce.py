"""Read-only baseline audit: run from repo root with Python + openpyxl.

Loads definitions without starting Tk. Writes synthetic workbooks only inside
TemporaryDirectory. Stubbed UI tests are not Windows/Tk/SMB acceptance tests.
"""
import ast
import importlib.util
import json
from pathlib import Path
import runpy
import tempfile
import types

import openpyxl
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

ROOT = Path(__file__).resolve().parents[2]
# Both diagnostic entry points use the same manifest and byte-level policy.
spec = importlib.util.spec_from_file_location("baseline_check", ROOT / "tools/baseline_check.py")
baseline = importlib.util.module_from_spec(spec)
spec.loader.exec_module(baseline)


class Var:
    def __init__(self, value=""):
        self.value = value

    def get(self):
        return self.value

    def set(self, value):
        self.value = value


class Tree:
    def get_children(self):
        return []

    def delete(self, *args):
        pass

    def insert(self, *args, **kwargs):
        pass


def instance(ns, folder):
    obj = ns["PriceCalculator"].__new__(ns["PriceCalculator"])
    obj.report_dir = str(folder)
    obj.image_dir = str(folder)
    for key, value in {"user_name": "客户", "order_no": "订单一", "wechat_no": "", "summary_income": "10.00", "summary_cost": "4.00", "summary_profit": "6.00", "summary_margin": "60%", "total_val": ""}.items():
        setattr(obj, key, Var(value))
    obj.validation_issues = []
    obj.last_config_snapshot = {}
    obj.set_feedback = lambda *a, **kw: None
    obj.render_issues = lambda *a: None
    obj.notify_validation_errors = lambda *a: None
    obj.refresh_item_selector = lambda: None
    obj.update_calc_row_states = lambda *a: None
    obj.reload_price_settings_from_json = lambda **kw: True
    obj.profit_tree = Tree()
    return obj


def calculate_rows(obj, quantities, price="10", cost="4"):
    obj.get_settings_snapshot = lambda: ({"商品": {"sale": price, "cost": cost, "k": "1", "single": True}}, [])
    obj.calc_rows = [{"display_name": "商品", "base_type": "商品", "row_kind": "normal", **{k: Var(v) for k, v in {"q": q, "h": "", "w": "", "discount": "", "res": "", "status": ""}.items()}} for q in quantities]
    obj.calculate(notify=False)


def main():
    result = {"openpyxl": openpyxl.__version__, "evidence": {}}
    manifest = json.loads((ROOT / "docs/baseline/manifest.json").read_text(encoding="utf-8"))
    for entry in manifest["files"]:
        label = {"manager": "A", "sales": "B"}[entry["role"]]
        filename = entry["path"]
        source = baseline.inspect_source(ROOT, entry)
        if source["status"] != "ok":
            raise SystemExit(f"Baseline changed: {filename}; re-audit before running these observations")
        data = baseline.source_path(ROOT, entry).read_bytes()
        ast.parse(data.decode("utf-8"))
        ns = runpy.run_path(str(baseline.source_path(ROOT, entry)), run_name="audit_only")
        # Functions keep runpy's globals; override message boxes there, never show UI.
        ns["PriceCalculator"].export_to_excel.__globals__["messagebox"] = types.SimpleNamespace(showinfo=lambda *a: None, showerror=lambda *a: None, showwarning=lambda *a: None)
        with tempfile.TemporaryDirectory(prefix="777-audit-") as temp:
            folder = Path(temp)
            obj = instance(ns, folder)
            out = {
                "lines": len(data.splitlines()),
                "blob": source["normalized_git_blob_sha1"],
                "worktree_blob": source["git_blob_sha1"],
                "baseline_match_kind": source["baseline_match_kind"],
            }
            calculate_rows(obj, ["1", "bad"])
            out["partial_calculation"] = {"valid_records": len(obj.last_export_records), "errors": sum(x["级别"] == "错误" for x in obj.validation_issues)}
            records = list(obj.last_export_records)
            # Export real method, retaining the already calculated invalid state.
            obj.calculate = lambda: None
            if label == "B":
                obj.get_default_report_dir = lambda: temp
                obj.export_dir_var = Var(temp)
            obj.export_to_excel()
            path = Path(obj.get_daily_excel_path())
            out["partial_export_file_created"] = path.exists()
            obj.last_export_records = records
            wb = openpyxl.Workbook()
            ws = wb.active
            obj.create_profit_sheet_template(ws, "2026-10-10 12:00:00")
            obj.append_profit_records_to_sheet(ws, "2026-10-10 12:00:00")
            obj.order_no.set("订单二")
            obj.append_profit_records_to_sheet(ws, "2026-10-10 12:00:00")
            hr, headers = obj.find_excel_table_header(ws)
            group = obj.get_excel_group_rows(ws, 7, headers, hr)[0]
            out["group_collision"] = {"rows": group, "orders": sorted({ws.cell(r, 3).value for r in group})}
            # Spreadsheet text assigned by actual writer is interpreted as formula.
            obj.user_name.set("=1+1")
            obj.append_profit_records_to_sheet(ws, "2026-10-10 12:00:01")
            out["customer_cell_type"] = ws.cell(ws.max_row - 1, 2).data_type
            wb.close()
            # Unknown schema is erased by the real append opener.
            wb = openpyxl.Workbook()
            wb.active.title = "利润明细"
            wb.active.append(["需要保留的旧数据"])
            wb.save(path)
            wb.close()
            _, wb, ws = obj.open_daily_workbook_for_append("2026-10-10 12:00:02")
            out["unknown_schema_old_cell_after_open"] = ws["A1"].value
            wb.close()
            # Two readers / two saves model the exact lost-update interleaving.
            first = openpyxl.load_workbook(path)
            second = openpyxl.load_workbook(path)
            first.active.append(["甲的新增订单"])
            first.save(path)
            second.active.append(["乙的新增订单"])
            second.save(path)
            first.close()
            second.close()
            final = openpyxl.load_workbook(path)
            out["lost_update_values"] = [r[0] for r in final.active.iter_rows(values_only=True)]
            final.close()
            out["bool_expression"] = ns["SafeNumberParser"].parse("True")
            obj = instance(ns, folder)
            calculate_rows(obj, ["1", "1"], price="0.005", cost="0")
            out["rounding"] = {"displayed_lines": [r["res"].get() for r in obj.calc_rows], "displayed_total": obj.summary_income.get()}
            calculate_rows(obj, ["1e308"], price="1e308", cost="0")
            out["overflow_display"] = obj.summary_income.get()
            if label == "A":
                bad = folder / "bad.xlsx"
                bad.write_text("not a zip")
                out["corrupt_summary"] = obj.parse_excel_summary(str(bad))
                # Current 26-column totals with no summary exercise fixed-index fallback.
                wb = openpyxl.Workbook()
                ws = wb.active
                obj.create_profit_sheet_template(ws, "t")
                obj.last_export_records = records
                obj.summary_income.set("10.00")
                obj.summary_cost.set("4.00")
                obj.summary_profit.set("6.00")
                obj.append_profit_records_to_sheet(ws, "t")
                p = folder / "fallback.xlsx"
                wb.save(p)
                wb.close()
                out["fallback_summary"] = obj.parse_excel_summary(str(p))
            else:
                out["string_false_single"] = obj.convert_price_items_from_json({"items": [{"name": "商品", "single": "false"}]})[0]["single"]
                # Instrument the installed library, not production code.
                counts = {}
                for n in [20, 40]:
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    obj.create_profit_sheet_template(ws, "t")
                    for i in range(n):
                        ws.append(["t", "客户", str(i), "", "商品"])
                    p = folder / f"rows{n}.xlsx"
                    wb.save(p)
                    wb.close()
                    calls = [0]
                    original = ReadOnlyWorksheet._get_source
                    def counted(self):
                        calls[0] += 1
                        return original(self)
                    ReadOnlyWorksheet._get_source = counted
                    try:
                        count = obj.count_visible_excel_rows(str(p))
                    finally:
                        ReadOnlyWorksheet._get_source = original
                    counts[str(n)] = {"visible": count, "xml_source_opens": calls[0]}
                out["read_only_cell_rescans"] = counts
            result["evidence"][label] = out
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
