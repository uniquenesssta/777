"""P03-03 user-visible storage, price protection and append-only snapshots."""
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from test_application_integration import application
from quote_files.images import save_png
from quote_files.settings import decode_settings, encode_settings, probe_directory
from quote_files.transactions import FileFailure, file_lock
from quote_app.connections import absolute_root
import p02_gui_check as gui


class StorageConfigurationTests(unittest.TestCase):
    def test_restart_restores_explicit_root_prices_and_directories_before_read(self):
        with application("manager") as (app, folder, *_):
            app.export_dir_var.set(str(folder / "中文 空格订单"))
            app.screenshot_dir_var.set(str(folder / "中文 图片"))
            self.assertTrue(app.apply_export_dir())
            app.backup_dir_var.set(str(folder / "备份登记"))
            self.assertTrue(app.apply_connection())
            app.initialize_quote_state()
            self.assertEqual(app.data_root, str(folder))
            self.assertEqual(app.report_dir, str(folder / "中文 空格订单"))
            self.assertEqual(app.image_dir, str(folder / "中文 图片"))
            self.assertEqual(app.backup_dir_var.get(), str(folder / "备份登记"))
            self.assertFalse((folder / "备份登记").exists())
            self.assertTrue(app.load_settings_to_ui())

    def test_unconfigured_start_never_opens_legacy_paths_or_creates_settings(self):
        for role in ("manager", "sales"):
            with application(role) as (app, folder, *_):
                Path(app.local_settings_file).unlink()
                before = gui.files(folder)
                app.initialize_quote_state()
                self.assertEqual(app.data_root, "")
                self.assertEqual(app._price_path(), "")
                self.assertFalse(app.load_settings_to_ui())
                self.assertFalse(app.export_to_excel())
                self.assertEqual(gui.files(folder), before)

                app.select_legacy_connection()
                self.assertEqual(app.data_source_var.get(), app.legacy_data_root)
                self.assertEqual(app.data_root, "")
                self.assertEqual(gui.files(folder), before)

    def test_gui_sandbox_rejects_outside_connection_before_reading_any_source(self):
        with application("manager") as (app, folder, ns, _):
            app.__class__ = gui.live_sandbox_type(ns, folder)
            app.data_source_var.set(str(folder.parent / "production-sentinel"))
            with patch("quote_app.connections.read_bytes", side_effect=AssertionError("outside source read")):
                with self.assertRaisesRegex(ValueError, "临时数据源"):
                    app.apply_connection()

    def test_bad_local_settings_are_not_overwritten_on_start_save_or_close(self):
        with application("manager") as (app, folder, *_):
            settings = Path(app.local_settings_file)
            settings.write_bytes(b'{"schema":')
            before = gui.files(folder)
            app.initialize_quote_state()
            app.data_source_var.set(str(folder))
            self.assertFalse(app.apply_connection())
            with patch.object(app.root, "destroy"):
                app.on_close()
            self.assertEqual(gui.files(folder), before)

    def test_connection_failure_retains_saved_root_and_live_paths(self):
        with application("manager") as (app, folder, *_):
            before = Path(app.local_settings_file).read_bytes()
            live = app.data_root, app.config_file, app.report_dir, app.image_dir
            app.data_source_var.set(str(folder / "不存在"))
            self.assertFalse(app.apply_connection())
            self.assertEqual((app.data_root, app.config_file, app.report_dir, app.image_dir), live)
            self.assertEqual(Path(app.local_settings_file).read_bytes(), before)

    def test_switch_source_changes_connection_without_moving_or_merging_old_data(self):
        with application("manager") as (app, folder, *_), tempfile.TemporaryDirectory() as target:
            target = Path(target)
            (folder / "reports" / "original.txt").write_text("original")
            gui.load_namespace(gui.ROOT, "manager", target)
            app.data_source_var.set(str(target))
            self.assertTrue(app.apply_connection())
            self.assertEqual(app.data_root, str(target))
            self.assertFalse((target / "reports" / "original.txt").exists())
            self.assertEqual((folder / "reports" / "original.txt").read_text(), "original")
            self.assertEqual((folder / "config.json").read_bytes(), json.dumps(gui.FIXTURE, ensure_ascii=False).encode())

    def test_prices_and_directory_changes_do_not_rewrite_each_other(self):
        with application("manager") as (app, folder, *_):
            original = (folder / "config.json").read_bytes()
            app.export_dir_var.set(str(folder / "独立订单"))
            self.assertTrue(app.apply_export_dir())
            self.assertEqual((folder / "config.json").read_bytes(), original)
            shared = (folder / "storage.json").read_bytes()
            app.setting_rows[0]["sale"].set("15")
            app.export_dir_var.set(str(folder / "未保存"))
            self.assertTrue(app.save_settings())
            self.assertEqual((folder / "storage.json").read_bytes(), shared)
            self.assertEqual(app.report_dir, str(folder / "独立订单"))
            self.assertEqual(set(json.loads((folder / "config.json").read_bytes())), {"items"})

    def test_new_empty_source_can_publish_explicit_prices_and_directories(self):
        with application("manager") as (app, folder, *_), tempfile.TemporaryDirectory() as new:
            app.data_source_var.set(new)
            self.assertTrue(app.apply_connection())
            self.assertIsNone(app.active_price_config)
            self.assertFalse(app._storage_valid)
            self.assertFalse(app.export_to_excel())
            app.setting_rows.clear()
            app.add_setting_row(name="首次明确价格", sale="10", cost="4", k="1", single=True)
            self.assertTrue(app.save_settings())
            self.assertTrue(app.apply_export_dir())
            self.assertEqual(app.config_file, str(Path(new) / "config.json"))
            app.initialize_quote_state()
            self.assertTrue(app.load_settings_to_ui())
            self.assertEqual(app.report_dir, str(Path(new) / "利润表"))

    def test_refresh_failure_after_settings_publish_reports_saved(self):
        with application("manager") as (app, folder, *_):
            app.export_dir_var.set(str(folder / "已保存但刷新失败"))
            with patch.object(app, "refresh_excel_files", side_effect=OSError("view failed")):
                self.assertTrue(app.apply_export_dir())
            self.assertIn("已保存", app.feedback)
            app.initialize_quote_state()
            self.assertEqual(app.report_dir, str(folder / "已保存但刷新失败"))

    def test_directory_change_alone_does_not_make_confirmed_prices_stale(self):
        with application("manager") as (app, folder, *_):
            app.add_fixture_row()
            app.calculate(notify=False)
            prices, result = app.quote_session.state.prices, app.quote_session.state.result
            app.export_dir_var.set(str(folder / "目录变化"))
            self.assertTrue(app.apply_export_dir())
            self.assertIs(app.quote_session.state.prices, prices)
            self.assertIs(app.quote_session.state.result, result)
            self.assertTrue(app.export_to_excel())

    def test_corrupt_prices_cannot_be_replaced_but_valid_storage_remains_readable(self):
        with application("manager") as (app, folder, *_):
            app.add_fixture_row()
            app.calculate(notify=False)
            path = app.export_to_excel()
            (folder / "config.json").write_bytes(b'{"items":')
            self.assertFalse(app.reload_price_settings_from_json())
            before = (folder / "config.json").read_bytes()
            app.setting_rows[0]["sale"].set("42")
            self.assertFalse(app.save_settings())
            self.assertEqual((folder / "config.json").read_bytes(), before)
            self.assertIn(path, app.get_report_excel_paths())
            self.assertFalse(app.export_to_excel())
            app.export_dir_var.set(str(folder / "价格损坏仍可设目录"))
            self.assertTrue(app.apply_export_dir())
            self.assertEqual((folder / "config.json").read_bytes(), before)

    def test_corrupt_shared_settings_block_output_and_cannot_be_overwritten(self):
        with application("manager") as (app, folder, *_):
            prior = app.report_dir, app.image_dir
            (folder / "storage.json").write_bytes(b'{"schema":')
            self.assertFalse(app.observe_storage())
            before = gui.files(folder)
            self.assertFalse(app.apply_export_dir())
            self.assertEqual(app.get_report_excel_paths(), [])
            self.assertFalse(app.export_to_excel())
            self.assertEqual((app.report_dir, app.image_dir), prior)
            self.assertEqual(gui.files(folder), before)

    def test_directory_probe_failure_keeps_saved_shared_settings(self):
        with application("manager") as (app, folder, *_):
            prior, live = (folder / "storage.json").read_bytes(), app.report_dir
            app.export_dir_var.set(str(folder / "无权限"))
            with patch("quote_app.connections.probe_directory", side_effect=PermissionError("denied")):
                self.assertFalse(app.apply_export_dir())
            self.assertEqual((folder / "storage.json").read_bytes(), prior)
            self.assertEqual(app.report_dir, live)

    def test_two_managers_cannot_replace_each_others_unseen_directory_change(self):
        with application("manager") as (first, folder, *_), application("manager") as (second, _, *_):
            second.data_source_var.set(str(folder))
            self.assertTrue(second.apply_connection())
            first.export_dir_var.set(str(folder / "第一台"))
            self.assertTrue(first.apply_export_dir())
            saved = (folder / "storage.json").read_bytes()
            second.export_dir_var.set(str(folder / "旧页面"))
            self.assertFalse(second.apply_export_dir())
            self.assertEqual((folder / "storage.json").read_bytes(), saved)
            self.assertTrue(second.observe_storage())
            self.assertEqual(second.report_dir, first.report_dir)

    def test_locked_settings_fail_without_permanently_blocking_a_safe_retry(self):
        with application("manager") as (app, folder, *_):
            app.export_dir_var.set(str(folder / "锁后重试"))
            with file_lock(folder / "storage.json"):
                self.assertFalse(app.apply_export_dir())
            self.assertIsNone(app._blocking_file_request)
            self.assertTrue(app.apply_export_dir())

    def test_unique_probes_clean_only_their_own_files(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            other = folder / "write_test.txt"
            other.write_bytes(b"someone else's file")
            probe_directory(folder)
            probe_directory(folder)
            self.assertEqual(list(folder.iterdir()), [other])
            self.assertEqual(other.read_bytes(), b"someone else's file")

    def test_unknown_settings_result_blocks_writes_then_queries_and_retries_same_request(self):
        import os
        with application("manager") as (app, folder, *_):
            original_replace = os.replace
            def fail_publish(source, target):
                if str(target) == str(folder / "storage.json"):
                    raise OSError("publish result unknown")
                return original_replace(source, target)
            app.add_fixture_row()
            app.calculate(notify=False)
            app.export_dir_var.set(str(folder / "结果核查"))
            with patch("quote_files.transactions.os.replace", side_effect=fail_publish):
                self.assertFalse(app.apply_export_dir())
            rid = app._blocking_file_request["request_id"]
            with patch("quote_app.integration.workbook_transaction", side_effect=AssertionError("write while unknown")):
                self.assertFalse(app.export_to_excel())
            self.assertEqual(app.check_last_write()["status"], "not_committed")
            self.assertTrue(app.apply_export_dir())
            self.assertEqual(app.last_file_transaction.request_id, rid)

    def test_unknown_png_is_queryable_and_prevents_another_capture_until_resolved(self):
        import os
        from PIL import Image
        with application("manager") as (app, folder, *_):
            original_replace = os.replace
            def fail_publish(source, target):
                if str(target).endswith(".png"):
                    raise OSError("publish result unknown")
                return original_replace(source, target)
            app.add_fixture_row()
            app.calculate(notify=False)
            with patch.object(app.root, "winfo_rootx", return_value=0), patch.object(app.root, "winfo_rooty", return_value=0), patch.object(app.root, "winfo_width", return_value=1), patch.object(app.root, "winfo_height", return_value=1), patch("PIL.ImageGrab.grab", return_value=Image.new("RGB", (1, 1))) as grab, patch("quote_files.transactions.os.replace", side_effect=fail_publish):
                self.assertFalse(app._capture_current_quote())
                self.assertIsNotNone(app._blocking_file_request)
                self.assertFalse(app._capture_current_quote())
                self.assertEqual(grab.call_count, 1)
            self.assertEqual(app.check_last_write()["status"], "not_committed")
            self.assertIsNone(app._blocking_file_request)
            self.assertFalse(list((folder / "screenshots").glob("*.png")))

    def test_duplicate_unknown_or_future_settings_are_rejected(self):
        valid = {"schema": "777-shared-v1", "price_file": "config.json", "export_dir": "订单", "screenshot_dir": "图片"}
        for raw in (b'{"schema":"777-shared-v1","schema":"777-shared-v1"}',
                    encode_settings(valid, "shared").replace(b'777-shared-v1', b'777-shared-v2'),
                    json.dumps({**valid, "unreviewed": True}).encode()):
            with self.assertRaises(FileFailure):
                decode_settings(raw, "shared")
        for root in ("relative", "E:relative", "\\rootrelative"):
            with self.assertRaises(FileFailure):
                absolute_root(root)
        self.assertEqual(absolute_root(r"D:\中文 空格"), r"D:\中文 空格")
        self.assertEqual(absolute_root(r"\\host\share\中文"), r"\\host\share\中文")

    def test_legacy_import_is_explicit_and_does_not_rewrite_price_source(self):
        with application("manager") as (app, folder, *_):
            before = gui.files(folder)
            self.assertTrue(app.import_legacy_directories())
            self.assertEqual(app.export_dir_var.get(), str(folder / "reports"))
            self.assertEqual(gui.files(folder), before)

    def test_each_order_keeps_its_original_complete_price_snapshot(self):
        from openpyxl import load_workbook
        with application("manager") as (app, folder, *_):
            app.add_fixture_row(product="P02面积")
            app.calculate(notify=False)
            path = app.export_to_excel()
            book = load_workbook(path)
            original = list(book["订单价格快照"].values)
            legacy = list(book["价格设定"].values)
            book.close()
            app.setting_rows[1]["sale"].set("20")
            self.assertTrue(app.save_settings())
            app.calculate(notify=False)
            app.confirm_current_prices()
            self.assertEqual(app.export_to_excel(), path)
            self.assertEqual(app.export_to_excel(), path)  # Replay adds no snapshot rows.
            book = load_workbook(path)
            try:
                history = list(book["订单价格快照"].values)
                self.assertEqual(history[:len(original)], original)
                self.assertEqual(len(history), 5)
                self.assertEqual(list(book["价格设定"].values), legacy)
                self.assertNotEqual(history[1][0], history[3][0])
                self.assertNotEqual(history[1][3], history[3][3])
                self.assertEqual([history[2][7], history[4][7]], ["10", "20"])
            finally:
                book.close()

    def test_same_timestamp_screenshots_have_unique_verified_pngs(self):
        from PIL import Image
        for role in ("manager", "sales"):
            with application(role) as (app, folder, *_):
                app.add_fixture_row()
                app.calculate(notify=False)
                with patch.object(app.root, "winfo_rootx", return_value=0), patch.object(app.root, "winfo_rooty", return_value=0), patch.object(app.root, "winfo_width", return_value=1), patch.object(app.root, "winfo_height", return_value=1), patch("PIL.ImageGrab.grab", return_value=Image.new("RGB", (1, 1))), patch("quote_app.integration.time.time_ns", return_value=1):
                    first = app._capture_current_quote()
                    original = Path(first).read_bytes()
                    second = app._capture_current_quote()
                    self.assertNotEqual(first, second)
                    self.assertEqual(Path(first).read_bytes(), original)
                    with Image.open(second) as check:
                        check.verify()

    def test_existing_screenshot_is_never_overwritten_even_with_same_image(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "original.png"
            image = Image.new("RGB", (1, 1))
            save_png(path, image)
            original = path.read_bytes()
            with self.assertRaises(FileFailure):
                save_png(path, image)
            self.assertEqual(path.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
