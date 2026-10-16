"""Contract tests against a real native Rust HTTPS process and old clients."""
import base64,hashlib,io,json,os,socket,ssl,subprocess,sys,threading,time,unittest,uuid
from pathlib import Path
from unittest.mock import patch
from native_host_fixture import Hosted,command,BINARY,ROOT,AUXILIARY
from quote_core import QuoteSnapshot,NormalLine,DraftDetails,PriceSnapshot,PriceItem,calculate_quote
from quote_service.protocol import APIError,calculated
from quote_service.client import CommitUnknown
from quote_service.journal import Journal,MemoryJournal
from quote_service.recovery import manifest,restore
from quote_files.transactions import file_lock,FileBusy
from quote_ledger.backup import hash_file
class NativeHostTests(unittest.TestCase):
    def setUp(self):self.host=Hosted();self.sales=self.host.client();self.manager=self.host.client('manager')
    def tearDown(self):self.host.close()
    def submit(self,client=None,quantity='4',journal=None):
        return (client or self.sales).submit(QuoteSnapshot('fixture-v1',(NormalLine('单件',quantity),)),DraftDetails('=中文','00003','微信'),journal or MemoryJournal())
    def wait(self,key,client=None):
        deadline=time.monotonic()+10
        while True:
            state=(client or self.sales).request('GET','/v1/outputs/'+key)
            if state['state'] in ('done','failed'):return state
            if time.monotonic()>deadline:raise RuntimeError('output timed out')
            time.sleep(.02)
    def directories(self):
        p=self.host.root/'public';p.mkdir(exist_ok=True)
        # Windows isolated test never pretends a local drive is an authorized
        # product UNC setting; both runtime adapters read an owned test record.
        with self.host.service.auth.connection() as c:c.execute('UPDATE preferences SET export_dir=?,image_dir=?,revision=revision+1 WHERE id=1',(str(p),str(p)))
        return p
    def output(self,order,kind='excel_quote',client=None):
        c=client or self.sales
        return c.mutate('/v1/outputs',{'order_id':order['order_id'],'expected_revision':order['revision'],'kind':kind},MemoryJournal())
    def test_standalone_initialization_is_independently_readable_and_native_selected(self):
        self.assertEqual(self.host.ready['host_backend'],'rust');self.assertEqual(command('inspect',root=self.host.host_root)['inventory'],self.host.service.ledger.inventory())
        self.assertEqual(self.manager.request('GET','/v1/capabilities')['protocol'],'777-service-v1')
        with self.assertRaises(APIError):self.host.service.handle('GET','/v1/me',self.manager.token)
    def test_role_projection_exact_calculation_front_zero_and_independent_order(self):
        result=self.sales.request('POST','/v1/calculate',{'snapshot':{'price_version_id':'fixture-v1','lines':[{'line_kind':'normal','product_id':'单件','q':'4'}]}})
        expected=calculated(calculate_quote(QuoteSnapshot('fixture-v1',(NormalLine('单件','4'),)),PriceSnapshot('fixture-v1',(PriceItem('单件','10','4','1',True),))),self.host.service.auth.authenticate(self.sales.token))
        self.assertEqual(result,expected)
        order=self.submit();data=self.sales.request('GET','/v1/orders/'+order['order_id']);self.assertEqual(data['details']['order_ref'],'00003');self.assertNotIn('cost_cents',data);self.assertNotIn('cost',data['lines'][0]['normalized'])
        full=self.manager.request('GET','/v1/orders/'+order['order_id']);self.assertEqual(full['cost_cents'],'1600');self.assertEqual(full,self.host.service.order(self.host.admin,order['order_id']))
    def test_same_request_two_threads_once_and_conflicting_payload_rejected(self):
        key=str(uuid.uuid4());data={'request_id':key,'snapshot':{'price_version_id':'fixture-v1','lines':[{'line_kind':'normal','product_id':'单件','q':'4'}]},'details':{'customer':'客户','order_ref':'00003','wechat':'微信'}}
        values=[];threads=[threading.Thread(target=lambda:values.append(self.sales.request('POST','/v1/orders',data))) for _ in range(2)]
        for t in threads:t.start()
        for t in threads:t.join()
        self.assertEqual(len(values),2);self.assertEqual(values[0],values[1]);self.assertEqual(self.host.service.ledger.inventory()['counts']['orders'],1)
        data['details']['customer']='changed'
        with self.assertRaises(APIError) as e:self.sales.request('POST','/v1/orders',data)
        self.assertEqual(e.exception.status,409)
    def test_native_commit_exit_original_journal_receipt_and_restart(self):
        j=Journal(self.host.root/'pending.dat') if not AUXILIARY else MemoryJournal();self.host.fault.write_text('after_commit')
        with self.assertRaises(CommitUnknown):self.submit(journal=j)
        self.assertIsNotNone(j.read())
        if not AUXILIARY:self.assertNotIn(b'00003',j.path.read_bytes())
        self.host.restart();result=self.sales.check_pending(j)['result'];self.assertEqual(result['income_cents'],'4000');self.assertIsNone(j.read());self.assertEqual(self.host.service.ledger.inventory()['counts']['orders'],1)
    def test_native_exit_before_commit_rolls_back_then_same_id_retries(self):
        j=MemoryJournal();self.host.fault.write_text('before_commit')
        with self.assertRaises(CommitUnknown):self.submit(journal=j)
        self.host.restart();self.assertIsNone(self.sales.request('GET','/v1/requests/'+j.read()['body']['request_id'])['result']);self.assertEqual(self.sales.check_pending(j,retry=True)['result']['income_cents'],'4000');self.assertEqual(self.host.service.ledger.inventory()['counts']['orders'],1)
    def test_sessions_revocation_role_change_last_admin_and_persistent_sessions(self):
        users=self.manager.request('GET','/v1/accounts')['accounts'];sale=next(u for u in users if u['username']=='sales-one');admin=next(u for u in users if u['username']=='manager')
        with self.assertRaises(APIError) as e:self.manager.request('POST','/v1/accounts/'+admin['id'],{'enabled':False})
        self.assertEqual(e.exception.status,409);self.host.restart();self.assertEqual(self.sales.request('GET','/v1/me')['role'],'sales')
        self.manager.request('POST','/v1/accounts/'+sale['id'],{'role':'admin'})
        with self.assertRaises(APIError):self.sales.request('GET','/v1/me')
        self.sales.login('sales-one',self.host.password);self.assertEqual(self.sales.request('GET','/v1/me')['role'],'admin');self.sales.request('POST','/v1/logout',{})
        with self.assertRaises(APIError):self.sales.request('GET','/v1/me')
    def test_account_casefold_scrypt_compatible_with_existing_auth(self):
        self.manager.request('POST','/v1/accounts',{'username':'Straße','password':'fixture-secret-123','role':'sales'});c=self.host.client('STRASSE',login=False);c.login('STRASSE','fixture-secret-123');self.assertEqual(c.request('GET','/v1/me')['role'],'sales')
        old=self.host.service.auth.login('straße','fixture-secret-123','independent');self.assertEqual(self.host.service.auth.authenticate(old['token']).subject,c.request('GET','/v1/me')['subject'])
    def test_sales_cannot_change_settings_prices_accounts_backup_or_other_owner(self):
        other=self.submit(self.host.client('sales-two'))
        for path in ['/v1/accounts','/v1/backup','/v1/orders/'+other['order_id'],'/v1/requests/'+other['request_id']]:
            with self.subTest(path=path),self.assertRaises(APIError) as e:self.sales.request('GET',path)
            self.assertIn(e.exception.status,(403,404))
        with self.assertRaises(APIError):self.sales.request('POST','/v1/settings',{'revision':1,'export_dir':'\\\\host\\share\\x','image_dir':'\\\\host\\share\\y'})
        with self.assertRaises(APIError):self.sales.request('POST','/v1/prices',{'request_id':str(uuid.uuid4()),'expected_price_version_id':'fixture-v1','prices':{'price_version_id':'bad','items':[]}})
    def test_settings_cas_sales_inherit_and_invalid_drive_directory(self):
        settings=self.manager.request('GET','/v1/settings');directory='\\\\host\\share\\reports' if os.name=='nt' else str(self.host.root/'reports')
        saved=self.manager.request('POST','/v1/settings',{'revision':settings['revision'],'export_dir':directory,'image_dir':directory});self.assertEqual(saved,self.sales.request('GET','/v1/settings'))
        with self.assertRaises(APIError) as e:self.manager.request('POST','/v1/settings',{'revision':settings['revision'],'export_dir':directory,'image_dir':directory})
        self.assertEqual(e.exception.status,409)
        with self.assertRaises(APIError):self.manager.request('POST','/v1/settings',{'revision':saved['revision'],'export_dir':'E:\\reports','image_dir':'E:\\images'})
    def test_pagination_snapshot_search_totals_unicode_and_cursor_restrictions(self):
        for _ in range(3):self.submit()
        other=self.submit(self.host.client('sales-two'));first=self.sales.request('GET','/v1/orders?limit=2&q=%E4%B8%AD%E6%96%87');self.assertEqual(first['total'],3);self.assertEqual(first['totals'],{'income_cents':'12000'});self.assertIsNotNone(first['next_cursor'])
        second=self.sales.request('GET','/v1/orders?limit=2&q=%E4%B8%AD%E6%96%87&cursor='+first['next_cursor']);self.assertEqual(len(second['orders']),1);self.assertNotIn(other['order_id'],[o['order_id'] for o in first['orders']+second['orders']])
        self.submit()
        with self.assertRaises(APIError) as e:self.sales.request('GET','/v1/orders?limit=2&q=%E4%B8%AD%E6%96%87&cursor='+first['next_cursor'])
        self.assertEqual(e.exception.status,409)
        for query in ['limit=201','limit=1&limit=2','format=future','q=%25_','offset=1000001','cursor=broken']:
            if query=='q=%25_':self.assertEqual(self.sales.request('GET','/v1/orders?'+query)['total'],0)
            else:
                with self.assertRaises(APIError):self.sales.request('GET','/v1/orders?'+query)
    def test_note_void_restore_cas_totals_preserve_receipts_and_outputs(self):
        order=self.submit();note=self.sales.mutate('/v1/notes',{'order_id':order['order_id'],'body':'=字面备注','expected_revision':1},MemoryJournal());self.assertEqual(note['revision'],2)
        void=self.manager.mutate('/v1/state',{'order_id':order['order_id'],'state':'void','reason':'测试作废','expected_revision':2},MemoryJournal());self.assertEqual(void['revision'],3);self.assertEqual(self.sales.request('GET','/v1/orders')['totals']['income_cents'],'0')
        self.manager.mutate('/v1/state',{'order_id':order['order_id'],'state':'active','reason':'恢复','expected_revision':3},MemoryJournal());self.assertEqual(self.sales.request('GET','/v1/orders')['totals']['income_cents'],'4000');self.assertEqual(self.sales.request('GET','/v1/requests/'+order['request_id'])['result'],order)
    def test_native_new_excel_full_independent_fields_styles_cost_and_public_directory(self):
        from openpyxl import load_workbook
        public=self.directories();order=self.submit();before=self.host.service.ledger.inventory()
        for c,kind in [(self.sales,'excel_quote'),(self.manager,'excel_admin')]:
            task=self.output(order,kind,c);self.assertEqual(self.wait(task['job_id'],c)['state'],'done');content=c.request('GET','/v1/outputs/'+task['job_id']+'/content');raw=base64.b64decode(content['data_base64']);self.assertEqual(hashlib.sha256(raw).hexdigest(),content['sha256']);book=load_workbook(io.BytesIO(raw));fields=dict(list(book['Order'].values)[1:]);self.assertEqual(fields['order_ref'],'00003');self.assertEqual(fields['income_cents'],'4000');self.assertEqual('cost_cents' in fields,kind=='excel_admin')
            for sheet in book:
                self.assertEqual(sheet.freeze_panes,'A2')
                for row in sheet:
                    for cell in row:self.assertEqual(cell.number_format,'@');self.assertNotEqual(cell.data_type,'f')
            book.close()
            if kind=='excel_admin':
                with self.assertRaises(APIError):self.sales.request('GET','/v1/outputs/'+task['job_id']+'/content')
                self.assertFalse((public/('777-'+order['order_id']+'-'+task['job_id']+'.xlsx')).exists())
            else:self.assertEqual((public/('777-'+order['order_id']+'-'+task['job_id']+'.xlsx')).read_bytes(),raw)
        self.assertEqual(self.host.service.ledger.inventory(),before)
    def test_failed_publication_keeps_frozen_blob_receipt_and_retry_not_rebook(self):
        order=self.submit();before=self.host.service.ledger.inventory();task=self.output(order);state=self.wait(task['job_id']);self.assertEqual(state['state'],'failed');content=self.sales.request('GET','/v1/outputs/'+task['job_id']+'/content');self.directories();retry=self.sales.mutate('/v1/outputs/'+task['job_id']+'/retry',{},MemoryJournal());self.assertEqual(self.wait(task['job_id'])['state'],'done');self.assertEqual(self.sales.request('GET','/v1/outputs/'+task['job_id']+'/content'),content);self.assertEqual(self.host.service.ledger.inventory(),before);self.assertEqual(self.sales.request('GET','/v1/requests/'+retry['request_id'])['result']['job_id'],task['job_id'])
    def test_financial_output_ids_never_cross_replay_domains(self):
        order=self.submit()
        with self.assertRaises(APIError) as e:self.sales.request('POST','/v1/outputs',{'request_id':order['request_id'],'order_id':order['order_id'],'expected_revision':1,'kind':'excel_quote'})
        self.assertEqual(e.exception.status,409)
    def test_daily_data_restore_preserves_every_ledger_field_and_output_without_old_sessions(self):
        from native_host_fixture import backup_pair,restore_data_backup
        from contextlib import closing
        import sqlite3
        self.directories();order=self.submit();task=self.output(order);self.assertEqual(self.wait(task['job_id'])['state'],'done')
        self.sales.mutate('/v1/notes',dict(order_id=order['order_id'],body='=中文原始备注',expected_revision=1),MemoryJournal())
        self.manager.mutate('/v1/state',dict(order_id=order['order_id'],state='void',reason='历史作废',expected_revision=2),MemoryJournal())
        self.manager.mutate('/v1/state',dict(order_id=order['order_id'],state='active',reason='历史恢复',expected_revision=3),MemoryJournal())
        entry=self.manager.request('POST','/v1/backup',{'kind':'daily'});folder=backup_pair(self.host,entry)
        data=json.loads((folder/'data.qbackup').read_bytes())['data'];before=self.host.service.ledger.inventory()
        target=self.host.root/'native-restored';native=restore_data_backup(self.host,entry,target)
        self.assertEqual(native['inventory'],before);self.assertEqual(command('inspect',root=target)['inventory'],before)
        with closing(sqlite3.connect(target/'private/ledger.sqlite3')) as c:
            for name,table in data['ledger'].items():self.assertEqual([list(r) for r in c.execute('SELECT * FROM '+name+' ORDER BY rowid')],table['rows'])
        with closing(sqlite3.connect(target/'private/auth.sqlite3')) as c:
            self.assertEqual(c.execute('SELECT count(*) FROM sessions').fetchone()[0],0)
            self.assertEqual(c.execute('SELECT username FROM users WHERE enabled=1').fetchall(),[('recovered-manager',)])
        for name,raw in data['blobs'].items():self.assertEqual((target/'private'/name).read_bytes(),base64.b64decode(raw))
        from openpyxl import load_workbook
        book=load_workbook(folder/'data.xlsx',read_only=True)
        try:
            for name,table in data['ledger'].items():
                rows=list(book['账本_'+name].values)
                self.assertEqual(list(rows[0]),table['columns'])
                self.assertEqual([list(r) for r in rows[1:]],[[('〈空值〉' if v is None else str(v) if not isinstance(v,str) else v or None) for v in r] for r in table['rows']])
        finally:book.close()
    def test_corrupt_or_future_data_refused_without_target_creation(self):
        from native_host_fixture import backup_pair,restore_data_backup
        entry=self.manager.request('POST','/v1/backup',{'kind':'daily'});file=backup_pair(self.host,entry)/'data.qbackup';raw=file.read_bytes()
        for field,value in [('sha256','0'*64),('format','777-data-backup-v999')]:
            data=json.loads(raw);data[field]=value;file.write_text(json.dumps(data),encoding='utf-8')
            with self.assertRaises(RuntimeError):restore_data_backup(self.host,entry,self.host.root/'refused')
            self.assertFalse((self.host.root/'refused').exists())
        file.write_bytes(raw)
    def test_owner_lock_blocks_old_python_and_second_rust_start(self):
        with self.assertRaises(FileBusy):
            with file_lock(self.host.host_root/'service-owner'):pass
        args=[str(BINARY),'run','--root',str(self.host.host_root)]+(['--auxiliary'] if AUXILIARY else [])
        p=subprocess.run(args,capture_output=True,timeout=15);self.assertNotEqual(p.returncode,0);self.assertIn(b'SERVICE_OWNER_BUSY',p.stderr)
    def test_raw_invalid_json_duplicate_origin_headers_and_oversized_body_fail(self):
        context=ssl.create_default_context(cafile=str(self.host.host_root/'service-cert.pem'))
        def send(body,extra=b'',headers=b'Content-Type: application/json\r\n'):
            with socket.create_connection(('127.0.0.1',self.host.port),timeout=10) as sock,context.wrap_socket(sock,server_hostname='127.0.0.1') as tls:
                tls.sendall(b'POST /v1/login HTTP/1.1\r\nHost: 127.0.0.1\r\nConnection: close\r\n'+headers+b'Content-Length: '+str(len(body)).encode()+b'\r\n'+extra+b'\r\n'+body);raw=b''
                while True:
                    data=tls.recv(65536)
                    if not data:break
                    raw+=data
                return int(raw.split(b' ',2)[1])
        for raw in [b'{"username":"x","username":"x","password":"bad"}',b'{"username":"x","password":NaN}',b'{"username":"\\ud800","password":"bad"}']:
            self.assertEqual(send(raw),422)
        self.assertEqual(send(b'{}',b'Origin: https://example.com\r\n'),403);self.assertEqual(send(b'{}',headers=b'Content-Type: application/json\r\nContent-Type: application/json\r\n'),400);self.assertEqual(send(b' '*(1024*1024+1)),413)

    def test_native_png_is_text_readable_on_windows_and_unavailable_font_is_explicit(self):
        from PIL import Image
        public=self.directories();order=self.submit();before=self.host.service.ledger.inventory();task=self.output(order,'quote_png');state=self.wait(task['job_id'])
        if AUXILIARY and not self.host.font:
            self.assertEqual(state['state'],'failed');self.assertEqual(state['error'],'OUTPUT_CJK_FONT_UNAVAILABLE')
        else:
            self.assertEqual(state['state'],'done');content=self.sales.request('GET','/v1/outputs/'+task['job_id']+'/content');raw=base64.b64decode(content['data_base64'])
            with Image.open(io.BytesIO(raw)) as image:
                self.assertEqual(image.format,'PNG');self.assertEqual(image.width,1200);self.assertEqual(image.info['777_format'],'777-output-v1');self.assertEqual(image.info['777_job_id'],task['job_id']);self.assertEqual(image.info['777_payload_sha256'],self._payload_hash(task['job_id']));image.verify()
            self.assertEqual((public/('777-'+order['order_id']+'-'+task['job_id']+'.png')).read_bytes(),raw)
        self.assertEqual(self.host.service.ledger.inventory(),before)
    def _payload_hash(self,key):
        with self.host.service.outputs.connection(readonly=True) as c:return c.execute('SELECT payload_sha256 FROM jobs WHERE job_id=?',(key,)).fetchone()[0]
    def test_offline_native_selection_and_rollback_require_ownership_and_preserve_inventory(self):
        before=self.host.service.ledger.inventory()
        with self.assertRaises(RuntimeError):command('select-host',root=self.host.host_root,secret=self.host.password,admin='manager',backend='python')
        self.host.stop();selected=command('select-host',root=self.host.host_root,secret=self.host.password,admin='manager',backend='python');self.assertEqual(selected['service_backend'],'python');self.assertEqual(command('inspect',root=self.host.host_root)['inventory'],before)
        selected=command('select-host',root=self.host.host_root,secret=self.host.password,admin='manager',backend='rust',notices=ROOT/'THIRD_PARTY_NOTICES.md');self.assertEqual(selected['service_backend'],'rust');self.host.start();self.assertEqual(self.sales.request('GET','/v1/orders')['total'],0);self.assertEqual(self.host.service.ledger.inventory(),before)
    def test_native_known_resource_hash_failure_never_launches_python_or_overwrites(self):
        self.host.stop();path=self.host.host_root/'private/rust-service-host.exe';path.write_bytes(b'corrupt known host')
        with self.assertRaises(RuntimeError):command('inspect',root=self.host.host_root)
        self.assertEqual(path.read_bytes(),b'corrupt known host')
    def test_login_rate_limit_and_tls_identity_mismatch_are_not_bypassed(self):
        bad=self.host.client('missing-user',login=False)
        for _ in range(5):
            with self.assertRaises(APIError) as e:bad.login('missing-user','wrong-secret-123')
            self.assertEqual(e.exception.status,401)
        with self.assertRaises(APIError) as e:bad.login('missing-user','wrong-secret-123')
        self.assertEqual(e.exception.status,429)
        bad.settings['certificate_sha256']='0'*64
        with self.assertRaises(APIError) as e:bad.request('GET','/v1/me')
        self.assertEqual(e.exception.status,403)
    def test_native_http_money_above_javascript_integer_range_is_exact_text(self):
        self.manager.request('POST','/v1/prices',{'request_id':str(uuid.uuid4()),'expected_price_version_id':'fixture-v1','prices':{'price_version_id':'large-v2','items':[{'product_id':'单件','sale':'100000000','cost':'0','k':'1','single':True}]}})
        snapshot=QuoteSnapshot('large-v2',(NormalLine('单件','922337203.6854775807'),));order=self.sales.submit(snapshot,DraftDetails('大整数','00003','微信'),MemoryJournal());self.assertEqual(order['income_cents'],'9223372036854775807');self.assertEqual(self.sales.request('GET','/v1/orders')['totals']['income_cents'],'9223372036854775807');self.assertEqual(self.host.service.ledger.inventory()['active_income_cents'],'9223372036854775807')
    def test_query_integer_unicode_digits_and_underscores_preserve_old_wire_semantics(self):
        self.assertEqual(self.sales.request('GET','/v1/orders?limit=%EF%BC%92')['limit'],2)
        self.assertEqual(self.sales.request('GET','/v1/orders?limit=2_0')['limit'],20)
        self.assertEqual(self.sales.request('GET','/v1/orders?limit=%D9%A2')['limit'],2)
        with self.assertRaises(APIError):self.sales.request('GET','/v1/orders?limit=2__0')
    def test_idle_tls_connections_are_bounded_and_capacity_is_released(self):
        context=ssl.create_default_context(cafile=str(self.host.host_root/'service-cert.pem'));connections=[]
        try:
            for _ in range(16):connections.append(context.wrap_socket(socket.create_connection(('127.0.0.1',self.host.port),timeout=3),server_hostname='127.0.0.1'))
            with self.assertRaises((ssl.SSLError,OSError)):
                sock=socket.create_connection(('127.0.0.1',self.host.port),timeout=3)
                try:context.wrap_socket(sock,server_hostname='127.0.0.1')
                finally:sock.close()
        finally:
            for c in connections:c.close()
        self.assertEqual(self.manager.request('GET','/v1/me')['role'],'admin')
    def test_demoted_cost_job_owner_cannot_read_private_retry_receipts_in_either_reader(self):
        users=self.manager.request('GET','/v1/accounts')['accounts'];other=next(u for u in users if u['username']=='sales-two');current=next(u for u in users if u['username']=='manager');self.manager.request('POST','/v1/accounts/'+other['id'],{'role':'admin'})
        order=self.submit(self.manager);job=self.output(order,'excel_admin',self.manager);self.assertEqual(self.wait(job['job_id'],self.manager)['state'],'done');retry=self.manager.mutate('/v1/outputs/'+job['job_id']+'/retry',{},MemoryJournal());self.assertEqual(self.wait(job['job_id'],self.manager)['state'],'done')
        promoted=self.host.client('sales-two');promoted.request('POST','/v1/accounts/'+current['id'],{'role':'sales'});self.manager.login('manager',self.host.password)
        with self.assertRaises(APIError) as e:self.manager.request('GET','/v1/requests/'+retry['request_id'])
        self.assertEqual(e.exception.status,404)
        actor=self.host.service.auth.authenticate(self.manager.token)
        with self.assertRaises(APIError):self.host.service.outputs.receipt(actor,retry['request_id'])


class NativeHostLifecycleTests(unittest.TestCase):
    """Real child pipes; separate from financial/HTTPS acceptance tests."""
    def child(self, source):
        host = Hosted.__new__(Hosted)
        process = subprocess.Popen([sys.executable, '-c', source],
                                   stdin=subprocess.PIPE, stdout=subprocess.PIPE,
                                   stderr=subprocess.PIPE)
        host.process = process
        def cleanup():
            if process.poll() is None:
                process.kill()
            process.communicate(timeout=10)
        self.addCleanup(cleanup)
        return host, process

    def test_graceful_stop_sends_newline_and_reaps_child(self):
        host, process = self.child("import sys; assert sys.stdin.readline() == '\\n'")
        host.stop()
        self.assertEqual(process.returncode, 0)
        self.assertIsNone(host.process)
        self.assertTrue(process.stdin.closed)
        host.stop()

    def test_closed_pipe_before_exit_retains_both_injected_exit_codes(self):
        for code in (17, 18):
            for newline in ('\n', '\r\n'):
                with self.subTest(exit_code=code, newline=repr(newline)):
                    host, process = self.child(
                        "import os,sys,time; os.close(0); "
                        "sys.stdout.reconfigure(newline=" + repr(newline) + "); "
                        "print('closed', flush=True); time.sleep(.1); sys.exit(" + str(code) + ")")
                    self.assertIn(process.stdout.readline(), (b'closed\n', b'closed\r\n'))
                    host.stop()
                    self.assertEqual(process.returncode, code)
                    self.assertIsNone(host.process)
                    self.assertTrue(process.stdin.closed)

    def test_already_exited_child_is_reaped_and_repeated_stop_is_safe(self):
        host, process = self.child('raise SystemExit(17)')
        process.wait(timeout=10)
        host.stop()
        host.stop()
        self.assertIsNone(host.process)
        self.assertTrue(process.stdin.closed)

    def test_unexpected_exit_is_not_silenced_by_pipe_cleanup(self):
        host, process = self.child("import sys; print('unexpected fault', file=sys.stderr); sys.exit(29)")
        with self.assertRaisesRegex(RuntimeError, 'unexpected fault'):
            host.stop()
        self.assertEqual(process.returncode, 29)
        self.assertIs(host.process, process)

    def test_timeout_retains_process_and_sandbox_for_diagnosis(self):
        host = Hosted.__new__(Hosted)
        from unittest.mock import Mock
        host.process = Mock()
        host.process.communicate.side_effect = subprocess.TimeoutExpired('host', 50)
        process = host.process
        with self.assertRaisesRegex(RuntimeError, 'sandbox retained'):
            host.stop()
        self.assertIs(host.process, process)
        process.kill.assert_not_called()
