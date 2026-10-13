"""Only synthetic legacy/current formats, using the accepted export protocol."""
from dataclasses import asdict
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
import io
import uuid

from openpyxl import Workbook
from quote_core import (PriceItem, PriceSnapshot, QuoteSnapshot, NormalLine, DiscountLine,
                        DraftDetails, calculate_quote)
from quote_core.numbers import money_text, settle, RULE_VERSION, NUMERIC_PROFILE
from quote_app.projection import append_quote
from quote_files.schema import DETAIL_HEADERS, PRICE_HISTORY_HEADERS, create_template
from quote_files.workbooks import (workbook_transaction, request, assign_order,
                                  table, open_book)
from quote_files.orders import edit_order

PRICES = PriceSnapshot('synthetic-historical-v1',
                       (PriceItem('单件', '10', '4', '', True), PriceItem('面积', '12.5', '5', '1', False)))


def recompute_summary(ws, h):
    totals = [sum(settle(Decimal(str(ws.cell(i, h[f]).value or 0))) for i in range(7, ws.max_row + 1)
                  if ws.cell(i, h['行类型']).value in ('normal', 'discount')
                  and ws.cell(i, h['记录状态']).value != 'void') for f in ('销售额', '总成本', '利润')]
    for col, amount in zip((2, 4, 6), totals):
        ws.cell(4, col, money_text(amount))


def current_book(path, orders=None):
    orders = orders or [(NormalLine('单件', '4', name='原单件'), DiscountLine('90%')),
                       (NormalLine('面积', '2', '1.5', '2', '原面积'),)]
    for lines in orders:
        snapshot = QuoteSnapshot(PRICES.price_version_id, tuple(lines))
        result = calculate_quote(snapshot, PRICES)
        assert result.status == 'ok'
        payload = SimpleNamespace(result=result, prices=PRICES, details=DraftDetails('同秒客户', '同秒订单', '原微信'))
        req = request('append_quote', asdict(snapshot), order_id=str(uuid.uuid4()))
        def create():
            book = Workbook()
            create_template(book.active, '2026-10-13 12:00:00')
            return book
        def mutate(book, ws, h):
            start = ws.max_row + 1
            append_quote(ws, '2026-10-13 12:00:00', payload, kind_column=h['行类型'], headers=h)
            assign_order(ws, start, ws.max_row, h, req)
            recompute_summary(ws, h)
            history = book['订单价格快照'] if '订单价格快照' in book.sheetnames else book.create_sheet('订单价格快照')
            if history.cell(1, 1).value is None:
                for column, field in enumerate(PRICE_HISTORY_HEADERS, 1):
                    history.cell(1, column, field)
            for p in PRICES.items:
                history.append([req.order_id, req.request_id, '2026-10-13 12:00:00', PRICES.price_version_id,
                                RULE_VERSION, NUMERIC_PROFILE, p.product_id, p.sale, p.cost, p.k, '是' if p.single else '否'])
            return {'order_revision': 1, 'price_version_id': PRICES.price_version_id}
        workbook_transaction(path, req, DETAIL_HEADERS, create, mutate)
    return Path(path)


def append_note_and_void(path):
    for operation, revision in [('append_remark', 1), ('void_order', 2)]:
        book = open_book(path.read_bytes())
        try:
            _, h = table(book['利润明细'])
            oid = book['利润明细'].cell(7, h['订单ID']).value
        finally:
            book.close()
        req = request(operation, {'synthetic': True}, order_id=oid)
        def mutate(book, ws, headers):
            result = edit_order(book, ws, headers, req, revision, text='=原始文字备注', reason='原作废理由')
            recompute_summary(ws, headers)
            return result
        workbook_transaction(path, req, DETAIL_HEADERS, lambda: None, mutate)
    return path


def legacy_bytes(raw, *, tagged=False, reorder=False, alias=False, footers=True):
    source = open_book(raw)
    try:
        ws = source['利润明细']
        _, old = table(ws)
        book = Workbook();dest = book.active;create_template(dest, '2026-10-13 12:00:00')
        fields = list(DETAIL_HEADERS) + (['行类型'] if tagged else [])
        if alias:
            fields[1] = '用户'
        if reorder:
            fields.reverse()
        for column, field in enumerate(fields, 1):
            dest.cell(6, column, field)
        for i in range(7, ws.max_row + 1):
            if not footers and ws.cell(i, old['行类型']).value == 'order_total':
                continue
            dest.append([ws.cell(i, old['淘宝ID' if field == '用户' else field]).value for field in fields])
        for col in (2, 4, 6):
            dest.cell(4, col, ws.cell(4, col).value)
        output = io.BytesIO();book.save(output);book.close()
        return output.getvalue()
    finally:
        source.close()


def mutate_bytes(raw, mutate):
    book = open_book(raw)
    try:
        mutate(book)
        stream = io.BytesIO();book.save(stream)
        return stream.getvalue()
    finally:
        book.close()
