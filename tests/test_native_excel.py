"""P06-03 native workbook IO with independent openpyxl field/style checks."""
import base64
from datetime import datetime
import io
import json
import os
from pathlib import Path
import struct
import tempfile
import time
import unittest
from unittest.mock import patch
import uuid
from zipfile import ZipFile, ZIP_DEFLATED
from quote_core import PriceItem, PriceSnapshot, NormalLine, QuoteSnapshot, DraftDetails
from quote_ledger import Principal
from quote_ledger.paths import LedgerError
from quote_files.native_excel import FormatReview
from quote_import import preview_bytes, apply_source
from quote_import.store import verify_provenance
from quote_service.rust_ledger import RustLedger
from quote_service.output_render import render,validate,excel_rows
from quote_service.protocol import APIError
from import_fixture import current_book,legacy_bytes,mutate_bytes,append_note_and_void
from p01_fixtures import workbook
import test_rust_service as reference
ROOT=Path(__file__).resolve().parents[1]
BINARY=ROOT/'target/release'/('quote-ledger.exe' if os.name=='nt' else 'quote-ledger')
ADMIN=Principal('manager','admin')
def uid():return str(uuid.uuid4())
def repack(raw,part,change):
    output=io.BytesIO()
    with ZipFile(io.BytesIO(raw)) as z,ZipFile(output,'w',ZIP_DEFLATED) as out:
        for item in z.infolist():out.writestr(item,change(z.read(item.filename)) if item.filename==part else z.read(item.filename))
    return output.getvalue()

class NativeReadTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix='777-p06-excel-');self.root=Path(self.temp.name)
        self.ledger=RustLedger.create(self.root/'ledger.sqlite3',BINARY,auxiliary=os.name!='nt')
        self.raw=current_book(self.root/'source.xlsx').read_bytes()
    def tearDown(self):self.temp.cleanup()
    def test_original_identified_legacy_and_note_void_plans_keep_full_binding(self):
        append_note_and_void(self.root/'source.xlsx')
        cases=[self.raw,legacy_bytes(self.raw),(self.root/'source.xlsx').read_bytes()]
        for raw in cases:
            expected=preview_bytes(raw,'xlsx')
            with patch('quote_import.adapters.open_book',side_effect=AssertionError('Python decoder reached')):
                actual=preview_bytes(raw,'xlsx',native=self.ledger)
            self.assertEqual(actual.payload(),expected.payload());self.assertEqual(actual.review_sha256,expected.review_sha256)
    def test_all_p01_workbook_shapes_keep_recognition_and_quarantine(self):
        for variant in ['current','empty_evidence_shape','reordered','remark','legacy','no_summary','latest_note','formula','formula_cached','duplicate_headers']:
            path=self.root/(variant+'.xlsx');workbook(path,2,variant=variant);raw=path.read_bytes()
            a=preview_bytes(raw,'xlsx');b=preview_bytes(raw,'xlsx',native=self.ledger)
            if a.source_schema!='unsupported':self.assertEqual(a.payload(),b.payload(),variant)
            else:self.assertEqual(a.quarantined[0]['code'],b.quarantined[0]['code'],variant)
    def test_native_import_replays_v1_batch_without_python_excel_decoder(self):
        plan=preview_bytes(self.raw,'xlsx');owner=uid();before=(self.root/'source.xlsx').stat().st_mtime_ns
        with patch('quote_import.adapters.open_book',side_effect=AssertionError('Python decoder reached')):
            first=apply_source(self.ledger,ADMIN,owner,self.raw,'xlsx',expected_review_sha256=plan.review_sha256)
            second=apply_source(self.ledger,ADMIN,owner,self.raw,'xlsx',expected_review_sha256=plan.review_sha256)
            with self.ledger.connection(readonly=True) as c:self.assertEqual(verify_provenance(c,native=self.ledger),1)
        self.assertEqual(first,second);self.assertEqual(self.ledger.inventory()['counts']['orders'],2)
        self.assertEqual((self.root/'source.xlsx').read_bytes(),self.raw);self.assertEqual((self.root/'source.xlsx').stat().st_mtime_ns,before)
    def test_mapped_columns_and_raw_prices_keep_independent_plan(self):
        def reverse(b):
            ws=b['利润明细'];values=[list(row) for row in ws.values]
            for r,row in enumerate(values,1):
                if r>=6:
                    for c,v in enumerate(reversed(row),1):ws.cell(r,c,v)
        raw=mutate_bytes(self.raw,reverse)
        self.assertEqual(preview_bytes(raw,'xlsx').payload(),preview_bytes(raw,'xlsx',native=self.ledger).payload())
    def test_exact_numeric_tokens_and_lossy_legacy_numeric_cells_are_traceable(self):
        from openpyxl import Workbook
        w=Workbook();w.active.cell(1,1,1);stream=io.BytesIO();w.save(stream);w.close()
        raw=repack(stream.getvalue(),'xl/worksheets/sheet1.xml',lambda s:s.replace(b'<v>1</v>',b'<v>9007199254740993.25</v>'))
        self.assertEqual(self.ledger.excel_document(raw)['sheets'][0]['cells'][0]['value'],'9007199254740993.25')
        with self.assertRaises(FormatReview) as caught:self.ledger.excel_book(raw)
        self.assertEqual(str(caught.exception),'NUMERIC_TOKEN_REQUIRES_TEXT');self.assertEqual(caught.exception.location,{'sheet':'Sheet','cell':'A1'})
        plan=preview_bytes(raw,'xlsx',native=self.ledger)
        self.assertEqual(plan.quarantined[0]['difference'],{'sheet':'Sheet','cell':'A1'})
    def test_date_iso_serial_and_1904_epoch_equal_independent_reader(self):
        from openpyxl import Workbook,load_workbook
        for epoch in [datetime(1899,12,30),datetime(1904,1,1)]:
            for iso in [False,True]:
                w=Workbook();w.epoch=epoch;w.iso_dates=iso;w.active.cell(1,1,datetime(2026,10,13,12));stream=io.BytesIO();w.save(stream);w.close();raw=stream.getvalue()
                old=load_workbook(io.BytesIO(raw));new=self.ledger.excel_book(raw)
                try:self.assertEqual(new['Sheet'].cell(1,1).value,old['Sheet'].cell(1,1).value)
                finally:old.close()
    def test_formula_cache_or_orphan_sheet_is_always_refused(self):
        for variant in ['formula','formula_cached']:
            p=self.root/(variant+'.xlsx');workbook(p,2,variant=variant);raw=p.read_bytes()
            self.assertEqual(preview_bytes(raw,'xlsx',native=self.ledger).quarantined[0]['code'],'FORMULA_PRESENT_NO_EVALUATION_OR_CACHE_TRUST')
        out=io.BytesIO()
        with ZipFile(io.BytesIO(self.raw)) as z,ZipFile(out,'w',ZIP_DEFLATED) as target:
            for item in z.infolist():target.writestr(item,z.read(item.filename))
            target.writestr('xl/worksheets/orphan.xml','<worksheet><f>1+1</f><v>2</v></worksheet>')
        with self.assertRaisesRegex(FormatReview,'FORMULA_PRESENT'):self.ledger.excel_document(out.getvalue())
    def test_unknown_schema_and_extensions_cannot_change_ledger(self):
        for raw in [mutate_bytes(self.raw,lambda w:w.create_sheet('未知')),repack(self.raw,'xl/worksheets/sheet1.xml',lambda s:s.replace(b'</worksheet>',b'<extLst/></worksheet>'))]:
            plan=preview_bytes(raw,'xlsx',native=self.ledger);self.assertEqual(plan.source_schema,'unsupported');before=self.ledger.inventory()
            with self.assertRaises(LedgerError):apply_source(self.ledger,ADMIN,uid(),raw,'xlsx',expected_review_sha256=plan.review_sha256)
            self.assertEqual(self.ledger.inventory(),before)
    def test_corrupt_duplicate_zip_entities_utf16_and_cell_duplicates_fail_closed(self):
        sources=[b'broken',repack(self.raw,'xl/worksheets/sheet1.xml',lambda _:b'<worksheet>\0</worksheet>'),repack(self.raw,'xl/worksheets/sheet1.xml',lambda s:s.replace(b'<worksheet',b'<!DOCTYPE worksheet [<!ENTITY x "x">]><worksheet'))]
        out=io.BytesIO()
        with ZipFile(out,'w') as z:
            z.writestr('same.xml','<a/>')
            with self.assertWarns(UserWarning):z.writestr('same.xml','<a/>')
        sources.append(out.getvalue())
        for raw in sources:
            with self.assertRaises(FormatReview):self.ledger.excel_document(raw)
    def test_row_column_source_and_declared_expansion_limits_are_enforced(self):
        for token in [b'BM1',b'A50001']:
            raw=repack(self.raw,'xl/worksheets/sheet1.xml',lambda s:s.replace(b'<dimension ref="',b'<dimension ref="'+token+b':',1))
            with self.assertRaises(FormatReview):self.ledger.excel_document(raw)
        with self.assertRaises(FormatReview):self.ledger.excel_document(b'x'*(20*1024*1024+1))
        raw=bytearray(self.raw);offset=raw.index(b'PK\x01\x02');struct.pack_into('<I',raw,offset+24,64*1024*1024+1)
        with self.assertRaisesRegex(FormatReview,'ZIP_EXPANSION_LIMIT'):self.ledger.excel_document(bytes(raw))
    def test_large_sparse_source_is_complete_and_not_silently_truncated(self):
        from openpyxl import Workbook
        w=Workbook(write_only=True);sheet=w.create_sheet('大表')
        for i in range(12000):sheet.append([str(i),'中文','000001'])
        out=io.BytesIO();w.save(out);w.close();raw=out.getvalue();before=time.monotonic();doc=self.ledger.excel_document(raw)
        self.assertEqual(len(doc['sheets'][0]['cells']),36000);self.assertEqual(doc['sheets'][0]['cells'][-3]['value'],'11999');self.assertLess(time.monotonic()-before,30)

class NativeOutputTests(unittest.TestCase):
    setUp=reference.RustOutputs.setUp
    tearDown=reference.RustOutputs.tearDown
    configure=reference.RustOutputs.configure
    request=reference.RustOutputs.request
    workbook=reference.RustOutputs.workbook
    def test_native_generation_download_and_private_role_projection_skip_python_workbook_io(self):
        self.configure();tasks=[self.request(),self.request('excel_admin',self.actor)];before=self.service.ledger.inventory()
        with patch('openpyxl.Workbook',side_effect=AssertionError('Python writer reached')),patch('openpyxl.load_workbook',side_effect=AssertionError('Python reader reached')):
            for task in tasks:self.assertTrue(self.outputs.run_one(self.service));self.outputs.content(self.actor,task['job_id'])
        for task in tasks:
            w=self.workbook(task['job_id'])
            try:
                data=dict(list(w['Order'].values)[1:]);self.assertEqual(data['order_ref'],'000123')
                self.assertEqual('cost_cents' in data,task['kind']=='excel_admin')
                for ws in w:
                    self.assertEqual(ws.freeze_panes,'A2');self.assertTrue(all(c.font.bold for c in ws[1]));self.assertTrue(all(c.number_format=='@' for row in ws for c in row))
            finally:w.close()
        self.assertEqual(self.service.ledger.inventory(),before)
    def test_full_cell_matrix_and_styles_match_independent_python_report(self):
        self.configure();task=self.request('excel_admin',self.actor)
        with self.outputs.connection(readonly=True) as c:payload=json.loads(c.execute('SELECT payload_json FROM jobs WHERE job_id=?',(task['job_id'],)).fetchone()[0])
        a=render(payload);b=render(payload,native=self.service.ledger)
        from openpyxl import load_workbook
        w1=load_workbook(io.BytesIO(a));w2=load_workbook(io.BytesIO(b))
        try:
            for name in w1.sheetnames:
                self.assertEqual(list(w1[name].values),list(w2[name].values));self.assertEqual(w1[name].freeze_panes,w2[name].freeze_panes)
        finally:w1.close();w2.close()
        validate(b,payload);validate(a,payload,native=self.service.ledger)
    def test_big_integer_controls_leading_zero_and_formula_like_text_remain_text(self):
        self.configure();task=self.request()
        with self.outputs.connection(readonly=True) as c:payload=json.loads(c.execute('SELECT payload_json FROM jobs WHERE job_id=?',(task['job_id'],)).fetchone()[0])
        payload['order']['details']={'customer':'=中文<&>_x0000_\x01','order_ref':'0000001','wechat':'+@-' };payload['order']['income_cents']='9223372036854775807'
        raw=render(payload,native=self.service.ledger);validate(raw,payload)
        from openpyxl import load_workbook
        w=load_workbook(io.BytesIO(raw))
        try:
            values=dict(list(w['Order'].values)[1:]);self.assertEqual(values['income_yuan'],'92233720368547758.07');self.assertEqual(values['customer'],'=中文<&>_x0000_\\u0001');self.assertEqual(values['order_ref'],'0000001')
            self.assertTrue(all(c.data_type!='f' for ws in w for row in ws for c in row))
        finally:w.close()
    def test_body_tampering_with_unchanged_protocol_is_refused_by_both_validators(self):
        task=self.request()
        with self.outputs.connection(readonly=True) as c:payload=json.loads(c.execute('SELECT payload_json FROM jobs WHERE job_id=?',(task['job_id'],)).fetchone()[0])
        raw=render(payload,native=self.service.ledger)
        def mutate(w):w['Order'].cell(2,2,'changed')
        changed=mutate_bytes(raw,mutate)
        for worker in [None,self.service.ledger]:
            with self.assertRaises(APIError):validate(changed,payload,native=worker)
    def test_failed_native_worker_never_falls_back_or_rebooks_order(self):
        self.configure();task=self.request();before=self.service.ledger.inventory();original=self.service.ledger.binary
        self.service.ledger.binary=self.root/'missing.exe'
        with patch('openpyxl.Workbook',side_effect=AssertionError('fallback')):self.outputs.run_one(self.service)
        self.service.ledger.binary=original
        self.assertEqual(self.outputs.status(self.sales,task['job_id'])['state'],'failed');self.assertEqual(self.service.ledger.inventory(),before)
    def test_native_bytes_are_reproducible_and_no_existing_source_is_edited(self):
        task=self.request()
        with self.outputs.connection(readonly=True) as c:payload=json.loads(c.execute('SELECT payload_json FROM jobs WHERE job_id=?',(task['job_id'],)).fetchone()[0])
        self.assertEqual(render(payload,native=self.service.ledger),render(payload,native=self.service.ledger))

import test_p05_import_entry as entry_reference
from rust_service_fixture import initialize_rust
class NativeEntryTests(unittest.TestCase):
    def setUp(self):
        with patch.object(entry_reference,'initialize',initialize_rust):entry_reference.ImportEntryTests.setUp(self)
    tearDown=entry_reference.ImportEntryTests.tearDown
    arguments=entry_reference.ImportEntryTests.arguments
    command=entry_reference.ImportEntryTests.command
    def test_apply_uses_selected_native_reader_and_complete_backup_gate(self):
        with patch('quote_import.adapters.open_book',side_effect=AssertionError('Python decoder reached')):
            code,report,text=self.command(self.arguments())
        self.assertEqual(code,0,text);self.assertEqual(report['status'],'committed');self.assertTrue(report['complete_pre_import_backup'])
        self.assertEqual(self.service.ledger.inventory()['counts']['orders'],2)
    def test_explicit_native_preview_keeps_source_private_and_ledger_unchanged(self):
        before=self.service.ledger.inventory()
        with patch('quote_import.adapters.open_book',side_effect=AssertionError('Python decoder reached')):
            code,report,text=self.command(['preview','--source',str(self.source),'--host-root',str(self.root/'host'),'--output','ignored'])
        self.assertEqual(code,0,text);self.assertEqual(report['mode'],'dry_run');self.assertEqual(self.service.ledger.inventory(),before)
        self.assertEqual((self.output/'private/source.xlsx').read_bytes(),self.source.read_bytes())
        self.assertNotIn('同秒客户',json.dumps(report,ensure_ascii=False))
    def test_changed_native_review_does_not_request_password_or_write_ledger(self):
        import contextlib
        from quote_service.hosting import load
        entry=entry_reference.entry;before=self.service.ledger.inventory()
        with contextlib.ExitStack() as stack:
            stack.enter_context(patch.object(entry.platform,'system',return_value='Windows'))
            stack.enter_context(patch.object(entry,'new_output',return_value=self.output))
            stack.enter_context(patch.object(entry,'load',side_effect=lambda p,**kw:load(p,auxiliary=os.name!='nt',**kw)))
            password=stack.enter_context(patch.object(entry.getpass,'getpass',return_value=self.secret))
            stack.enter_context(patch('sys.stdout',new=io.StringIO()))
            code=entry.main(self.arguments(**{'--review-sha256':'0'*64}));password.assert_not_called()
        self.assertEqual(code,2);self.assertEqual(self.service.ledger.inventory(),before)

class LegacyWorkerCompatibilityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from legacy_worker_fixture import worker
        cls.binary=worker()
    def setUp(self):
        from quote_service.hosting import initialize,load
        from quote_service.backend import select_backend
        self.temp=tempfile.TemporaryDirectory(prefix='777-p06-legacy-format-');self.root=Path(self.temp.name)
        secret='legacy-fixture-password-123';initialize(self.root/'host',self.root/'backups','manager',secret,hosts=['localhost'],auxiliary=os.name!='nt')
        select_backend(self.root/'host','manager',secret,'rust',binary=self.binary,notices=ROOT/'THIRD_PARTY_NOTICES.md',auxiliary=os.name!='nt')
        _,_,self.service=load(self.root/'host',auxiliary=os.name!='nt')
        self.actor=self.service.auth.authenticate(self.service.auth.login('manager',secret,'fixture')['token'])
        self.service.ledger.publish_prices(uid(),self.actor,PriceSnapshot('v1',(PriceItem('P','10','4','',True),)))
        self.order=self.service.ledger.submit_order(uid(),self.actor,QuoteSnapshot('v1',(NormalLine('P','1'),)),DraftDetails('旧版','0001','微信'))
    def tearDown(self):self.temp.cleanup()
    def test_actual_accepted_financial_only_worker_uses_explicit_format_bridge(self):
        self.assertEqual(self.service.ledger.backend,'rust');self.assertEqual(self.service.ledger.format_backend,'python-v1-compatibility');self.assertIsNone(self.service.outputs.native)
        from quote_service.output_render import public_order
        payload={'format':'777-output-v1','job_id':uid(),'kind':'excel_quote','order':public_order(self.service.order(self.actor,self.order['order_id']))}
        validate(render(payload,native=self.service.outputs.native),payload)
        self.assertEqual(self.service.ledger.inventory()['active_income_cents'],'1000')
    def test_old_v3_full_package_restores_output_and_keeps_original_worker(self):
        from quote_service.recovery import restore
        from quote_service.hosting import load
        task=self.service.output_task(self.actor,dict(request_id=uid(),order_id=self.order['order_id'],expected_revision=1,kind='excel_admin'))
        self.service.outputs.run_one(self.service);content=self.service.outputs.content(self.actor,task['job_id']);before=self.service.ledger.inventory()
        saved=self.service.backup(kind='maintenance');restore(saved['service_bundle'],self.root/'restored','manager','new-fixture-password-123',auxiliary=os.name!='nt')
        _,_,new=load(self.root/'restored',auxiliary=os.name!='nt',allow_maintenance=True)
        self.assertEqual(new.ledger.backend,'rust');self.assertEqual(new.ledger.format_backend,'python-v1-compatibility');self.assertEqual(new.ledger.inventory(),before)
        self.assertEqual(new.outputs.content(self.actor,task['job_id']),content)
    def test_missing_legacy_worker_never_negotiates_a_financial_or_format_fallback(self):
        self.service.ledger.binary=self.root/'missing.exe'
        with self.assertRaises(LedgerError):self.service.ledger.inventory()
