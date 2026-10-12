"""Measurement boundaries and real synthetic read counts, not Windows claims."""
from copy import deepcopy
import ast
import io
import json
from pathlib import Path
import sys
import tempfile
import time
import unittest
from types import SimpleNamespace
from unittest.mock import Mock, patch

import openpyxl
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet.worksheet import Worksheet
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / 'tools')]

import p01_fixtures as fixtures
from p04_metrics import Heartbeat, LatencyTrace, distribution, healthy_heartbeat, input_hashes, instrumentation
import p04_performance_check as perf
from quote_files.documents import parse_document


def configured_notebook(app, source):
    """Execute the actual four entrypoints' notebook-add statements unchanged."""
    calculator = next(n for n in ast.parse(source).body if isinstance(n, ast.ClassDef) and n.name == 'PriceCalculator')
    statements = [n for n in ast.walk(calculator) if isinstance(n, ast.Expr) and isinstance(n.value, ast.Call)
        and isinstance(n.value.func, ast.Attribute) and n.value.func.attr == 'add'
        and isinstance(n.value.func.value, ast.Attribute) and n.value.func.value.attr == 'notebook']
    class Notebook:
        def __init__(self):self.labels = {};self.selected = None
        def add(self, page, text):self.labels[page] = text
        def tabs(self):return tuple(self.labels)
        def tab(self, page, option):return self.labels[page]
        def select(self, page):
            if page not in self.labels:raise ValueError('Unknown notebook page')
            self.selected = page
    app.notebook = Notebook()
    for statement in statements:setattr(app, statement.value.args[0].attr, object())
    exec(compile(ast.Module(body=statements, type_ignores=[]), '<actual notebook configuration>', 'exec'), {'self': app})
    return app.notebook


class PerformanceMeasurementTests(unittest.TestCase):
    def test_read_turn_budget_resets_after_wakeup_and_restores_nested_error_scope(self):
        from quote_files import scheduling as s
        clock = [0.0]
        def pause(seconds):
            self.assertEqual(seconds, .001)
            clock[0] += .020  # An oversleep does not cause immediate re-sleep.
        with patch.object(s, 'WINDOWS', True), patch.object(s.time, 'perf_counter', side_effect=lambda: clock[0]), \
             patch.object(s.time, 'sleep', side_effect=pause) as sleeps:
            s.checkpoint();sleeps.assert_not_called()
            with self.assertRaisesRegex(RuntimeError, 'read failed'):
                with s.read_turns():
                    outer = s._state.turn
                    clock[0] = .007;s.checkpoint();sleeps.assert_not_called()
                    clock[0] = .008;s.checkpoint();self.assertEqual(sleeps.call_count, 1)
                    self.assertAlmostEqual(outer.deadline, .036)
                    with s.read_turns():
                        self.assertIs(s._state.turn, outer)
                        s.checkpoint();self.assertEqual(sleeps.call_count, 1)
                    raise RuntimeError('read failed')
            self.assertIsNone(s._state.turn)
            clock[0] = 100;s.checkpoint();self.assertEqual(sleeps.call_count, 1)
        with patch.object(s, 'WINDOWS', False), patch.object(s.time, 'sleep') as sleeps, s.read_turns():
            s.checkpoint();sleeps.assert_not_called();self.assertIsNone(s._state.turn)

    def test_jobs_read_turn_is_thread_local_and_never_reaches_writer_or_owner(self):
        from threading import Event
        from quote_app.jobs import Jobs
        from quote_files import scheduling as s
        held, release, written = Event(), Event(), Event()
        observed = {}
        def read(context):
            observed['read'] = s._state.turn
            held.set()
            if not release.wait(3):raise TimeoutError('read gate')
            raise ValueError('synthetic read failure')
        def write(context):
            observed['write'] = getattr(s._state, 'turn', None)
            written.set()
        jobs = Jobs(readers=1)
        try:
            with patch.object(s, 'WINDOWS', True):
                reader = jobs.submit('read', read)
                self.assertTrue(held.wait(3))
                writer = jobs.submit('write', write, write=True)
                self.assertTrue(written.wait(3))
                self.assertIsNone(getattr(s._state, 'turn', None))
                self.assertIsNone(observed['write']);self.assertIsNotNone(observed['read'])
                release.set()
                follow = jobs.submit('follow', lambda _: getattr(s._state, 'turn', None))
                deadline = time.monotonic() + 3
                events = []
                while follow.state in {'queued', 'running'} and time.monotonic() < deadline:
                    events.extend(jobs.drain());time.sleep(.002)
                events.extend(jobs.drain())
                self.assertEqual(reader.state, 'failed');self.assertEqual(writer.state, 'completed')
                value = next(v for t, state, v in events if t is follow and state == 'completed')
                self.assertIsNotNone(value);self.assertIsNot(value, observed['read'])
        finally:
            release.set()
            deadline = time.monotonic() + 3
            while jobs.writing and time.monotonic() < deadline:jobs.drain();time.sleep(.002)
            jobs.close();self.assertTrue(jobs.wait_closed(3))

    def test_actual_budgeted_decode_and_views_keep_exact_rows_amounts_and_single_open(self):
        from threading import Thread
        from quote_files import scheduling as s
        from quote_files.views import read_view
        from quote_app.pagination import filter_view
        path = self.root / 'budgeted.xlsx';fixtures.workbook(path, 1000)
        raw = path.read_bytes();reference = parse_document(raw)
        self.assertEqual(reference.result.status, 'ok')
        baseline = {role: read_view(path, None, role, SimpleNamespace(check=lambda: None), document=reference)
                    for role in ('manager', 'sales')}
        results, failures, pauses = [], [], []
        original_pause = s.ReadTurn.pause
        def pause(turn):
            pauses.append(time.perf_counter());return original_pause(turn)
        def read():
            try:
                with s.read_turns(), instrumentation() as counts:
                    doc = parse_document(raw)
                    self.assertEqual(dict(counts), {'workbook_opens': 1, 'worksheet_xml_source_opens': 1,
                                                    'iter_rows_calls': 1, 'yielded_rows': 1006})
                    views = {role: read_view(path, None, role, SimpleNamespace(check=lambda: None), document=doc)
                             for role in ('manager', 'sales')}
                    last = {role: filter_view(view, '合成订单000999') for role, view in views.items()}
                    self.assertEqual(counts['workbook_opens'], 1)
                    results.append((doc, views, last))
            except BaseException as exc:failures.append(exc)
        with patch.object(s, 'WINDOWS', True), patch.object(s.ReadTurn, 'pause', pause):
            thread = Thread(target=read);thread.start();thread.join(10)
            self.assertFalse(thread.is_alive())
        if failures:raise failures[0]
        self.assertTrue(pauses)
        doc, views, last = results[0]
        self.assertEqual(doc.result.rows, reference.result.rows)
        self.assertEqual(doc.result.totals, reference.result.totals)
        self.assertEqual((doc.result.totals.income_cents, doc.result.totals.cost_cents,
                          doc.result.totals.profit_cents), (1000000, 400000, 600000))
        self.assertEqual(doc.result.issues, reference.result.issues)
        for role in views:
            self.assertEqual(views[role], baseline[role]);self.assertEqual(last[role], filter_view(baseline[role], '合成订单000999'))
            self.assertEqual(len(last[role]), 1);self.assertEqual(views[role].rows[last[role][0]][0], 1006)
        self.assertEqual(path.read_bytes(), raw)
        path.unlink()  # The detached model does not retain a source handle.

    def test_latency_trace_only_wraps_cooperative_pause_and_restores_it_after_error(self):
        from threading import Thread
        from quote_files import scheduling as s
        sleep, pause = time.sleep, s.ReadTurn.pause
        trace = LatencyTrace()
        with self.assertRaisesRegex(ValueError, 'trace failure'), patch.object(s, 'WINDOWS', True), trace:
            trace.start(time.perf_counter())
            def read():
                with s.read_turns():s._state.turn.pause()
            thread = Thread(target=read);thread.start();thread.join(3)
            self.assertFalse(thread.is_alive())
            time.sleep(.001)  # An unrelated owner sleep is not a read turn.
            raise ValueError('trace failure')
        self.assertIs(time.sleep, sleep);self.assertIs(s.ReadTurn.pause, pause)
        pauses = [e for e in trace.events if e['name'] == 'cooperative_pause']
        self.assertEqual(len(pauses), 1);self.assertEqual(pauses[0]['thread'], 'worker')

    def test_real_diagnostic_inputs_select_two_files_and_keep_source_gate_and_matching_bytes(self):
        cases = perf.make_inputs(self.root, ['small'], ['manager-rows-100', 'sales-rows-100'])
        self.assertEqual([c['id'] for c in cases], ['manager-rows-100', 'sales-rows-100'])
        self.assertFalse((self.root / 'manager-files-10').exists())
        self.assertEqual(cases[0]['input_sha256'], cases[1]['input_sha256'])
        for case in cases:
            checked, _ = perf.get_case(self.root, case['id'])
            self.assertTrue(checked['latency_diagnostic'])
            self.assertEqual(checked['files'], 1)

    def test_latency_diagnostic_rejects_shared_or_legacy_options_before_any_output_access(self):
        for flag in (['--shared-root', '\\\\synthetic\\share'], ['--compare-legacy']):
            with self.subTest(flag=flag), patch.object(perf.original, 'new_output', side_effect=AssertionError('output accessed')), \
                 patch('sys.stderr', new=io.StringIO()), self.assertRaises(SystemExit) as stopped:
                perf.main(['--diagnose-latency', '--output', 'ignored', *flag])
            self.assertEqual(stopped.exception.code, 2)

    def test_latency_trace_actual_decode_worker_gc_and_error_restore_all_hooks(self):
        import gc
        from threading import Thread
        from quote_files import documents, cache
        path = self.root / 'trace.xlsx';fixtures.workbook(path, 3)
        callbacks = list(gc.callbacks)
        originals = documents.parse_document, cache.parse_document
        results = []
        trace = LatencyTrace()
        with self.assertRaisesRegex(RuntimeError, 'synthetic trace error'):
            with instrumentation() as counts, trace:
                trace.start(time.perf_counter())
                worker = Thread(target=lambda: results.append(documents.parse_document(path.read_bytes())))
                worker.start();worker.join(5)
                self.assertFalse(worker.is_alive())
                gc.collect()
                raise RuntimeError('synthetic trace error')
        self.assertEqual(gc.callbacks, callbacks)
        self.assertEqual((documents.parse_document, cache.parse_document), originals)
        self.assertEqual(results[0].result.detail_count, 3)
        self.assertEqual(counts['workbook_opens'], 1)
        self.assertTrue(any(e['name'] == 'decode_workbook' and e['thread'] == 'worker' for e in trace.events))
        self.assertTrue(any(e['name'].startswith('gc.') and e['thread'] == 'owner' for e in trace.events))
        self.assertTrue(all(e['wall_seconds'] >= 0 and (e['thread_cpu_seconds'] is None or e['thread_cpu_seconds'] >= 0) for e in trace.events))

    def test_latency_overlap_is_aligned_bounded_and_never_a_causal_or_acceptance_claim(self):
        trace = LatencyTrace(max_events=2);trace.start(0)
        trace.record('render_page', .02, .10, .04)
        trace.record('gc.gen2', .04, .08)
        trace.record('discarded_span', .01, .02)
        data = trace.snapshot({'gaps_ms': [120, 50]})
        self.assertEqual(data['events_dropped'], 1)
        self.assertEqual(len(data['over_100_ms_gaps']), 1)
        self.assertAlmostEqual(data['over_100_ms_gaps'][0]['gc_overlap_ms'], 40)
        self.assertIn('render_page', data['over_100_ms_gaps'][0]['owner_spans'])
        self.assertIn('not proof of cause', data['scope'])

    def test_short_diagnostic_selects_only_two_large_files_and_never_accepts_one_round(self):
        def generated(root, sizes, case_ids):
            self.assertEqual(sizes, ['large'])
            self.assertEqual(case_ids, ['manager-rows-10000', 'sales-rows-10000'])
            cases = [{'id': name, 'files': 1, 'input_sha256': {}, 'latency_diagnostic': True} for name in case_ids]
            for c in cases:(root / c['id'] / 'inputs').mkdir(parents=True)
            (root / 'p04-run.json').write_text(json.dumps({'cases': cases}), encoding='utf-8')
            return cases
        def measured(root, case, version, native, timeout):
            self.assertTrue(native);self.assertEqual(version, 'current')
            hb = {'samples': 21, 'gaps_ms': [50] * 21, 'p95_gap_ms': 50, 'max_gap_ms': 50, 'status': 'passed'}
            return {'status': 'measured', 'case_id': case['id'], 'version': version, 'inputs_unchanged': True,
                'measurements': [{'phase': phase, 'seconds': .1, 'python_allocation_peak_bytes': 1,
                    'heartbeat': hb, 'counts': {'workbook_opens': 1 if phase == 'first_refresh' else 0},
                    'observed': {'correct': True}} for phase in ('first_refresh', 'warm_refresh', 'open_view', 'all_pages', 'filter_last_detail')]}
        with patch.object(perf, 'run_checks', return_value={'status': 'passed', 'tests_run': 0}), \
             patch.object(perf.original, 'new_output', return_value=self.root), \
             patch.object(perf.platform, 'system', return_value='Windows'), \
             patch.object(perf, 'make_inputs', side_effect=generated), \
             patch.object(perf, 'launch_worker', side_effect=measured) as workers, patch('sys.stdout', new=io.StringIO()):
            code = perf.main(['--diagnose-latency', '--output', 'ignored'])
        report = json.loads((self.root / 'report.json').read_text(encoding='utf-8'))
        self.assertEqual(code, 0);self.assertEqual(workers.call_count, 2)
        self.assertTrue(report['latency_diagnostic_only'])
        self.assertEqual(report['status'], 'latency_diagnostic_completed_acceptance_pending')
        self.assertFalse(report['storages'][0]['summary']['full_scale_five_rounds'])
        self.assertFalse(report['storages'][0]['summary']['Windows_healthy_performance_passed'])

    def test_latency_diagnostic_blocks_non_windows_before_input_or_gui_access(self):
        with patch.object(perf, 'run_checks', return_value={'status': 'passed', 'tests_run': 0}), \
             patch.object(perf.original, 'new_output', return_value=self.root), \
             patch.object(perf.platform, 'system', return_value='Linux'), \
             patch.object(perf, 'make_inputs', side_effect=AssertionError('inputs accessed')), patch('sys.stdout', new=io.StringIO()):
            self.assertEqual(perf.main(['--diagnose-latency', '--output', 'ignored']), 2)
        report = json.loads((self.root / 'report.json').read_text(encoding='utf-8'))
        self.assertEqual(report['status'], 'blocked');self.assertEqual(report['storages'], [])

    def test_normal_native_performance_still_rejects_one_round(self):
        with patch.object(perf, 'run_checks', side_effect=AssertionError('checks ran')), \
             patch('sys.stderr', new=io.StringIO()), self.assertRaises(SystemExit) as stopped:
            perf.main(['--gui-heartbeat', '--rounds', '1', '--output', 'ignored'])
        self.assertEqual(stopped.exception.code, 2)
    def test_paging_retains_user_column_width_and_reconfigures_changed_headings(self):
        from dataclasses import replace
        from test_application_integration import application
        from test_background_io import asynchronous, wait_for
        from test_pagination_images import page_controls
        with application('sales') as (app, folder, _, _), asynchronous(app):
            page_controls(app)
            widths = {}
            app.excel_view_tree.column = lambda name, **options: widths.update({name: options['width']})
            path = folder / 'reports' / 'synthetic.xlsx';fixtures.workbook(path, 401)
            task = app.open_excel_path(str(path))
            wait_for(lambda: task.number not in app._io_callbacks, app._poll_io)
            widths['c0'] = 321
            app.change_excel_page(1)
            self.assertEqual(widths['c0'], 321)
            self.assertEqual(len(app.excel_view_row_map), 200)
            app._page_view = replace(app._page_view, headings=('换名', *app._page_view.headings[1:]))
            app._render_page()
            self.assertEqual(widths['c0'], 140)

    def test_windows_thread_policy_preserves_shorter_slice_and_other_platform_settings(self):
        from quote_app.background import configure_ui_thread_scheduling
        with patch.object(sys, 'platform', 'win32'), patch.object(sys, 'getswitchinterval', return_value=.005), patch.object(sys, 'setswitchinterval') as setter:
            configure_ui_thread_scheduling();setter.assert_called_once_with(.001)
        with patch.object(sys, 'platform', 'win32'), patch.object(sys, 'getswitchinterval', return_value=.0005), patch.object(sys, 'setswitchinterval') as setter:
            configure_ui_thread_scheduling();setter.assert_called_once_with(.0005)
        with patch.object(sys, 'platform', 'linux'), patch.object(sys, 'setswitchinterval') as setter:
            configure_ui_thread_scheduling();setter.assert_not_called()

    def test_bulk_tcl_page_insert_keeps_scriptlike_unicode_cells_as_data(self):
        import tkinter as tk
        from quote_app.pagination import insert_page_items
        interpreter = tk.Tcl()
        interpreter.eval('set captured {}; set injected 0; proc page_tree {command parent where id_flag iid values_flag values} {global captured; lappend captured [list $iid $values]}')
        tree = SimpleNamespace(tk=interpreter.tk, _w='page_tree')
        values = ('中文换行\n第二行', '[set injected 1]; $captured', 'a {b} "c"', '000123')
        insert_page_items(tree, [('row_7', values), ('row_8', ('', '尾行'))])
        captured = interpreter.tk.splitlist(interpreter.getvar('captured'))
        self.assertEqual(interpreter.tk.splitlist(captured[0]), ('row_7', values))
        self.assertEqual(interpreter.tk.splitlist(captured[1]), ('row_8', ('', '尾行')))
        self.assertEqual(interpreter.getvar('injected'), '0')

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='777-p04-measurement-test-')
        self.root = Path(self.temp.name)
        codec = patch('io.text_encoding', side_effect=lambda encoding, stacklevel=2: 'gbk' if encoding in (None, 'locale') else encoding)
        codec.start();self.addCleanup(codec.stop)
    def tearDown(self):self.temp.cleanup()

    def test_p95_cannot_hide_single_563ms_maximum(self):
        clock = 0.;heart = Heartbeat();heart.start(clock)
        for gap in [50.] * 39 + [563.]:
            clock += gap / 1000;heart.pulse(clock)
        data = heart.finish(clock)
        self.assertLess(data['p95_gap_ms'], 100)
        self.assertAlmostEqual(data['max_gap_ms'], 563.)
        self.assertEqual(data['over_500_ms'], 1)
        self.assertEqual(data['status'], 'failed')
        self.assertFalse(healthy_heartbeat(data))

    def test_phase_boundary_includes_work_before_first_callback_and_rejects_short_sampling(self):
        heart = Heartbeat();heart.start(0.)
        data = heart.finish(.8)
        self.assertEqual(data['gaps_ms'], [800.])
        self.assertEqual(data['status'], 'insufficient_samples')
        self.assertFalse(healthy_heartbeat(data))

    def test_stage_reset_excludes_preparation_but_preserves_every_stage_gap(self):
        heart = Heartbeat();heart.pulse(100.)
        heart.start(101.)
        for i in range(1, 22):heart.pulse(101. + i * .05)
        first = heart.finish(102.1)
        self.assertTrue(healthy_heartbeat(first))
        heart.pulse(200.)  # no active phase: preprocessing is not UI work
        heart.start(201.)
        data = heart.finish(201.7)
        self.assertAlmostEqual(data['max_gap_ms'], 700.)
        self.assertGreater(len(first['gaps_ms']), 20)

    def test_summary_recomputes_raw_gaps_instead_of_trusting_passed_flag(self):
        data = {'samples': 21, 'gaps_ms': [50.] * 20 + [563.], 'p95_gap_ms': 50., 'max_gap_ms': 50., 'status': 'passed'}
        self.assertFalse(healthy_heartbeat(data))
        data['max_gap_ms'] = 563.
        self.assertFalse(healthy_heartbeat(data))
        self.assertEqual(distribution([]), {'samples': 0, 'p50': None, 'p95': None, 'max': None})

    def test_negative_and_nonfinite_clock_samples_are_errors(self):
        for now in (-1., float('nan'), float('inf')):
            heart = Heartbeat();heart.start(0.)
            with self.assertRaises(ValueError):heart.pulse(now)

    def test_real_current_decode_is_one_open_one_xml_and_one_source_iteration(self):
        path = self.root / '合成.xlsx';fixtures.workbook(path, 3)
        with instrumentation() as counts:doc = parse_document(path.read_bytes())
        self.assertEqual(doc.result.detail_count, 3)
        self.assertEqual(counts['workbook_opens'], 1)
        self.assertEqual(counts['worksheet_xml_source_opens'], 1)
        self.assertEqual(counts['iter_rows_calls'], 1)
        self.assertEqual(counts['yielded_rows'], 9)

    def test_readonly_cell_reopens_are_observed_and_all_hooks_restore_after_error(self):
        path = self.root / '合成.xlsx';fixtures.workbook(path, 3)
        before = (openpyxl.load_workbook, Worksheet.iter_rows, ReadOnlyWorksheet.iter_rows, ZipFile.open)
        with self.assertRaisesRegex(RuntimeError, 'diagnostic'):
            with instrumentation() as counts:
                book = openpyxl.load_workbook(path, read_only=True)
                try:
                    for row in range(7, 10):book.active.cell(row, 5).value
                    self.assertEqual(counts['workbook_opens'], 1)
                    self.assertGreaterEqual(counts['worksheet_xml_source_opens'], 3)
                    raise RuntimeError('diagnostic')
                finally:book.close()
        self.assertEqual(before, (openpyxl.load_workbook, Worksheet.iter_rows, ReadOnlyWorksheet.iter_rows, ZipFile.open))

    def test_fixture_bytes_are_identical_for_manager_sales_and_rounds_are_not_regenerated(self):
        cases = perf.make_inputs(self.root, ['small'])
        manager = next(c for c in cases if c['id'] == 'manager-rows-100')
        sales = next(c for c in cases if c['id'] == 'sales-rows-100')
        self.assertEqual(manager['input_sha256'], sales['input_sha256'])
        before = input_hashes(self.root / manager['id'] / 'inputs')
        for _ in range(2):
            case, folder = perf.get_case(self.root, manager['id'])
            self.assertEqual(case, manager)
            self.assertEqual(input_hashes(folder / 'inputs'), before)

    def test_changed_synthetic_inputs_refuse_measurement(self):
        cases = perf.make_inputs(self.root, ['small']);case = cases[0]
        path = next((self.root / case['id'] / 'inputs').glob('*.xlsx'))
        path.write_bytes(b'changed')
        with self.assertRaisesRegex(ValueError, 'SYNTHETIC_INPUT_CHANGED'):perf.get_case(self.root, case['id'])

    def test_actual_current_small_read_counts_and_exact_amounts_with_fresh_round_reset(self):
        cases = perf.make_inputs(self.root, ['small'])
        for case in (cases[0], cases[2]):
            result = perf.worker(self.root, case['id'], 'current')
            cold, warm = result['measurements']
            self.assertEqual(cold['counts']['workbook_opens'], case['files'])
            self.assertEqual(warm['counts'].get('workbook_opens', 0), 0)
            self.assertEqual(cold['observed']['income_cents'], case['income_cents'])
            self.assertEqual(cold['observed']['cost_cents'], case['cost_cents'])
            self.assertTrue(result['inputs_unchanged'])
            ns = perf.prepare_current(self.root / case['id'], case['role'])
            shared = json.loads((self.root / case['id'] / 'storage.json').read_text(encoding='utf-8'))
            self.assertEqual(shared['export_dir'], 'reports')

    def test_actual_original_matched_small_preserves_known_cost_defect_and_five_n_reads(self):
        case = perf.make_inputs(self.root, ['small'])[0]
        result = perf.worker(self.root, case['id'], 'original')
        for row in result['measurements']:
            self.assertEqual(row['counts']['workbook_opens'], 5 * case['files'])
            self.assertFalse(row['observed']['correct'])
            self.assertIn('known_baseline_issue', row['observed'])
        self.assertTrue(result['inputs_unchanged'])

    def test_non_windows_gui_block_precedes_all_input_and_widget_access(self):
        with patch.object(perf.platform, 'system', return_value='Linux'), patch.object(perf, 'get_case', side_effect=AssertionError('input accessed')):
            with self.assertRaisesRegex(ValueError, 'WINDOWS_GUI_ONLY'):perf.worker(self.root, 'ignored', 'current', native=True)
        with patch.object(perf.platform, 'system', return_value='Linux'):
            with self.assertRaisesRegex(ValueError, 'EXPLICIT_WINDOWS_UNC'):perf.shared_parent('\\\\synthetic\\share')

    def test_timeout_is_unavailable_sample_and_cannot_be_matched_baseline(self):
        case = perf.make_inputs(self.root, ['small'])[0]
        with patch.object(perf.subprocess, 'run', side_effect=perf.subprocess.TimeoutExpired('worker', 5)):
            result = perf.launch_worker(self.root, case, 'original', False, 5)
        self.assertEqual(result['status'], 'timeout')
        self.assertNotIn('measurements', result)
        self.assertFalse(perf.legacy_complete([result], [case], 5, ['small']))

    def test_missing_native_phases_or_headless_samples_cannot_pass_windows_performance(self):
        case = perf.make_inputs(self.root, ['small'])[0]
        result = perf.worker(self.root, case['id'], 'current')
        runs = []
        for _ in range(5):
            for index in range(9):runs.append({**deepcopy(result), 'case_id': str(index)})
        cases = [{**case, 'id': str(index)} for index in range(9)]
        auxiliary = perf.summarize(runs, cases, 5, False)
        self.assertTrue(auxiliary['current_correctness_and_read_counts_passed'])
        self.assertFalse(auxiliary['Windows_healthy_performance_passed'])
        native = perf.summarize(runs, cases, 5, True)
        self.assertFalse(native['Windows_healthy_performance_passed'])
        self.assertFalse(native['current_correctness_and_read_counts_passed'])

    def test_actual_current_and_archived_notebooks_expose_existing_excel_page_with_padded_labels(self):
        manifest = json.loads((ROOT / 'docs/baseline/manifest.json').read_text(encoding='utf-8'))
        for entry in manifest['files']:
            for path in (ROOT / entry['path'], perf.baseline_check.source_path(ROOT, entry)):
                with self.subTest(role=entry['role'], path=path.name):
                    app = SimpleNamespace();notebook = configured_notebook(app, path.read_bytes())
                    self.assertEqual(notebook.tab(app.excel_page, 'text'), ' Excel记录 ')
                    self.assertEqual(len(notebook.tabs()), 5 if entry['role'] == 'manager' else 3)

    def test_requested_full_scale_without_completed_rounds_is_not_measured_full_scale(self):
        cases = [{'id': str(index), 'files': 1} for index in range(9)]
        runs = [{'case_id': c['id'], 'version': 'current', 'status': 'worker_failed', 'round': 1} for c in cases]
        summary = perf.summarize(runs, cases, 5, True)
        self.assertFalse(summary['full_scale_five_rounds'])
        self.assertFalse(summary['Windows_healthy_performance_passed'])

    def test_worker_failure_preserves_entire_original_exception_traceback(self):
        stderr = 'Root failure: real notebook page missing\n' + 'context\n' * 1000 + 'RuntimeError: startup failed\n'
        child = SimpleNamespace(returncode=1, stdout='', stderr=stderr)
        with patch.object(perf.subprocess, 'run', return_value=child):
            result = perf.launch_worker(self.root, {'id': 'synthetic'}, 'current', True, 5)
        self.assertEqual(result['status'], 'worker_failed')
        self.assertEqual(result['stderr'], stderr)

    def test_sales_background_file_list_without_manager_amount_controls(self):
        from test_application_integration import application
        from test_background_io import asynchronous, wait_for
        with application('sales') as (app, folder, _, _), asynchronous(app):
            fixtures.workbook(folder / 'reports' / 'synthetic.xlsx', 3)
            for name in ('count', 'income', 'cost', 'profit', 'margin'):
                delattr(app, 'excel_total_' + name)
            task = app.refresh_excel_files()
            wait_for(lambda: task.number not in app._io_callbacks, app._poll_io)
            self.assertEqual(len(app.excel_file_tree.get_children()), 1, getattr(app, 'feedback', ''))
            self.assertIs(app._io_last_files, app.excel_read_results)
            self.assertEqual(next(iter(app.excel_file_tree.rows.values()))[2], 3)
            self.assertFalse(any(hasattr(app, 'excel_total_' + name) for name in ('income', 'cost', 'profit')))

    def test_observation_waits_for_actual_twenty_callbacks_under_windows_timer_rounding(self):
        heart = Heartbeat();heart.start(0)
        for index in range(1, 18):heart.pulse(index * .064)
        self.assertFalse(heart.ready(1.11))
        for index in range(18, 21):heart.pulse(index * .064)
        self.assertTrue(heart.ready(1.281))
        self.assertTrue(healthy_heartbeat(heart.finish(1.281)))

    def test_readonly_worker_timeout_preserves_phase_diagnostics(self):
        import subprocess
        stderr = b'P04_PHASE {"phase":"open_view","state":"started"}\n'
        failure = subprocess.TimeoutExpired('synthetic', 5, stderr=stderr)
        with patch.object(perf.subprocess, 'run', side_effect=failure):
            result = perf.launch_worker(self.root, {'id': 'synthetic'}, 'current', True, 5)
        self.assertEqual(result['status'], 'timeout')
        self.assertEqual(result['stderr'], stderr.decode('utf-8'))
        self.assertEqual(result['last_phase'], 'open_view')

    def test_gui_sequence_preflight_uses_real_background_adapters_and_all_phase_gates(self):
        # Tcl/replacement controls validate the diagnostic workflow only. No
        # Tk widgets, Windows heartbeat/DPI or native performance acceptance.
        from test_application_integration import application
        from test_background_io import asynchronous
        from test_pagination_images import page_controls
        from quote_app.background import BackgroundApplicationMixin
        from quote_files import scheduling
        for role, failing_delivery in (('manager', False), ('sales', False), ('sales', True)):
            with self.subTest(role=role, failing_delivery=failing_delivery), application(role) as (app, folder, ns, _), asynchronous(app):
                inputs = folder / 'inputs';inputs.mkdir();fixtures.workbook(inputs / 'synthetic.xlsx', 401)
                case = {'role': role, 'files': 1, 'detail_rows_per_file': 401, 'latency_diagnostic': True,
                        'income_cents': 401000, 'cost_cents': 160400, 'profit_cents': 240600}
                page_controls(app)
                app.excel_next_button.invoke = lambda a=app: a.change_excel_page(1)
                app.excel_filter_button = Mock();app.excel_filter_button.invoke = app.filter_excel_view
                if role == 'sales':
                    for name in ('count', 'income', 'cost', 'profit', 'margin'):
                        delattr(app, 'excel_total_' + name)
                notebook = configured_notebook(app, (ROOT / next(e['path'] for e in json.loads(
                    (ROOT / 'docs/baseline/manifest.json').read_text(encoding='utf-8'))['files'] if e['role'] == role)).read_bytes())
                alive = [True]
                window = app.root
                # Real entrypoints queue an initial refresh for 300ms; real
                # Windows 50ms timers commonly deliver around 64ms.
                real_after = window.after
                window.after = lambda ms, *args: real_after(64 if ms == 50 else ms, *args)
                app.schedule_quote_callback(app.refresh_excel_files, delay=300)
                if failing_delivery:
                    app._insert_excel_file = Mock(side_effect=RuntimeError('synthetic render failure'))
                original_tcl = window.tk
                class TclProxy:
                    def call(self, *args):
                        if args == ('tk', 'scaling'):return 1.0
                        return original_tcl.call(*args)
                    def __getattr__(self, name):return getattr(original_tcl, name)
                window.tk = TclProxy()
                def loop():
                    deadline = time.monotonic() + 15
                    while alive[0] and time.monotonic() < deadline:
                        window.tk.eval('update');time.sleep(.002)
                    if alive[0]:raise TimeoutError('auxiliary workflow did not finish')
                window.mainloop = loop;window.quit = lambda: alive.__setitem__(0, False)
                import tkinter as tk
                window.attributes = Mock();window.protocol = Mock()
                window.destroy = Mock(side_effect=[None, tk.TclError('root already destroyed')])
                # Exercise actual scheduled poll callbacks rather than the
                # test fixture's manual pump wrapper.
                app._poll_io = BackgroundApplicationMixin._poll_io.__get__(app)
                class Factory:
                    def __new__(cls, _):return app
                try:
                    with patch('tkinter.Tk', return_value=window), patch.object(perf.gui, 'live_sandbox_type', return_value=Factory), \
                         patch.object(scheduling, 'WINDOWS', True), patch('sys.stderr', new=io.StringIO()):
                        if failing_delivery:
                            with self.assertRaisesRegex(RuntimeError, 'first_refresh: RuntimeError: APPLICATION_READ_FAILED: .*synthetic render failure'):
                                perf.gui_measurements(ns, folder, case, True)
                        else:
                            measured, _ = perf.gui_measurements(ns, folder, case, True)
                    window.attributes.assert_called_once_with('-disabled', True)
                    self.assertEqual(window.destroy.call_count, 2)
                finally:
                    # The shared fixture pumps its closed runtime on teardown.
                    window.destroy.side_effect = None
                if failing_delivery:continue
                self.assertIs(notebook.selected, app.excel_page)
                self.assertEqual([m['phase'] for m in measured], ['first_refresh', 'warm_refresh', 'open_view', 'all_pages', 'filter_last_detail'])
                self.assertEqual(measured[0]['counts']['workbook_opens'], 1)
                self.assertEqual(measured[1]['counts'].get('workbook_opens', 0), 0)
                self.assertTrue(all(m['heartbeat']['samples'] >= 20 for m in measured))
                self.assertTrue(all('latency' in m for m in measured))
                self.assertTrue(any(e['name'] == 'cooperative_pause' and e['thread'] == 'worker'
                                    for m in measured for e in m['latency']['events']))
                self.assertFalse(any(e['name'] == 'cooperative_pause' and e['thread'] == 'owner'
                                     for m in measured for e in m['latency']['events']))
                self.assertTrue(any(e['thread'] == 'owner' and e['name'] == 'render_page'
                                    for m in measured for e in m['latency']['events']))
                self.assertEqual(measured[3]['observed']['visible_rows'], 407 if role == 'manager' else 401)
                self.assertEqual(measured[3]['observed']['pages'], 3)
                self.assertEqual(measured[4]['observed']['source_row'], 407)

    def test_cli_off_windows_reports_blocked_without_probing_shared_root(self):
        with patch.object(perf, 'run_checks', return_value={'status': 'passed', 'tests_run': 0}), \
             patch.object(perf.original, 'new_output', return_value=self.root), \
             patch.object(perf.platform, 'system', return_value='Linux'), \
             patch.object(perf, 'shared_parent', side_effect=AssertionError('share accessed')), \
             patch('sys.stdout', new=io.StringIO()):
            code = perf.main(['--gui-heartbeat', '--shared-root', '\\\\synthetic\\share', '--output', 'ignored'])
        report = json.loads((self.root / 'report.json').read_text(encoding='utf-8'))
        self.assertEqual(code, 2)
        self.assertEqual(report['status'], 'blocked')
        self.assertFalse(report['production_paths_accessed'])
        self.assertEqual(report['storages'], [])

    def test_preparation_failure_emits_failed_report_instead_of_successful_checks_snapshot(self):
        with patch.object(perf, 'run_checks', return_value={'status': 'passed', 'tests_run': 0}), \
             patch.object(perf.original, 'new_output', return_value=self.root), \
             patch.object(perf, 'make_inputs', side_effect=PermissionError('synthetic preparation refused')), \
             patch('sys.stdout', new=io.StringIO()):
            code = perf.main(['--auxiliary', '--output', 'ignored'])
        report = json.loads((self.root / 'report.json').read_text(encoding='utf-8'))
        self.assertEqual(code, 2)
        self.assertEqual(report['status'], 'failed')
        self.assertIn('PermissionError', report['error'])
        self.assertEqual(report['storages'], [])


if __name__ == '__main__':unittest.main()
