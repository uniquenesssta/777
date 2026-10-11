"""Current manager/sales methods, real Tcl variable traces, synthetic files.

These owner-thread adapter checks do not replace Windows widget/screen checks.
"""
import ast
from contextlib import contextmanager
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import tkinter as tk
import unittest
from unittest.mock import Mock, patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))
import baseline_check
import p02_gui_check as gui
from quote_app import QuoteApplicationMixin
from quote_app.integration import config_mapping


class Tree:
    def __init__(self):
        self.rows = []
    def get_children(self):
        return list(range(len(self.rows)))
    def delete(self, _):
        self.rows.pop()
    def insert(self, *_, values):
        self.rows.append(values)


@contextmanager
def application(role):
    with tempfile.TemporaryDirectory(prefix="777-p02-ui-adapter-") as folder:
        sandbox = Path(folder)
        ns, entry = gui.load_namespace(ROOT, role, sandbox)
        current = ns["PriceCalculator"]
        class Harness(current):
            def add_setting_row(self, **item):
                row = {k: tk.BooleanVar(self.root, v) if k == "single" else tk.StringVar(self.root, v)
                       for k, v in item.items()}
                row["frame"] = None
                self.setting_rows.append(row)
                self.watch_quote_row(row, settings=True)
            def render_issues(self, issues):
                self.rendered_issues = issues
            def refresh_item_selector(self):
                pass
            def update_calc_row_states(self, config):
                pass
            def refresh_excel_files(self):
                pass
            def notify_validation_errors(self, *args):
                pass
            def set_feedback(self, text, is_error=False):
                self.feedback = text
        app = Harness.__new__(Harness)
        app.root = tk.Tcl()
        app.default_items = []
        app.setting_rows, app.calc_rows = [], []
        app.config_file = str(sandbox / "config.json")
        app.report_dir = app.default_report_dir = str(sandbox / "reports")
        app.image_dir = str(sandbox / "screenshots")
        app.export_dir_var = tk.StringVar(app.root, app.report_dir)
        app.profit_tree = Tree()
        app.last_export_records, app.validation_issues, app.last_config_snapshot = [], [], {}
        for key in ("user_name", "order_no", "wechat_no", "total_val", "summary_income",
                    "summary_cost", "summary_profit", "summary_margin", "quote_status_var"):
            setattr(app, key, tk.StringVar(app.root, ""))
        app.initialize_quote_state()
        for var in (app.user_name, app.order_no, app.wechat_no):
            var.trace_add("write", app.quote_input_changed)
        app.reload_price_settings_from_json(first_load=True, silent=True)
        def add(kind="normal", product="P02单件", q="1", h="1", w="1", percent="90%"):
            row = {"row_kind": kind, "base_type": product, "display_name": product,
                   "frame": Mock()}
            for key, value in (("q", q), ("h", h), ("w", w), ("discount_choice", percent),
                               ("discount_custom", "90"), ("discount", "—"), ("res", ""), ("status", "")):
                row[key] = tk.BooleanVar(app.root, value) if type(value) is bool else tk.StringVar(app.root, value)
            if kind == "discount":
                row["discount"] = row["discount_choice"]
            app.calc_rows.append(row)
            app.watch_quote_row(row)
            return row
        app.add_fixture_row = add
        with patch("tkinter.messagebox.showinfo"), patch("tkinter.messagebox.showerror"), patch("tkinter.messagebox.showwarning"):
            yield app, sandbox, ns, entry
        app.calc_rows.clear()
        app.setting_rows.clear()


class ApplicationIntegrationTests(unittest.TestCase):
    def test_current_entrypoints_share_core_and_original_archive_is_exact(self):
        manifest = json.loads(baseline_check.MANIFEST.read_text(encoding="utf-8"))
        for e in manifest["files"]:
            archive = baseline_check.source_path(ROOT, e)
            self.assertEqual(hashlib.sha256(archive.read_bytes().replace(b"\r\n", b"\n")).hexdigest(), e["sha256"])
            tree = ast.parse((ROOT / e["path"]).read_text(encoding="utf-8"))
            classes = {n.name: n for n in tree.body if isinstance(n, ast.ClassDef)}
            self.assertNotIn("SafeNumberParser", classes)
            self.assertNotIn("DiscountParser", classes)
            names = {n.name for n in classes["PriceCalculator"].body if isinstance(n, ast.FunctionDef)}
            for name in ("calculate", "export_to_excel", "save_screenshot", "get_settings_snapshot"):
                self.assertNotIn(name, names)
            with application(e["role"]) as (app, *_):
                self.assertIs(app.calculate.__func__, QuoteApplicationMixin.calculate)

    def test_both_live_adapters_match_all_shared_golden_amounts_or_whole_rejection(self):
        suite = json.loads((ROOT / "docs/baseline/p01-04-golden.json").read_text(encoding="utf-8"))
        for role in ("manager", "sales"):
            with application(role) as (app, folder, *_):
                for case in suite["cases"]:
                    with self.subTest(role=role, case=case["id"]):
                        app.calc_rows.clear()
                        app.setting_rows.clear()
                        app.initialize_quote_state()
                        data = {"items": [{"name": name, **values} for name, values in case["price_snapshot"]["items"].items()]}
                        (folder / "config.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
                        app.reload_price_settings_from_json(silent=True)
                        for raw in case["input"]:
                            app.add_fixture_row(kind=raw["kind"], product=raw.get("product_id", "折扣"),
                                q=raw.get("q", ""), h=raw.get("h", ""), w=raw.get("w", ""), percent=raw.get("percent", "90%"))
                        result = app.calculate(notify=False)
                        if case["expected"]["status"] == "ok":
                            self.assertEqual(result.display(), case["expected"])
                        else:
                            self.assertIsNone(result.totals)
                            self.assertIsNone(result.lines)
                            self.assertFalse(app.quote_session.prepare_output("excel").allowed)

    def test_real_variable_edits_and_identity_clear_previous_amounts_immediately(self):
        for role in ("manager", "sales"):
            with application(role) as (app, *_):
                row = app.add_fixture_row(product="P02面积")
                for var in (row["q"], row["h"], row["w"], app.user_name, app.order_no, app.wechat_no):
                    app.calculate(notify=False)
                    self.assertIsNotNone(app.quote_session.state.result)
                    var.set(var.get() + "0")
                    self.assertIsNone(app.quote_session.state.result)
                    self.assertEqual(app.summary_income.get(), "—")
                    self.assertEqual(app.last_export_records, [])
                    var.set("1")
                discount = app.add_fixture_row(kind="discount")
                app.calculate(notify=False)
                discount["discount_choice"].set("自定义")
                self.assertIsNone(app.quote_session.state.result)
                app.calculate(notify=False)
                discount["discount_custom"].set("89.123456789012")
                self.assertIsNone(app.quote_session.state.result)

    def test_invalid_or_stale_current_exports_and_screenshots_never_invoke_io(self):
        for role in ("manager", "sales"):
            with application(role) as (app, folder, *_):
                app.add_fixture_row()
                bad = app.add_fixture_row(q="bad")
                before = gui.files(folder)
                with patch("quote_app.integration.workbook_transaction", side_effect=AssertionError("opened")), patch("PIL.ImageGrab.grab", side_effect=AssertionError("captured")):
                    self.assertFalse(app.export_to_excel())
                    app.calculate(notify=False)
                    self.assertFalse(app.export_to_excel())
                    self.assertFalse(app._capture_current_quote())
                    self.assertIsNone(app.quote_session.state.result.totals)
                    self.assertEqual(app.quote_session.state.result.issues[0].line_index, 2)
                    bad["q"].set("1")
                    app.calculate(notify=False)
                    bad["q"].set("2")
                    self.assertFalse(app.export_to_excel())
                    self.assertFalse(app._capture_current_quote())
                self.assertEqual(gui.files(folder), before)

    def test_price_observation_does_not_silently_rebind_then_refresh_requires_confirmation(self):
        for role in ("manager", "sales"):
            with application(role) as (app, folder, *_):
                app.add_fixture_row()
                app.calculate(notify=False)
                old = app.quote_session.state.prices
                data = json.loads((folder / "config.json").read_text(encoding="utf-8"))
                data["items"][0]["sale"] = "20"
                (folder / "config.json").write_text(json.dumps(data), encoding="utf-8")
                self.assertFalse(app.export_to_excel())
                self.assertIs(app.quote_session.state.prices, old)
                self.assertEqual(app.quote_session.prepare_output("excel").status, "PRICE_STALE")
                app.calculate(notify=False)
                self.assertEqual(app.quote_session.state.result.totals.income_cents, 1)
                self.assertFalse(app.quote_session.prepare_output("excel").allowed)
                app.reload_price_settings_from_json(silent=True)
                self.assertTrue(app.price_changes)
                self.assertTrue(app.quote_session.state.price_confirmation_required)
                app.confirm_current_prices()
                self.assertFalse(app.quote_session.prepare_output("excel").allowed)
                app.calculate(notify=False)
                app.confirm_current_prices()
                self.assertTrue(app.quote_session.prepare_output("excel").allowed)
                self.assertEqual(app.quote_session.state.result.totals.income_cents, 2000)

    def test_invalid_or_unavailable_file_retains_snapshot_and_blocks_preview_outputs(self):
        for role in ("manager", "sales"):
            with application(role) as (app, folder, *_):
                app.add_fixture_row()
                active = app.active_price_config
                for content in ('{"items":[{"single":"false"}]}', None):
                    path = folder / "config.json"
                    if content is None:
                        path.unlink()
                    else:
                        path.write_text(content, encoding="utf-8")
                    before = gui.files(folder)
                    self.assertFalse(app.reload_price_settings_from_json())
                    self.assertIs(app.active_price_config, active)
                    result = app.calculate(notify=False)
                    self.assertEqual(result.totals.income_cents, 1)
                    self.assertFalse(app.export_to_excel())
                    self.assertFalse(app._capture_current_quote())
                    self.assertEqual(gui.files(folder), before)

    def test_first_bad_configuration_has_no_defaults_and_close_has_no_save(self):
        for role in ("manager", "sales"):
            with application(role) as (app, folder, *_):
                app.active_price_config = None
                app.initialize_quote_state()
                app.setting_rows.clear()
                app.default_items = [{"name": "不得兜底", "sale": "100", "cost": "1"}]
                path = folder / "config.json"
                path.write_text('{"items":[false]}', encoding="utf-8")
                before = gui.files(folder)
                self.assertFalse(app.load_settings_to_ui())
                self.assertIsNone(app.active_price_config)
                self.assertEqual(app.setting_rows, [])
                self.assertFalse(app.export_to_excel())
                self.assertEqual(app.quote_session.state.price_source, "invalid")
                with patch.object(app, "save_settings", side_effect=AssertionError("autosaved")), patch.object(app.root, "destroy"):
                    app.on_close()
                self.assertEqual(gui.files(folder), before)

    def test_blank_unused_editor_price_allows_calculation_but_blocks_whole_config_save(self):
        with application("manager") as (app, folder, *_):
            app.add_fixture_row()
            app.add_setting_row(name="未使用空价格", sale="", cost="", k="", single=False)
            result = app.calculate(notify=False)
            self.assertEqual(result.totals.income_cents, 1)
            before = gui.files(folder)
            self.assertFalse(app.save_settings())
            self.assertFalse(app.export_to_excel())
            self.assertEqual(gui.files(folder), before)

    def test_valid_editor_save_applies_whole_snapshot_and_requires_recompute_confirm(self):
        with application("manager") as (app, folder, *_):
            row = app.add_fixture_row()
            app.calculate(notify=False)
            app.setting_rows[0]["sale"].set("30")
            self.assertIsNone(app.quote_session.state.result)
            self.assertTrue(app.save_and_apply_settings())
            self.assertEqual(app.active_price_config.prices.items[0].sale, "30")
            self.assertFalse(app.export_to_excel())
            app.calculate(notify=False)
            app.confirm_current_prices()
            self.assertTrue(app.quote_session.prepare_output("excel").allowed)

    def test_restoring_unsaved_editor_prices_still_requires_new_calculation(self):
        with application("manager") as (app, *_):
            app.add_fixture_row()
            app.setting_rows[0]["sale"].set("30")
            app.calculate(notify=False)
            self.assertFalse(app.quote_session.prepare_output("excel").allowed)
            app.setting_rows[0]["sale"].set("0.005")
            app.calculate(notify=False)
            app.confirm_current_prices()
            self.assertTrue(app.quote_session.prepare_output("excel").allowed)

    def test_explicit_same_file_refresh_invalidates_and_requires_new_confirmation(self):
        for role in ("manager", "sales"):
            with application(role) as (app, *_):
                app.add_fixture_row()
                app.calculate(notify=False)
                old = app.quote_session.state.result
                app.reload_price_settings_from_json()
                self.assertIsNone(app.quote_session.state.result)
                self.assertTrue(app.quote_session.state.price_confirmation_required)
                self.assertFalse(app.quote_session.prepare_output("screenshot").allowed)
                app.calculate(notify=False)
                app.confirm_current_prices()
                self.assertTrue(app.quote_session.prepare_output("screenshot").allowed)

    def test_valid_excel_uses_frozen_result_not_tampered_widgets_or_cached_records(self):
        from openpyxl import load_workbook
        for role in ("manager", "sales"):
            with application(role) as (app, folder, *_):
                app.user_name.set("000123")
                app.order_no.set("0000456")
                app.add_fixture_row()
                app.add_fixture_row()
                app.calculate(notify=False)
                app.summary_income.set("9999")
                app.summary_cost.set("9999")
                app.last_export_records.clear()
                path = app.export_to_excel()
                self.assertTrue(path)
                wb = load_workbook(path)
                try:
                    rows = list(wb["利润明细"].iter_rows(min_row=7, values_only=True))
                    self.assertEqual(wb["利润明细"].cell(4, 2).value, "0.02")
                    self.assertEqual(set(wb.sheetnames), {"利润明细", "输入问题", "价格设定", "_写入事务"})
                    self.assertEqual(wb["_写入事务"].sheet_state, "hidden")
                    self.assertEqual([(r[1], r[2], r[16], r[17], r[18]) for r in rows],
                        [("000123", "0000456", "0.01", "0.01", "0.00"),
                         ("000123", "0000456", "0.01", "0.01", "0.00"),
                         ("000123", "0000456", "0.02", "0.02", "0.00")])
                finally:
                    wb.close()

    def test_explicit_export_line_kind_keeps_goods_named_like_summary_or_remark(self):
        from openpyxl import load_workbook
        for role in ("manager", "sales"):
            with application(role) as (app, folder, *_):
                row = app.add_fixture_row(product="P02面积")
                row["display_name"] = "--- 正常商品"
                second = app.add_fixture_row(product="P02面积")
                second["display_name"] = "备注"
                app.calculate(notify=False)
                first = app.export_to_excel()
                self.assertTrue(first)
                app.calculate(notify=False)
                self.assertTrue(app.export_to_excel())
                wb = load_workbook(first)
                try:
                    sheet = wb["利润明细"]
                    self.assertEqual(sheet.cell(4, 2).value, "48.00")
                    self.assertEqual(sheet.cell(6, 27).value, "行类型")
                    self.assertEqual([sheet.cell(r, 27).value for r in range(7, 13)],
                                     ["normal", "normal", "order_total"] * 2)
                finally:
                    wb.close()

    def test_editor_duplicates_and_reserved_names_are_errors_not_silent_skips(self):
        with application("manager") as (app, folder, *_):
            app.add_fixture_row()
            for name in ("P02单件", "折扣"):
                app.add_setting_row(name=name, sale="1", cost="1", k="1", single=True)
                before = gui.files(folder)
                self.assertIsNone(app.calculate(notify=False).totals)
                self.assertFalse(app.save_settings())
                self.assertEqual(gui.files(folder), before)
                app.setting_rows.pop()

    def test_discount_focus_loss_preserves_full_original_precision(self):
        for role in ("manager", "sales"):
            with application(role) as (app, *_):
                row = app.add_fixture_row(kind="discount", percent="自定义")
                row["discount_custom"].set("89.123456789012")
                app.normalize_discount_entry(row["discount_custom"])
                self.assertEqual(row["discount_custom"].get(), "89.123456789012")

    def test_discount_result_render_does_not_mutate_aliased_input_or_invalidate_itself(self):
        for role in ("manager", "sales"):
            with application(role) as (app, *_):
                app.add_fixture_row(product="P02面积")
                row = app.add_fixture_row(kind="discount", percent="自定义")
                row["discount_custom"].set("89.123456789012")
                result = app.calculate(notify=False)
                self.assertIs(app.quote_session.state.result, result)
                self.assertEqual(row["discount_choice"].get(), "自定义")
                self.assertEqual(row["discount_custom"].get(), "89.123456789012")
                self.assertTrue(app.quote_session.prepare_output("excel").allowed)
                row["discount_choice"].set("95%")
                result = app.calculate(notify=False)
                self.assertIs(app.quote_session.state.result, result)
                self.assertEqual(row["discount_choice"].get(), "95%")

    def test_pending_screenshot_callback_rechecks_after_intervening_edit(self):
        for role in ("manager", "sales"):
            with application(role) as (app, folder, *_):
                row = app.add_fixture_row()
                app.calculate(notify=False)
                app.notebook = Mock()
                app.calc_page = Mock()
                before = gui.files(folder)
                with patch.object(app, "_write_quote_screenshot", side_effect=AssertionError("captured")):
                    app.save_screenshot()
                    row["q"].set("2")
                    app.root.eval("update idletasks")
                self.assertEqual(gui.files(folder), before)

    def test_excel_refresh_failure_is_reported_as_saved_not_write_failure(self):
        with application("manager") as (app, *_):
            app.add_fixture_row()
            app.calculate(notify=False)
            with patch.object(app, "refresh_excel_files", side_effect=OSError("view failed")):
                path = app.export_to_excel()
            self.assertTrue(Path(path).is_file())
            self.assertIn("已保存", app.feedback)

    def test_price_change_dialog_for_sales_never_displays_cost_values(self):
        with application("sales") as (app, *_):
            app.price_changes = (("商品", {"sale": "10", "cost": "secret-old", "k": "1", "single": True},
                                 {"sale": "11", "cost": "secret-new", "k": "1", "single": True}),)
            with patch("tkinter.messagebox.showinfo") as dialog:
                app.show_price_changes()
            text = str(dialog.call_args)
            self.assertNotIn("secret", text)
            self.assertIn("10 → 11", text)

    def test_row_deletion_reset_and_setting_removal_invalidate_without_auto_calculate(self):
        for role in ("manager", "sales"):
            with application(role) as (app, *_):
                row = app.add_fixture_row()
                app.add_fixture_row()
                app.calculate(notify=False)
                app.remove_calc_row(row["frame"])
                self.assertIsNone(app.quote_session.state.result)
                self.assertEqual(len(app.calc_rows), 1)
                app.calculate(notify=False)
                app.reset()
                self.assertIsNone(app.quote_session.state.result)
                self.assertEqual(app.calc_rows, [])
                app.add_fixture_row()
                app.calculate(notify=False)
                app.setting_rows[0]["frame"] = Mock()
                app.remove_setting_row(app.setting_rows[0]["frame"])
                self.assertIsNone(app.quote_session.state.result)
                self.assertIsNone(app.calculate(notify=False).totals)

    def test_valid_screenshot_writer_and_blocked_callback_use_same_current_gate(self):
        from PIL import Image
        for role in ("manager", "sales"):
            with application(role) as (app, folder, *_):
                app.add_fixture_row()
                app.calculate(notify=False)
                with patch.object(app.root, "winfo_rootx", return_value=0), patch.object(app.root, "winfo_rooty", return_value=0), patch.object(app.root, "winfo_width", return_value=1), patch.object(app.root, "winfo_height", return_value=1), patch("PIL.ImageGrab.grab", return_value=Image.new("RGB", (1, 1))) as grab:
                    path = app._capture_current_quote()
                    self.assertTrue(Path(path).is_relative_to(folder))
                    self.assertTrue(Path(path).is_file())
                    self.assertEqual(grab.call_count, 1)
                    app.order_no.set("changed")
                    before = gui.files(folder)
                    self.assertFalse(app._capture_current_quote())
                    self.assertEqual(grab.call_count, 1)
                    self.assertEqual(gui.files(folder), before)

    def test_archived_source_tampering_is_rejected_before_execution(self):
        manifest = json.loads(baseline_check.MANIFEST.read_text(encoding="utf-8"))
        e = manifest["files"][0]
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            archive = root / "docs/baseline/original/manager.py"
            archive.parent.mkdir(parents=True)
            marker = root / "executed"
            data = baseline_check.source_path(ROOT, e).read_bytes()
            archive.write_bytes(data + f"\nfrom pathlib import Path\nPath({str(marker)!r}).touch()\n".encode("utf-8"))
            (root / "docs/baseline/source-location.json").write_text(json.dumps({"paths": {e["path"]: str(archive.relative_to(root))}}), encoding="utf-8")
            self.assertEqual(baseline_check.inspect_source(root, e)["status"], "changed")
            with self.assertRaises(ValueError):
                baseline_check.sandbox_namespace(root, e, root / "sandbox")
            self.assertFalse(marker.exists())

    def test_gui_cli_blocks_non_windows_before_loading_current_entrypoints(self):
        with tempfile.TemporaryDirectory() as folder, patch.object(gui.platform, "system", return_value="Linux"), patch.object(gui, "run_role", side_effect=AssertionError("GUI invoked")), patch.object(gui.baseline_check, "emit"):
            self.assertEqual(gui.main(["--gui-smoke", "--output", str(Path(folder) / "report.json")]), 2)


if __name__ == "__main__":
    unittest.main()
