"""Static transfer fixtures for Rust; not an additional business calculator."""
import io
import json
from pathlib import Path
import sys
import unittest
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from quote_service.output_render import FORMAT,KINDS,render,public_order
from quote_service.protocol import APIError
from quote_core.numbers import money_text


class ProtocolTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.fixture=json.loads((Path(__file__).resolve().parents[1]/'docs/verification/P05-04-protocol.json').read_text(encoding='utf-8'))
    def payload(self,cents):
        order={'order_id':self.fixture['request']['order_id'],'revision':1,'state':'active','price_version_id':'historical-v1','rule_version':'business-v1','numeric_profile':'numeric-v1','created_at':'2026-01-01T00:00:00+00:00','details':{'customer':'=SUM(1,2)','order_ref':'000001','wechat':'@identifier'},'income_cents':cents,'cost_cents':'0','profit_cents':cents,'notes':[],'lines':[]}
        return {'format':FORMAT,'job_id':self.fixture['request']['request_id'],'kind':'excel_quote','order':public_order(order)}
    def test_literal_format_and_kind_extensions_match_static_contract(self):
        self.assertEqual(FORMAT,self.fixture['output_format']);self.assertEqual(KINDS,{k:v['extension'] for k,v in self.fixture['kinds'].items()})
    def test_precision_over_javascript_integer_range_remains_exact_excel_text(self):
        from openpyxl import load_workbook
        for sample in self.fixture['precision_samples']:
            self.assertEqual(money_text(int(sample['income_cents'])),sample['expected_yuan'])
            raw=render(self.payload(sample['income_cents']));w=load_workbook(io.BytesIO(raw))
            try:
                self.assertEqual(w.sheetnames,self.fixture['public_workbook_sheets']);values=dict(list(w['Order'].values)[1:])
                self.assertEqual(values['income_cents'],sample['income_cents']);self.assertEqual(values['income_yuan'],sample['expected_yuan'])
                self.assertEqual(w['Order']['B9'].data_type,'s')
            finally:w.close()
    def test_future_format_is_not_silently_rendered_as_v1(self):
        with self.assertRaises(APIError):render(self.payload('3600')|{'format':self.fixture['invalid_future_output_format']})
    def test_public_projection_discards_future_fields_and_admin_finances(self):
        value=self.payload('3600')['order']|{'cost_cents':'1','profit_cents':'3599','owner_id':'secret','future_private':'secret'}
        safe=public_order(value)
        self.assertFalse(set(safe)&{'cost_cents','profit_cents','owner_id','future_private'})
        from openpyxl import load_workbook
        payload=self.payload('3600');payload['order']['details']['customer']='literal\x01text'
        w=load_workbook(io.BytesIO(render(payload)))
        try:self.assertEqual(dict(list(w['Order'].values)[1:])['customer'],'literal\\u0001text')
        finally:w.close()
