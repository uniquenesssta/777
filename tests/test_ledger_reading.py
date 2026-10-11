"""P03-02 synthetic layouts and current adapters; no production or SMB paths."""
from contextlib import ExitStack
from decimal import Decimal
import json
from pathlib import Path
import tempfile
import tkinter as tk
import unittest
from unittest.mock import Mock, patch
from zipfile import ZipFile, ZIP_DEFLATED

import openpyxl
from test_application_integration import application
from test_workbook_transactions import export_one
import p01_fixtures as fixtures
from quote_app.file_reads import FileReadingMixin
from quote_files.reading import analyze_raw, read_ledger
from quote_files.schema import DETAIL_HEADERS, create_template
from quote_files.transactions import FileFailure
from quote_files.workbooks import open_book, table


def edit(path, change):
    book = openpyxl.load_workbook(path)
    try:
        change(book)
        book.save(path)
    finally:
        book.close()


def codes(result):
    return {issue.code for issue in result.issues}


class ViewTree:
    def __init__(self):
        self.rows = {}
    def get_children(self):
        return tuple(self.rows)
    def delete(self, *indices):
        for index in indices:
            self.rows.pop(index)
    def insert(self, *_, iid, values):
        self.rows[iid] = values
    def configure(self, **_):
        pass
    def heading(self, *_args, **_kwargs):
        pass
    def column(self, *_args, **_kwargs):
        pass


class LedgerReadingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="777-p03-format-")
        self.folder = Path(self.temp.name)
        self.path = self.folder / "ledger.xlsx"
    def tearDown(self):
        self.temp.cleanup()
    def fixture(self, count=2, variant="current", summary=True):
        fixtures.workbook(self.path, count, variant, summary)
        return self.path
    def totals(self, result):
        self.assertEqual(result.status, "ok", result.issues)
        return tuple(result[key] for key in ("income", "cost", "profit"))

    def test_all_immutable_p01_excel_formats_have_explicit_results_and_unchanged_bytes(self):
        directory = self.folder / "formats"
        entries = fixtures.formats(directory)
        for entry in entries:
            if entry["kind"] != "excel":
                continue
            with self.subTest(path=entry["path"]):
                path = directory / entry["path"]
                before = path.read_bytes()
                result = read_ledger(path)
                expected = ("unsupported" if "duplicate_headers" in path.name else
                            "error" if any(k in path.name for k in ("formula", "corrupt")) else "ok")
                self.assertEqual(result.status, expected, result.issues)
                self.assertEqual(path.read_bytes(), before)
                if expected == "ok":
                    amounts = (0, 0, 0) if "empty_evidence" in path.name else (20, 8, 12)
                    self.assertEqual(self.totals(result), amounts)
                else:
                    self.assertIsNone(result.totals)
                    with self.assertRaises(ValueError):
                        result["income"]

    def test_missing_summary_uses_named_financial_columns_and_never_double_counts_footer(self):
        self.fixture(count=1, variant="no_summary", summary=False)
        result = read_ledger(self.path)
        self.assertEqual(self.totals(result), (10, 4, 6))
        self.assertEqual(result.summary_source, "details")
        self.assertEqual(result.detail_count, 1)
        self.assertIn("UNBOUND_LEGACY_TOTAL", codes(result))

    def test_zero_summary_is_present_and_nonzero_details_are_rejected(self):
        self.fixture(count=0)
        result = read_ledger(self.path)
        self.assertEqual(self.totals(result), (0, 0, 0))
        self.assertEqual(result.summary_source, "header")
        self.fixture(count=1)
        edit(self.path, lambda b: [setattr(b.active.cell(4, c), "value", 0) for c in (2, 4, 6)])
        result = read_ledger(self.path)
        self.assertEqual(result.status, "error")
        self.assertIn("SUMMARY_MISMATCH", codes(result))
        self.assertIsNone(result.totals)

    def test_missing_fields_alias_collisions_and_unknown_columns_are_unsupported(self):
        for column, value in ((18, None), (23, None), (3, "用户"), (26, "未知字段")):
            with self.subTest(column=column, value=value):
                self.fixture()
                edit(self.path, lambda b: setattr(b.active.cell(6, column), "value", value))
                result = read_ledger(self.path)
                self.assertEqual(result.status, "unsupported")
                self.assertIsNone(result.totals)

    def test_incomplete_summary_and_invalid_amounts_never_become_zero(self):
        self.fixture()
        edit(self.path, lambda b: setattr(b.active["E4"], "value", None))
        self.assertIn("INCOMPLETE_SUMMARY", codes(read_ledger(self.path)))
        for value in (True, "nan", "inf", "bad", "1e999", None):
            with self.subTest(value=value):
                self.fixture()
                edit(self.path, lambda b: setattr(b.active["Q7"], "value", value))
                result = read_ledger(self.path)
                self.assertEqual(result.status, "error")
                self.assertIn("INVALID_AMOUNT", codes(result))
                self.assertIsNone(result.totals)

    def test_footer_without_detail_is_error_and_first_row_header_has_no_preamble(self):
        self.fixture(0, "no_summary", False)
        self.assertIn("ORPHAN_ORDER_TOTAL", codes(read_ledger(self.path)))
        self.fixture(1, summary=False)
        edit(self.path, lambda b: b.active.delete_rows(1, 5))
        result = read_ledger(self.path)
        self.assertEqual(self.totals(result), (10, 4, 6))
        self.assertFalse(result.schema.writable)

    def test_formula_without_cache_and_inconsistent_cache_are_explicit_errors(self):
        self.fixture(1, "formula")
        self.assertIn("FORMULA_WITHOUT_CACHE", codes(read_ledger(self.path)))
        self.fixture(1, "formula_cached")
        result = read_ledger(self.path)
        self.assertEqual(result.status, "error")
        self.assertTrue({"FORMULA_CACHE_USED", "AMOUNT_MISMATCH"}.issubset(codes(result)))

    def test_consistent_synthetic_formula_cache_is_readonly_and_not_evaluated(self):
        self.fixture(1, "formula", False)
        edit(self.path, lambda b: (setattr(b.active["R7"], "value", 0), setattr(b.active["S7"], "value", 2)))
        with ZipFile(self.path) as package:
            parts = [(info, package.read(info.filename)) for info in package.infolist()]
        with ZipFile(self.path, "w", ZIP_DEFLATED) as package:
            for info, data in parts:
                if info.filename == "xl/worksheets/sheet1.xml":
                    data = fixtures.synthetic_formula_cache(data)
                package.writestr(info, data)
        before = self.path.read_bytes()
        result = read_ledger(self.path)
        self.assertEqual(self.totals(result), (2, 0, 2))
        self.assertFalse(result.schema.writable)
        self.assertIn("FORMULA_CACHE_USED", codes(result))
        self.assertEqual(self.path.read_bytes(), before)

    def test_same_second_distinct_order_references_are_separate_and_collisions_unresolved(self):
        self.fixture()
        result = read_ledger(self.path)
        self.assertEqual(sorted(result.group_totals.values()), [Decimal(10), Decimal(10)])
        edit(self.path, lambda b: setattr(b.active["C8"], "value", b.active["C7"].value))
        result = read_ledger(self.path)
        self.assertEqual(self.totals(result), (20, 8, 12))
        self.assertIn("AMBIGUOUS_LEGACY_GROUP", codes(result))
        self.assertEqual(result.group_totals, {})
        self.assertNotEqual(result.group_keys[7], result.group_keys[8])
        self.assertEqual(result.detail_count, 2)

    def test_legacy_export_boundaries_allow_same_prefix_without_assigning_ids(self):
        self.fixture(1, summary=False)
        def blocks(book):
            ws = book.active
            values = [c.value for c in ws[7]]
            footer = values[:4] + ["--- 本次总计 ---"] + [None] * 11 + [10, 4, 6]
            ws.append(footer)
            ws.append(values)
            ws.append(footer)
        edit(self.path, blocks)
        before = self.path.read_bytes()
        result = read_ledger(self.path)
        self.assertEqual(self.totals(result), (20, 8, 12))
        self.assertEqual(len(result.group_totals), 2)
        self.assertNotIn("AMBIGUOUS_LEGACY_GROUP", codes(result))
        self.assertFalse(result.schema.identified)
        self.assertEqual(self.path.read_bytes(), before)

    def test_unknown_sheet_or_package_extension_is_readable_but_not_writable(self):
        self.fixture()
        edit(self.path, lambda b: b.create_sheet("其他资料"))
        result = read_ledger(self.path)
        self.assertEqual(self.totals(result), (20, 8, 12))
        self.assertFalse(result.schema.writable)
        self.fixture()
        with ZipFile(self.path, "a") as package:
            package.writestr("customXml/item1.xml", "<synthetic/>")
        result = read_ledger(self.path)
        self.assertEqual(self.totals(result), (20, 8, 12))
        self.assertFalse(result.schema.writable)

    def test_missing_or_corrupt_files_are_errors_and_workbooks_close_on_failed_analysis(self):
        self.assertEqual(read_ledger(self.path).status, "error")
        primary, cached = Mock(), Mock()
        ws = Mock(max_row=7, max_column=26)
        ws.iter_rows.return_value = [(Mock(data_type="f"),)]
        primary.worksheets = [ws]
        with patch("quote_files.reading.open_book", side_effect=[primary, cached]), patch("quote_files.reading.analyze", side_effect=ValueError("synthetic")):
            result = analyze_raw(b"synthetic")
        self.assertEqual(result.status, "error")
        primary.close.assert_called_once()
        cached.close.assert_called_once()


class LedgerAdapterTests(unittest.TestCase):
    def test_current_mapped_and_legacy_alias_headers_append_without_rewriting_old_values_or_ids(self):
        for role in ("manager", "sales"):
            with self.subTest(role=role), application(role) as (app, *_):
                path = export_one(app)
                def reorder(book):
                    ws = book["利润明细"]
                    width = ws.max_column
                    for row in range(6, ws.max_row + 1):
                        values = [ws.cell(row, col).value for col in range(1, width + 1)][::-1]
                        for col, value in enumerate(values, 1):
                            ws.cell(row, col).value = "用户" if row == 6 and value == "淘宝ID" else value
                edit(path, reorder)
                original = open_book(path.read_bytes())
                try:
                    ws = original["利润明细"]
                    _, h = table(ws)
                    before = {name: ws.cell(7, col).value for name, col in h.items()}
                finally:
                    original.close()
                app.calculate(notify=False)
                self.assertTrue(app.export_to_excel(), app.feedback)
                result = read_ledger(path)
                self.assertEqual(result.status, "ok", result.issues)
                self.assertEqual(result.schema.name, "identified-v1-mapped")
                self.assertEqual(result["income"], 24)
                book = open_book(path.read_bytes())
                try:
                    ws = book["利润明细"]
                    _, h = table(ws)
                    self.assertEqual({name: ws.cell(7, h[name]).value for name in before}, before)
                    self.assertEqual(ws.cell(9, h["销售额"]).value, "12.00")
                    self.assertNotEqual(ws.cell(7, h["订单ID"]).value, ws.cell(9, h["订单ID"]).value)
                    self.assertEqual(ws.cell(9, h["销售额"]).number_format, "#,##0.00")
                finally:
                    book.close()

    def test_all_new_identity_raw_and_metadata_strings_are_explicit_text(self):
        for role in ("manager", "sales"):
            with self.subTest(role=role), application(role) as (app, *_):
                app.user_name.set("=1+1")
                app.order_no.set("000001")
                app.wechat_no.set("=synthetic")
                app.add_fixture_row(product="P02面积", q="1+1", h="1+1", w="1")
                app.calculate(notify=False)
                path = Path(app.export_to_excel())
                book = open_book(path.read_bytes())
                try:
                    self.assertFalse(any(c.data_type == "f" for ws in book.worksheets for row in ws.iter_rows() for c in row))
                    ws = book["利润明细"]
                    _, h = table(ws)
                    for coordinate, expected in (("B3", "=1+1"), ("D3", "000001"), ("F3", "=synthetic")):
                        self.assertEqual(ws[coordinate].value, expected)
                        self.assertEqual(ws[coordinate].data_type, "s")
                    self.assertEqual(ws.cell(7, h["数量原文"]).value, "1+1")
                    self.assertEqual(ws.cell(7, h["高度原文"]).value, "1+1")
                finally:
                    book.close()

    def test_existing_template_is_never_cleared_and_fresh_header_stays_at_row_six(self):
        book = openpyxl.Workbook()
        try:
            create_template(book.active, "synthetic-time")
            self.assertEqual(tuple(c.value for c in book.active[6]), DETAIL_HEADERS)
            self.assertEqual(book.active["A1"].value, "六六派对利润明细")
            before = tuple(tuple(c.value for c in row) for row in book.active.iter_rows())
            with self.assertRaisesRegex(FileFailure, "禁止模板重置"):
                create_template(book.active, "later")
            self.assertEqual(tuple(tuple(c.value for c in row) for row in book.active.iter_rows()), before)
        finally:
            book.close()

    def test_post_mutation_amount_check_refuses_publish_and_preserves_original_and_receipts(self):
        from quote_app.integration import append_quote
        with application("manager") as (app, *_):
            path = export_one(app)
            before = path.read_bytes()
            app.calculate(notify=False)
            def invalid_footer(ws, *args, **kwargs):
                append_quote(ws, *args, **kwargs)
                ws.cell(ws.max_row, kwargs["headers"]["利润"], "999.00")
            with patch("quote_app.integration.append_quote", side_effect=invalid_footer):
                self.assertFalse(app.export_to_excel())
            self.assertIn("写后工作簿核对失败", app.feedback)
            self.assertEqual(path.read_bytes(), before)

    def test_readonly_files_refuse_append_and_edit_with_unchanged_original(self):
        for variant in ("legacy", "formula", "unknown_sheet", "extension"):
            with self.subTest(variant=variant), application("manager") as (app, *_):
                path = export_one(app)
                if variant == "legacy":
                    fixtures.workbook(path)
                elif variant == "formula":
                    fixtures.workbook(path, 1, "formula")
                elif variant == "unknown_sheet":
                    edit(path, lambda b: b.create_sheet("未知资料"))
                else:
                    with ZipFile(path, "a") as package:
                        package.writestr("customXml/item1.xml", "<synthetic/>")
                before = path.read_bytes()
                app.current_excel_path = str(path)
                app.excel_sheet_var = tk.StringVar(app.root, "利润明细")
                app.get_selected_excel_row_index = lambda: 7
                book = app.read_excel_view(path, "利润明细")
                book.close()
                self.assertEqual(app._excel_view_orders, {})
                with self.assertRaises(FileFailure):
                    app._edit_order("append_remark", "synthetic")
                app.calculate(notify=False)
                self.assertFalse(app.export_to_excel())
                self.assertEqual(path.read_bytes(), before)

    def test_legitimate_zero_counts_and_failures_display_dash_with_partial_aggregation(self):
        with application("manager") as (app, folder, *_):
            zero, broken, unknown = [folder / (name + ".xlsx") for name in ("zero", "broken", "unknown")]
            fixtures.workbook(zero, 0)
            broken.write_bytes(b"synthetic corrupt")
            fixtures.workbook(unknown, variant="duplicate_headers")
            app.get_report_excel_paths = lambda: [str(zero), str(broken), str(unknown)]
            app.excel_file_tree = ViewTree()
            for key in ("count", "income", "cost", "profit", "margin"):
                setattr(app, "excel_total_" + key, tk.StringVar(app.root))
            result = app.aggregate_all_excels(silent=True)
            self.assertEqual((result["status"], result["valid_files"], result["errors"], result["unsupported"]), ("partial", 1, 1, 1))
            self.assertEqual(app.excel_total_income.get(), "0.00")
            self.assertIn("部分", app.excel_total_count.get())
            self.assertEqual(app.excel_file_tree.rows[str(broken)][2], "—")
            self.assertIn("读取失败", app.excel_file_tree.rows[str(broken)][-1])
            self.assertIn("格式不支持", app.excel_file_tree.rows[str(unknown)][-1])
            app.get_report_excel_paths = lambda: [str(broken)]
            app.aggregate_all_excels(silent=True)
            self.assertEqual(app.excel_total_income.get(), "—")
            app.get_report_excel_paths = lambda: []
            app.aggregate_all_excels(silent=True)
            self.assertEqual(app.excel_total_income.get(), "0.00")

    def test_sales_list_count_and_unresolved_group_total_are_visible_not_zero(self):
        with application("sales") as (app, folder, *_):
            path = folder / "collision.xlsx"
            fixtures.workbook(path)
            edit(path, lambda b: setattr(b.active["C8"], "value", b.active["C7"].value))
            app.get_report_excel_paths = lambda: [str(path)]
            app.excel_file_tree = ViewTree()
            FileReadingMixin.refresh_excel_files(app)
            self.assertEqual(app.excel_file_tree.rows[str(path)][2], 2)
            self.assertIn("未合并总销售额", app.excel_file_tree.rows[str(path)][-1])
            app.current_excel_path = str(path)
            app.excel_sheet_var = tk.StringVar(app.root, "利润明细")
            app.excel_view_tree = ViewTree()
            app.load_current_excel_sheet()
            self.assertEqual(len(app.excel_view_tree.rows), 2)
            self.assertTrue(all(row[-1] == "—（历史归属待核对）" for row in app.excel_view_tree.rows.values()))

    def test_current_page_builders_wire_every_list_heading_and_read_status_column(self):
        for role in ("manager", "sales"):
            with self.subTest(role=role), application(role) as (app, _folder, ns, _entry):
                app.excel_page = Mock()
                for key in ("excel_selected_file", "excel_sheet_var", "excel_total_count", "excel_total_income", "excel_total_cost", "excel_total_profit", "excel_total_margin"):
                    setattr(app, key, tk.StringVar(app.root))
                app.install_storage_controls = Mock()
                app.show_excel_list = Mock()
                app.schedule_quote_callback = Mock()
                widgets = []
                def tree(*args, **kwargs):
                    widget = Mock()
                    widget.columns = kwargs.get("columns", ())
                    widgets.append(widget)
                    return widget
                with ExitStack() as stack:
                    for name in ("Frame", "Label", "Button", "LabelFrame", "Scrollbar", "Combobox"):
                        stack.enter_context(patch.object(ns["ttk"], name))
                    stack.enter_context(patch.object(ns["ttk"], "Treeview", side_effect=tree))
                    stack.enter_context(patch.object(ns["tk"], "Text"))
                    app.setup_excel_page()
                self.assertIn("status", widgets[0].columns)
                self.assertEqual({call.args[0] for call in widgets[0].heading.call_args_list}, set(widgets[0].columns))
                self.assertEqual({call.args[0] for call in widgets[0].column.call_args_list}, set(widgets[0].columns))

    def test_current_gui_diagnostic_report_can_serialize_typed_summary_after_shared_adapter_run(self):
        import p02_gui_check as gui
        with application("manager") as (manager, folder, *_), application("sales") as (sales, _, ns, _):
            ns["LAN_PRICE_CONFIG_FILE"] = manager.config_file
            for app in (manager, sales):
                app.excel_file_tree = ViewTree()
                app.refresh_excel_files = lambda a=app: FileReadingMixin.refresh_excel_files(a)
            manager.excel_selected_file = tk.StringVar(manager.root)
            manager.excel_sheet_var = tk.StringVar(manager.root)
            manager.excel_sheet_combo = Mock()
            manager.excel_view_tree = ViewTree()
            manager.show_excel_detail = Mock()
            manager.load_current_excel_remark = Mock()
            sales.storage_entries = {"synthetic": Mock()}
            sales.storage_entries["synthetic"].cget.return_value = "readonly"
            sales.add_calc_row = lambda _display, base, _single: sales.add_fixture_row(product=base)
            sales._write_quote_screenshot = lambda _payload: str(Path(sales.image_dir) / "synthetic.png")
            try:
                report = gui.shared_directory_smoke(manager, sales, folder)
                decoded = json.loads(json.dumps(report, ensure_ascii=False))
                self.assertEqual(decoded["manager_summary_income_cost_profit"]["income"], "12.00")
                self.assertEqual(decoded["manager_summary_income_cost_profit"]["cost"], "4.00")
            finally:
                for app in (manager, sales):
                    for callback in app.root.tk.call("after", "info"):
                        app.root.after_cancel(callback)


if __name__ == "__main__":
    unittest.main()
