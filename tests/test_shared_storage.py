"""The two adapters share persisted storage; no production drive/share access."""
import json
from pathlib import Path
import unittest
from unittest.mock import patch

from test_application_integration import application
from quote_app.storage import publish_directory, resolve_directory
from quote_core import parse_price_config
import p02_gui_check as gui


class SharedStorageTests(unittest.TestCase):
    def test_local_manager_and_unc_sales_resolve_same_published_relative_location(self):
        local = r"E:\14T共享盘\葛宁\PyInstaller\config.json"
        remote = r"\\192.168.4.38\14T共享盘\葛宁\PyInstaller\config.json"
        for child in ("利润表", r"新位置\订单", "截图"):
            saved = publish_directory(
                local.rsplit("\\", 1)[0] + "\\" + child, local, child)
            self.assertEqual(saved, child)
            self.assertEqual(resolve_directory(saved, remote, child, sales=True),
                             remote.rsplit("\\", 1)[0] + "\\" + child)

    def test_external_unc_and_share_root_are_explicit_portable_locations(self):
        config = r"E:\data\config.json"
        for path in (r"\\other\share", r"\\other\share\利润表"):
            self.assertEqual(publish_directory(path, config, "利润表"), path)
            self.assertEqual(resolve_directory(path, r"\\host\share\config.json", "利润表", sales=True), path)

    def test_local_drive_cannot_be_published_outside_shared_root_or_used_on_sales(self):
        with self.assertRaises(ValueError):
            publish_directory(r"D:\orders", r"E:\data\config.json", "利润表")
        with self.assertRaises(ValueError):
            resolve_directory(r"E:\data\利润表", r"\\host\share\config.json", "利润表", sales=True)
        for path in (r"E:orders", r"\orders", r"..\orders"):
            with self.subTest(path=path), self.assertRaises(ValueError):
                resolve_directory(path, r"E:\data\config.json", "利润表")

    def test_optional_screenshot_directory_is_whole_validated_for_both_schemas(self):
        for data in (gui.FIXTURE, {"P": [1, 0, 1], "screenshot_dir": "截图"}):
            result = parse_price_config(json.dumps(data), price_version_id="v", legacy_single_names=())
            self.assertEqual(result.status, "ok")
            self.assertIsNotNone(result.config.screenshot_dir)
            for invalid in (False, None, 123, "bad\x00"):
                bad = {**data, "screenshot_dir": invalid}
                result = parse_price_config(json.dumps(bad), price_version_id="v", legacy_single_names=())
                self.assertIsNone(result.config)
                self.assertEqual(result.issues[0].field, "screenshot_dir")

    def test_manager_saved_directories_survive_reload_and_unsaved_fields_cannot_redirect_output(self):
        with application("manager") as (app, folder, *_):
            reports, images = folder / "新订单", folder / "新图片"
            app.export_dir_var.set(str(reports))
            app.screenshot_dir_var.set(str(images))
            self.assertTrue(app.apply_export_dir())
            saved = json.loads((folder / "config.json").read_text(encoding="utf-8"))
            self.assertEqual(saved["export_dir"], "新订单")
            self.assertEqual(saved["screenshot_dir"], "新图片")
            app.report_dir = "wrong"
            app._storage_version = None  # Reconstruct startup binding, without GUI.
            app.reload_price_settings_from_json()
            self.assertEqual(app.report_dir, str(reports))
            self.assertEqual(app.image_dir, str(images))
            app.add_fixture_row(product="P02面积")
            app.calculate(notify=False)
            app.confirm_current_prices()
            app.export_dir_var.set(str(folder / "unsaved"))
            path = app.export_to_excel()
            self.assertEqual(Path(path).parent, reports)
            self.assertFalse((folder / "unsaved").exists())

    def test_sales_export_is_found_and_read_by_manager_after_directory_change(self):
        with application("manager") as (manager, folder, *_), application("sales") as (sales, _, ns, _):
            ns["LAN_PRICE_CONFIG_FILE"] = manager.config_file
            manager.export_dir_var.set(str(folder / "共同订单"))
            manager.screenshot_dir_var.set(str(folder / "共同图片"))
            self.assertTrue(manager.apply_export_dir())
            sales.get_report_excel_paths()
            self.assertEqual(sales.report_dir, manager.report_dir)
            self.assertEqual(sales.image_dir, manager.image_dir)
            sales.reload_price_settings_from_json()
            sales.add_fixture_row(product="P02面积")
            sales.user_name.set("000-sales")
            sales.calculate(notify=False)
            sales.confirm_current_prices()
            # Even accidental programmatic changes to the display cannot redirect.
            sales.export_dir_var.set(str(folder / "sales-override"))
            path = sales.export_to_excel()
            self.assertIn(path, manager.get_report_excel_paths())
            summary = manager.parse_excel_summary(path)
            self.assertEqual((summary["income"], summary["cost"], summary["profit"]), (12, 4, 8))
            from openpyxl import load_workbook
            with open(path, "rb") as stream:
                book = load_workbook(stream, read_only=True)
                self.assertEqual(book["利润明细"].cell(7, 2).value, "000-sales")
                book.close()
            self.assertFalse((folder / "sales-override").exists())

    def test_sales_cannot_save_select_or_apply_storage(self):
        with application("sales") as (app, folder, *_):
            before = gui.files(folder)
            app.export_dir_var.set(str(folder / "override"))
            app.screenshot_dir_var.set(str(folder / "override-image"))
            with patch("tkinter.filedialog.askdirectory", side_effect=AssertionError("sales chooser ran")):
                self.assertFalse(app.choose_export_dir())
                self.assertFalse(app.choose_storage_dir(app.screenshot_dir_var))
                self.assertFalse(app.apply_export_dir())
                self.assertFalse(app.reset_export_dir_to_local())
                self.assertFalse(app.save_settings())
            self.assertEqual(gui.files(folder), before)

    def test_bad_directory_retains_snapshot_and_blocks_writes_and_list_fallback(self):
        for role in ("manager", "sales"):
            with application(role) as (app, folder, *_):
                retained = app.active_price_config
                directories = app.report_dir, app.image_dir
                bad = {**gui.FIXTURE, "export_dir": "../unconfirmed"}
                (folder / "config.json").write_text(json.dumps(bad), encoding="utf-8")
                before = gui.files(folder)
                self.assertFalse(app.reload_price_settings_from_json())
                self.assertIs(app.active_price_config, retained)
                self.assertEqual((app.report_dir, app.image_dir), directories)
                self.assertEqual(app.get_report_excel_paths(), [])
                self.assertFalse(app.export_to_excel())
                self.assertEqual(gui.files(folder), before)

    def test_invalid_price_cannot_claim_a_successful_directory_save(self):
        with application("manager") as (app, folder, *_):
            before = gui.files(folder)
            app.setting_rows[0]["sale"].set("bad")
            app.export_dir_var.set(str(folder / "new"))
            self.assertFalse(app.apply_export_dir())
            self.assertEqual(gui.files(folder), before)
            self.assertNotEqual(app.report_dir, str(folder / "new"))

    def test_close_cancels_only_owned_delayed_callbacks_and_never_saves(self):
        with application("manager") as (app, *_):
            unrelated = app.root.after(10000, lambda: None)
            app.schedule_quote_callback(lambda: self.fail("closed refresh ran"), delay=10000)
            app.schedule_quote_callback(lambda: self.fail("closed screenshot ran"))
            with patch.object(app.root, "destroy") as destroy, patch.object(app, "save_settings", side_effect=AssertionError("close saved")):
                app.on_close()
                destroy.assert_called_once()
            self.assertEqual(set(app.root.tk.call("after", "info")), {unrelated})
            app.root.eval("update idletasks")
            app.root.after_cancel(unrelated)

    def test_executed_callback_is_removed_and_background_errors_fail_report(self):
        with application("manager") as (app, *_):
            called = []
            app.schedule_quote_callback(lambda: called.append("ok"))
            app.root.eval("update idletasks")
            self.assertEqual(called, ["ok"])
            self.assertFalse(app._after_jobs)
            errors = gui.track_callback_errors(app.root)
            app.root.tk.call("after", 0, "missing_refresh_excel_command")
            result = {"status": "passed"}
            gui.finish_callbacks(app.root, result, errors)
            self.assertEqual(result["status"], "failed")
            self.assertIn("missing_refresh_excel_command", errors[0])


if __name__ == "__main__":
    unittest.main()
