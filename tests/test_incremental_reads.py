"""Synthetic cache correctness/counts; only Windows runs are product evidence."""
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
import io
import os
from pathlib import Path
import tempfile
from threading import Event
import time
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_STORED

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from test_application_integration import application
import p01_fixtures as fixtures
from test_background_io import asynchronous, wait_for
from test_workbook_transactions import export_one
from quote_app.jobs import Cancelled
from quote_files import cache as cached
from quote_files.cache import LedgerCache
from quote_files.documents import parse_document
from quote_files.metadata import FileVersion, canonical, file_version
from quote_files.reading import read_ledger
from quote_files.views import read_view


class Context:
    def __init__(self):self.cancelled = False
    def check(self):
        if self.cancelled:raise Cancelled()


@contextmanager
def workbook_counts():
    with patch('openpyxl.load_workbook', wraps=openpyxl.load_workbook) as opened:
        yield opened


def rewrite_xml(raw, before, after):
    out = io.BytesIO()
    with ZipFile(io.BytesIO(raw)) as source, ZipFile(out, 'w', ZIP_STORED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == 'xl/worksheets/sheet1.xml':data = data.replace(before, after)
            info.compress_type = ZIP_STORED;target.writestr(info, data)
    return out.getvalue()


class IncrementalReadTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='777-p04-cache-')
        self.folder = Path(self.temp.name)
        self.path = self.folder / '中文账本.xlsx'
        self.context = Context()
    def tearDown(self):self.temp.cleanup()
    def fixture(self, count=20, name=None, **kwargs):
        path = self.folder / name if name else self.path
        fixtures.workbook(path, count, **kwargs)
        return path
    def scan(self, cache, paths):
        rows = {str(p): (cache.get(p, self.context, detail=False).ledger(), 'synthetic') for p in paths}
        cache.prune(self.folder, paths)
        return cache.directory_result(self.folder, rows)
    def stored_payloads(self):
        self.fixture(1)
        book = openpyxl.load_workbook(self.path);book.active['B7'] = 'AAA';book.save(self.path);book.close()
        old = rewrite_xml(self.path.read_bytes(), b'AAA', b'AAA')
        new = rewrite_xml(old, b'AAA', b'BBB')
        self.assertEqual(len(old), len(new));self.path.write_bytes(old)
        return old, new

    def test_each_source_sheet_is_iterated_once_for_summary_and_both_views(self):
        self.fixture()
        original = Worksheet.iter_rows;iterations = []
        def iterate(ws, *args, **kwargs):iterations.append(ws.title);return original(ws, *args, **kwargs)
        cache = LedgerCache()
        with workbook_counts() as opened, patch.object(Worksheet, 'iter_rows', iterate):
            doc = cache.get(self.path, self.context)
            manager = read_view(self.path, None, 'manager', self.context, doc)
            sales = read_view(self.path, None, 'sales', self.context, doc)
            self.assertEqual(opened.call_count, 1)
            self.assertEqual(iterations, ['利润明细'])
        self.assertEqual(doc.result.totals.income_cents, 20000)
        self.assertEqual(len(sales.rows), 20)
        self.assertEqual(len(manager.rows), 26)

    def test_twenty_and_forty_rows_do_not_reopen_worksheet_xml_per_cell(self):
        counts = []
        original = ZipFile.open
        for size in (20, 40):
            self.fixture(size);opened = []
            def xml(package, name, *args, **kwargs):
                path = getattr(name, 'filename', name)
                if path == 'xl/worksheets/sheet1.xml':opened.append(path)
                return original(package, name, *args, **kwargs)
            with patch.object(ZipFile, 'open', xml):
                doc = parse_document(self.path.read_bytes())
                view = read_view(self.path, None, 'sales', self.context, doc)
            self.assertEqual(len(view.rows), size)
            counts.append(len(opened))
        self.assertEqual(counts, [1, 1])

    def test_directory_warm_index_survives_full_document_lru_eviction(self):
        paths = [self.fixture(20, f'{i}.xlsx') for i in range(5)]
        cache = LedgerCache(max_files=2)
        with workbook_counts() as opened:
            cold = self.scan(cache, paths);self.assertEqual(opened.call_count, 5)
            opened.reset_mock();warm = self.scan(cache, paths)
            self.assertEqual(opened.call_count, 0)
        self.assertEqual(warm.aggregate['income_cents'], 100000)
        self.assertEqual(warm.aggregate['changed_files'], 0)
        self.assertGreater(cache.metrics()['evictions'], 0)
        self.assertEqual(cold.aggregate['income_cents'], warm.aggregate['income_cents'])

    def test_one_changed_file_adjusts_only_its_amounts_without_reading_other_books(self):
        paths = [self.fixture(20, f'{i}.xlsx') for i in range(3)]
        cache = LedgerCache();self.scan(cache, paths)
        fixtures.workbook(paths[1], 21)
        with workbook_counts() as opened:result = self.scan(cache, paths)
        self.assertEqual(opened.call_count, 1)
        self.assertEqual(result.aggregate['changed_files'], 1)
        self.assertEqual(result.aggregate['income_cents'], 61000)
        self.assertEqual(result.aggregate['cost_cents'], 24400)

    def test_removed_file_leaves_totals_and_index_without_reparsing_remaining_files(self):
        paths = [self.fixture(2, f'{i}.xlsx') for i in range(3)]
        cache = LedgerCache();self.scan(cache, paths);paths[1].unlink()
        with workbook_counts() as opened:result = self.scan(cache, [paths[0], paths[2]])
        self.assertEqual(opened.call_count, 0)
        self.assertEqual(result.aggregate['income_cents'], 4000)
        self.assertEqual(result.aggregate['changed_files'], 1)
        self.assertEqual(cache.metrics()['index_files'], 2)
        self.assertEqual(cache.get(paths[1], self.context).result.status, 'error')

    def test_same_size_in_place_write_with_restored_mtime_is_detected(self):
        old, new = self.stored_payloads();cache = LedgerCache()
        previous = cache.get(self.path, self.context);stamp = self.path.stat()
        self.path.write_bytes(new);os.utime(self.path, ns=(stamp.st_atime_ns, stamp.st_mtime_ns))
        current = cache.get(self.path, self.context)
        self.assertEqual(current.book['利润明细']['B7'].value, 'BBB')
        self.assertNotEqual(previous.result.source_sha256, current.result.source_sha256)

    def test_same_size_atomic_replacement_with_restored_mtime_is_detected(self):
        old, new = self.stored_payloads();cache = LedgerCache();cache.get(self.path, self.context)
        stamp = self.path.stat();other = self.folder / '替换文件.xlsx';other.write_bytes(new)
        os.utime(other, ns=(stamp.st_atime_ns, stamp.st_mtime_ns));os.replace(other, self.path)
        self.assertEqual(cache.get(self.path, self.context).book['利润明细']['B7'].value, 'BBB')

    def test_weak_metadata_verifies_bytes_and_keeps_full_view_separate_from_summary(self):
        old, new = self.stored_payloads();cache = LedgerCache()
        weak = FileVersion(len(old), 0, ('unverified',), False, 'synthetic-weak')
        with patch.object(cached, 'file_version', return_value=weak), workbook_counts() as opened:
            cache.get(self.path, self.context)
            summary = cache.get(self.path, self.context, detail=False)
            full = cache.get(self.path, self.context)
            self.assertIsNotNone(full.book);self.assertEqual(summary.result.detail_count, 1)
            self.assertEqual(opened.call_count, 1)
            self.path.write_bytes(new)
            changed = cache.get(self.path, self.context)
            self.assertEqual(changed.book['利润明细']['B7'].value, 'BBB')
            self.assertEqual(opened.call_count, 2)
        self.assertGreater(cache.metrics()['weak_hash_probes'], 0)

    def test_mtime_only_change_revalidates_same_bytes_without_decoding_again(self):
        self.fixture();cache = LedgerCache();cache.get(self.path, self.context)
        stamp = self.path.stat();os.utime(self.path, ns=(stamp.st_atime_ns, stamp.st_mtime_ns+1000000))
        with workbook_counts() as opened:cache.get(self.path, self.context)
        self.assertEqual(opened.call_count, 0)

    def test_parser_version_change_rebuilds_even_when_bytes_and_metadata_match(self):
        self.fixture();cache = LedgerCache();cache.get(self.path, self.context)
        with patch.object(cached, 'PARSER_VERSION', 'synthetic-parser-v2'), workbook_counts() as opened:
            self.assertEqual(cache.get(self.path, self.context).result.status, 'ok')
            self.assertEqual(opened.call_count, 1)

    def test_failures_and_unsupported_formats_remain_partial_until_actual_correction(self):
        good = self.fixture(2);bad = self.folder / '损坏.xlsx';bad.write_bytes(b'bad')
        unsupported = self.fixture(1, '不支持.xlsx', variant='duplicate_headers')
        cache = LedgerCache();partial = self.scan(cache, [good, bad, unsupported]).aggregate
        self.assertEqual((partial['valid_files'], partial['errors'], partial['unsupported']), (1,1,1))
        self.assertEqual(partial['income_cents'], 2000)
        fixtures.workbook(bad, 1);fixtures.workbook(unsupported, 1)
        fixed = self.scan(cache, [good, bad, unsupported]).aggregate
        self.assertEqual((fixed['status'],fixed['valid_files'],fixed['income_cents']), ('ok',3,4000))

    def test_formula_cache_is_counted_separately_and_missing_cache_stays_error(self):
        self.fixture(1, variant='formula')
        with workbook_counts() as opened:doc = LedgerCache().get(self.path, self.context)
        self.assertEqual(opened.call_count, 2)
        self.assertEqual(doc.formula_cache_opens, 1)
        self.assertEqual(doc.result.status, 'error')
        self.assertIn('FORMULA_WITHOUT_CACHE', {i.code for i in doc.result.issues})

    def test_parallel_scan_and_view_share_one_decode_and_all_handles_finish(self):
        self.fixture();cache = LedgerCache();start, release = Event(), Event()
        original = cached.parse_document
        def slow(raw):start.set();release.wait(3);return original(raw)
        try:
            with patch.object(cached, 'parse_document', slow), workbook_counts() as opened, ThreadPoolExecutor(max_workers=2) as pool:
                first = pool.submit(cache.get, self.path, self.context)
                self.assertTrue(start.wait(1))
                second = pool.submit(cache.get, self.path, self.context)
                wait_for(lambda: cache.metrics().get('coalesced_waits',0)>0)
                release.set();a,b = first.result(2),second.result(2)
                self.assertEqual(opened.call_count, 1)
                self.assertEqual(a.result.source_sha256,b.result.source_sha256)
        finally:release.set()

    def test_write_invalidation_discards_inflight_old_result_and_retries_new_bytes(self):
        old,new = self.stored_payloads();cache = LedgerCache();start,release = Event(),Event()
        original = cached.parse_document;count = 0
        def slow(raw):
            nonlocal count
            count += 1
            if count==1:start.set();release.wait(3)
            return original(raw)
        try:
            with patch.object(cached,'parse_document',slow),ThreadPoolExecutor(max_workers=1) as pool:
                job=pool.submit(cache.get,self.path,self.context);self.assertTrue(start.wait(1))
                cache.invalidate(self.path);self.path.write_bytes(new);release.set();doc=job.result(2)
            self.assertEqual(doc.book['利润明细']['B7'].value,'BBB')
            self.assertEqual(cache.metrics()['discarded_changes'],1)
            self.assertEqual(cache.metrics()['workbook_opens'],2)
        finally:release.set()

    def test_cancelled_load_never_publishes_partial_cache_or_closes_owner_ui(self):
        self.fixture();cache=LedgerCache();start,release=Event(),Event();original=cached.parse_document
        def slow(raw):start.set();release.wait(3);return original(raw)
        try:
            with patch.object(cached,'parse_document',slow),ThreadPoolExecutor(max_workers=1) as pool:
                job=pool.submit(cache.get,self.path,self.context);self.assertTrue(start.wait(1))
                self.context.cancelled=True;release.set()
                with self.assertRaises(Cancelled):job.result(2)
            self.assertEqual(cache.metrics()['files'],0)
            self.context.cancelled=False;self.assertEqual(cache.get(self.path,self.context).result.status,'ok')
        finally:release.set()

    def test_memory_budget_evicts_full_documents_but_retains_exact_summary(self):
        self.fixture();cache=LedgerCache(max_bytes=1)
        first=cache.get(self.path,self.context,detail=False)
        with workbook_counts() as opened:second=cache.get(self.path,self.context,detail=False)
        self.assertEqual(opened.call_count,0)
        self.assertEqual(first.result.totals,second.result.totals)
        self.assertEqual(cache.metrics()['estimated_bytes'],0)
        self.assertEqual(cache.metrics()['index_files'],1)

    def test_clear_scope_cannot_reuse_an_old_directory_index(self):
        self.fixture();cache=LedgerCache();cache.get(self.path,self.context);cache.clear()
        with workbook_counts() as opened:cache.get(self.path,self.context)
        self.assertEqual(opened.call_count,1)
        self.assertEqual(cache.metrics()['clears'],1)

    def test_cached_ui_ledger_maps_are_copies_and_cannot_change_next_projection(self):
        self.fixture();cache=LedgerCache();doc=cache.get(self.path,self.context)
        original=dict(doc.result.group_totals);shown=doc.ledger();shown.group_totals.clear()
        self.assertEqual(cache.get(self.path,self.context).ledger().group_totals,original)


class IncrementalAdapterTests(unittest.TestCase):
    def test_windows_diagnostic_cache_steps_use_real_adapters_and_exact_amounts(self):
        from test_order_management import display
        import p04_cache_check as check
        with application('manager') as (manager,folder,*_),application('sales') as (sales,*_):
            path=export_one(manager);manager.calculate(notify=False);manager.export_to_excel()
            display(manager,path)
            sales.data_source_var.set(str(folder));self.assertTrue(sales.apply_connection())
            button=manager.excel_note_button;button._state='normal'
            def configure(**values):button._state=values.get('state',button._state)
            button.configure.side_effect=configure
            button.cget.side_effect=lambda name:button._state
            button.invoke=manager.save_current_excel_remark
            with asynchronous(manager),asynchronous(sales):
                def await_task(app,task):
                    self.assertTrue(task);yield lambda:task.number not in app._io_callbacks
                report={'checks':{}}
                for ready in check.cache_steps(manager,sales,folder,report,await_task):
                    wait_for(ready,lambda:(manager._poll_io(),sales._poll_io()))
                self.assertEqual(len(report['checks']),4)
                self.assertEqual(set(report['checks'].values()),{'passed'})
                self.assertEqual(report['cache_profile']['cold_workbook_opens'],4)
                self.assertEqual(report['cache_profile']['warm_workbook_opens'],0)

    def test_actual_background_write_reloads_only_target_and_preserves_other_file_index(self):
        with application('manager') as (app,folder,*_):
            target=export_one(app);app.calculate(notify=False)
            extra=folder/'reports'/'其他历史.xlsx';fixtures.workbook(extra,20)
            with asynchronous(app):
                task=app.refresh_excel_files();wait_for(lambda: not app._io_callbacks,app._poll_io)
                before=app._ledger_cache.metrics()['workbook_opens']
                task=app.export_to_excel();wait_for(lambda: not app._io_callbacks,app._poll_io)
                self.assertEqual(app._ledger_cache.metrics()['workbook_opens']-before,1)
                self.assertEqual(read_ledger(target).totals.income_cents,2400)
                self.assertEqual(app.excel_total_income.get(),'224.00')
                self.assertEqual(app.excel_read_results.aggregate['changed_files'],1)

    def test_external_change_refreshes_open_view_and_deletion_revokes_old_edit_lease(self):
        with application('manager') as (app,folder,*_):
            path=export_one(app)
            with asynchronous(app):
                task=app.refresh_excel_files();wait_for(lambda: not app._io_callbacks,app._poll_io)
                task=app.open_excel_path(str(path));wait_for(lambda: not app._io_callbacks,app._poll_io)
                before=app._excel_view_hash
                book=openpyxl.load_workbook(path);book['利润明细']['B7']='外部改名';book.save(path);book.close()
                task=app.refresh_excel_files();wait_for(lambda: not app._io_callbacks,app._poll_io)
                self.assertNotEqual(app._excel_view_hash,before)
                self.assertIn('外部改名',app.excel_view_tree.rows['row_7'])
                path.unlink();task=app.refresh_excel_files();wait_for(lambda: not app._io_callbacks,app._poll_io)
                self.assertIsNone(app._excel_view_hash)
                self.assertEqual(app._excel_view_orders,{})
                self.assertEqual(app.excel_total_income.get(),'0.00')

    def test_windows_gui_check_is_blocked_off_target_without_constructing_tk(self):
        import p04_cache_check as check
        with patch.object(check,'run_checks',return_value={'status':'passed'}),patch.object(check.platform,'system',return_value='Linux'),patch.object(check,'run_gui') as gui,patch('baseline_check.emit'):
            self.assertEqual(check.main(['--gui-smoke']),2);gui.assert_not_called()


if __name__=='__main__':unittest.main()
