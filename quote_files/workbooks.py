"""Stable identity and replay receipts in one workbook transaction.

Known tagged layouts use canonical column mapping; historical untyped or
nonconforming files remain read-only. No cost/privacy claims: P05 owns auth.
"""
from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
import io
import json
import uuid
from zipfile import ZipFile


from .transactions import FileConflict, FileFailure, transact
from .schema import DETAIL_HEADERS, IDENTITY_HEADERS, force_text, inspect_schema, table_header

LOG_SHEET = "_写入事务"
LOG_HEADERS = ("请求ID", "操作", "内容哈希", "订单ID", "文件版本", "提交时间", "结果JSON")
KINDS = {"normal", "discount", "order_total"}


@dataclass(frozen=True)
class Request:
    request_id: str
    fingerprint: str
    operation: str
    order_id: str | None = None


def request(operation, value, *, order_id=None):
    encoded = json.dumps({"operation": operation, "order_id": order_id, "value": value}, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return Request(str(uuid.uuid4()), hashlib.sha256(encoded).hexdigest(), operation, order_id)


def open_book(raw, *, data_only=False, for_write=True):
    if raw is None:
        raise FileFailure("文件不存在")
    with ZipFile(io.BytesIO(raw)) as package:
        entries = package.infolist()
        if (len(entries) > 2048 or sum(x.file_size for x in entries) > 64 * 1024 * 1024
                or len({x.filename for x in entries}) != len(entries)):
            raise FileFailure("工作簿容量/压缩包结构不支持写入")
        for entry in entries:
            name = entry.filename.lower()
            if for_write and not data_only and any(token in name for token in ("externallinks", "vbaproject", "customxml", "embeddings", "drawings", "vml")):
                raise FileFailure("含未支持扩展/外链/宏/图片的工作簿只读保留")
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(raw), data_only=data_only, keep_links=for_write)


def table(ws):
    return table_header(ws)


def writable(book, base_headers):
    schema = inspect_schema(book)
    if tuple(base_headers) != DETAIL_HEADERS:
        raise FileFailure("未知写入字段契约，原件只读保留")
    if not schema.writable:
        raise FileFailure("；".join(schema.readonly_reasons))
    from .reading import analyze
    result = analyze(book)
    if result.status != "ok":
        raise FileFailure("；".join(i.message for i in result.issues))
    return book["利润明细"], schema.header_row, schema.columns


def validate_identity(ws, row, headers):
    present_identity = set(IDENTITY_HEADERS) & set(headers)
    if not present_identity:
        return
    if present_identity != set(IDENTITY_HEADERS):
        raise FileFailure("订单身份列不完整，原件只读保留")
    ids, groups = set(), {}
    for index in range(row + 1, ws.max_row + 1):
        values = [c.value for c in ws[index]]
        if all(v is None for v in values):
            continue
        if "行类型" not in headers or ws.cell(index, headers["行类型"]).value not in KINDS:
            raise FileFailure("历史记录没有明确行类型；禁止猜测后追加，原件只读保留")
        if "订单ID" in headers and ws.cell(index, headers["订单ID"]).value:
            for key in ("订单ID", "明细ID", "请求ID"):
                value = ws.cell(index, headers[key]).value
                if type(value) is not str or str(uuid.UUID(value)) != value:
                    raise FileFailure("订单身份损坏，原件只读保留")
            line_id = ws.cell(index, headers["明细ID"]).value
            if line_id in ids:
                raise FileFailure("明细编号重复，原件只读保留")
            ids.add(line_id)
            version = ws.cell(index, headers["订单版本"]).value
            if type(version) is not int or version < 1:
                raise FileFailure("订单版本无效")
            if ws.cell(index, headers["记录状态"]).value not in {"active", "void"}:
                raise FileFailure("记录状态无效")
            order_id = ws.cell(index, headers["订单ID"]).value
            signature = (version, ws.cell(index, headers["记录状态"]).value,
                         ws.cell(index, headers["请求ID"]).value)
            group = groups.setdefault(order_id, {"signature": signature, "totals": 0})
            if group["signature"] != signature:
                raise FileFailure("同一订单的版本/状态/原始请求不一致")
            group["totals"] += ws.cell(index, headers["行类型"]).value == "order_total"
        elif present_identity and any(ws.cell(index, headers[k]).value is not None for k in IDENTITY_HEADERS):
            raise FileFailure("部分订单身份损坏，原件只读保留")
    if any(group["totals"] != 1 for group in groups.values()):
        raise FileFailure("订单总计身份不完整，原件只读保留")


def receipts(book):
    if LOG_SHEET not in book.sheetnames:
        return {}
    sheet = book[LOG_SHEET]
    if tuple(c.value for c in sheet[1]) != LOG_HEADERS:
        raise FileFailure("写入回执格式损坏")
    records = {}
    for expected, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 1):
        rid, operation, fingerprint, order_id, revision, timestamp, encoded = row
        if rid in records or type(revision) is not int or revision != expected:
            raise FileFailure("写入回执编号或版本损坏")
        if str(uuid.UUID(rid)) != rid or operation not in {"append_quote", "append_remark", "void_order"}:
            raise FileFailure("写入回执身份或操作无效")
        if type(fingerprint) is not str or len(fingerprint) != 64 or any(c not in "0123456789abcdef" for c in fingerprint):
            raise FileFailure("写入回执哈希无效")
        data = json.loads(encoded)
        if data.get("fingerprint") != fingerprint or data.get("order_id") != order_id or data.get("revision") != revision:
            raise FileFailure("写入回执内容不一致")
        records[rid] = data
    return records


def inspect_receipt(raw, request_id):
    if raw is None:
        return None
    book = open_book(raw)
    try:
        return receipts(book).get(request_id)
    finally:
        book.close()


def identity_columns(ws, row, headers):
    for name in ("行类型", *IDENTITY_HEADERS):
        if name not in headers:
            headers[name] = ws.max_column + 1
            ws.cell(row, headers[name], name)
        if name in IDENTITY_HEADERS[:-1]:
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(headers[name])].hidden = True
    return headers


def assign_order(ws, start, end, headers, req):
    for index in range(start, end + 1):
        values = {"订单ID": req.order_id, "明细ID": str(uuid.uuid5(uuid.UUID(req.order_id), str(index - start))),
                  "订单版本": 1, "请求ID": req.request_id, "记录状态": "active"}
        for name, value in values.items():
            ws.cell(index, headers[name], value)
        for cell in ws[index]:
            if type(cell.value) is str:
                cell.data_type = "s"  # Raw identity/expressions are text, never formulas.


def workbook_transaction(path, req, base_headers, create, mutate, *, expected_hash=None, check_hash=False):
    def transform(raw):
        book = open_book(raw) if raw is not None else create()
        try:
            if raw is None:
                force_text(book)  # The trusted new template contains user labels.
            ws, row, headers = writable(book, base_headers)
            prior = receipts(book)
            headers = identity_columns(ws, row, headers)
            detail = mutate(book, ws, headers)
            result = {"fingerprint": req.fingerprint, "order_id": req.order_id,
                      "revision": len(prior) + 1, **detail}
            log = book[LOG_SHEET] if LOG_SHEET in book.sheetnames else book.create_sheet(LOG_SHEET)
            if log.cell(1, 1).value is None:
                for col, value in enumerate(LOG_HEADERS, 1):
                    log.cell(1, col, value)
            log.sheet_state = "hidden"
            log.append([req.request_id, req.operation, req.fingerprint, req.order_id,
                        result["revision"], datetime.now(timezone.utc).isoformat(), json.dumps(result, ensure_ascii=False)])
            force_text(book)
            from .reading import analyze
            checked = analyze(book)
            if checked.status != "ok":
                raise FileFailure("写后工作簿核对失败：" + "；".join(i.message for i in checked.issues))
            buffer = io.BytesIO()
            book.save(buffer)
            return buffer.getvalue(), result
        finally:
            book.close()
    options = {"expected_hash": expected_hash} if check_hash else {}
    return transact(path, req.request_id, req.fingerprint, req.operation, transform, inspect_receipt, **options)


def view_identity(book, sheet_name):
    if sheet_name != "利润明细":
        return {}
    try:
        row, headers = table(book[sheet_name])
    except FileFailure:
        return {}
    if not set(IDENTITY_HEADERS).issubset(headers):
        return {}
    return {index: (book[sheet_name].cell(index, headers["订单ID"]).value,
                    book[sheet_name].cell(index, headers["订单版本"]).value)
            for index in range(row + 1, book[sheet_name].max_row + 1)}


def target_order(ws, headers, order_id, expected_revision):
    rows = [index for index in range(1, ws.max_row + 1)
            if ws.cell(index, headers["订单ID"]).value == order_id]
    if not order_id or not rows:
        raise FileConflict("选中记录无稳定订单身份；旧记录暂不可编辑")
    if any(ws.cell(index, headers["订单版本"]).value != expected_revision for index in rows):
        raise FileConflict("订单版本已变化；请刷新后再操作")
    return rows
