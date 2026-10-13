"""Hydrate bounded Rust evidence into the existing read-only layout projection.

No workbook decoding or saving in Python. Numeric compatibility is checked
against the original token before preserving the legacy-v1 scalar binding.
"""
from datetime import datetime
from decimal import Decimal, InvalidOperation
import math

from .documents import Book, Sheet, Cell, Columns

class FormatReview(ValueError):
    def __init__(self, code, sheet='', cell=''):
        super().__init__(code)
        self.location = {'sheet': sheet, 'cell': cell}


def value(data, styles, epoch):
    raw = data['value'];kind = data['kind']
    if raw is None: return None, 'n'
    if kind == 's': return raw, 's'
    if kind == 'b': return raw == '1', 'b'
    if kind == 'd':
        from openpyxl.utils.datetime import from_ISO8601
        return from_ISO8601(raw), 'd'
    if kind != 'n': raise FormatReview('CELL_TYPE_UNSUPPORTED')
    try:
        exact = Decimal(raw)
        result = float(raw) if '.' in raw or 'e' in raw.lower() else int(raw)
    except (ValueError, InvalidOperation, OverflowError):
        raise FormatReview('INVALID_NUMERIC_TOKEN') from None
    if not exact.is_finite() or isinstance(result, float) and (not math.isfinite(result) or Decimal(str(result)) != exact):
        raise FormatReview('NUMERIC_TOKEN_REQUIRES_TEXT')
    # Retain P05-v1 timestamp interpretation and binding, without workbook IO.
    from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format, is_timedelta_format
    style = styles[data['style']];fmt = style['format_code'] or BUILTIN_FORMATS.get(style['num_fmt_id'], '')
    if is_date_format(fmt):
        from openpyxl.utils.datetime import from_excel
        try: result = from_excel(result, epoch, timedelta=is_timedelta_format(fmt))
        except (ValueError, OverflowError): raise FormatReview('DATE_VALUE_REVIEW') from None
        return result, 'd'
    return result, 'n'


def book(document):
    if document.get('format') != '777-excel-read-v1': raise FormatReview('FORMAT_VERSION_CONFLICT')
    epoch = datetime(1904, 1, 1) if document['date_1904'] else datetime(1899, 12, 30)
    sheets=[]
    for data in document['sheets']:
        sheet=object.__new__(Sheet);sheet.title=data['name'];sheet.sheet_state=data['state']
        sheet.max_row=data['max_row'];sheet.max_column=data['max_column']
        rows=[[None]*sheet.max_column for _ in range(sheet.max_row)]
        for item in data['cells']:
            row,col=item['row'],item['column'];coordinate=Cell(None,'n',row,col).coordinate
            try: scalar,kind=value(item,document['styles'],epoch)
            except FormatReview as exc: raise FormatReview(str(exc),sheet.title,coordinate) from None
            if scalar is not None: rows[row-1][col-1]=Cell(scalar,kind,row,col)
        sheet.rows=tuple(tuple(row) for row in rows);sheet.weight=128*sum(len(r) for r in rows)
        sheet.has_formulas=False;sheet._images=sheet._charts=False
        # Columns expects letter names. Cell coordinates supply the bounded map.
        sheet.column_dimensions=Columns(Cell(None,'n',1,c).coordinate[:-1] for c in data['hidden_columns'])
        sheets.append(sheet)
    result=object.__new__(Book);result.worksheets=tuple(sheets);result.sheetnames=tuple(s.title for s in sheets)
    result._sheets=dict(zip(result.sheetnames,sheets));result.weight=sum(s.weight for s in sheets)
    return result
