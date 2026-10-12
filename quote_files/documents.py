"""One workbook decode into detached rows shared by all read-only projections."""
from dataclasses import dataclass, replace
import io
import sys
from zipfile import ZipFile

from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

from .schema import UnsupportedSchema
from .transactions import digest

PARSER_VERSION = "ledger-document-v1"


@dataclass(frozen=True, slots=True)
class Cell:
    value: object
    data_type: str
    row: int
    column: int

    @property
    def coordinate(self):
        return f"{get_column_letter(self.column)}{self.row}"


@dataclass(frozen=True, slots=True)
class Column:
    hidden: bool = False


class Columns:
    def __init__(self, hidden):
        self.hidden = frozenset(hidden)

    def __getitem__(self, name):
        return Column(name in self.hidden)


class Sheet:
    """Read-only worksheet interface; cell/row access never opens XML again."""
    def __init__(self, source):
        self.title, self.sheet_state = source.title, source.sheet_state
        self.max_row, self.max_column = source.max_row, source.max_column
        if self.max_row > 100000 or self.max_column > 64:
            raise UnsupportedSchema("工作表超过当前100000行/64列读取范围，原件保留")
        rows, weight, formulas = [], 0, False
        for row_number, cells in enumerate(source.iter_rows(), 1):
            values = []
            for column, cell in enumerate(cells, 1):
                formulas |= cell.data_type == "f"
                value = cell.value
                values.append(Cell(value, cell.data_type, row_number, column) if value is not None else None)
                weight += 8 if value is None else 128 + sys.getsizeof(value)
            rows.append(tuple(values))
        self.rows = tuple(rows)
        self.weight = weight + len(rows) * 64
        self.has_formulas = formulas
        hidden = set()
        for name, column in source.column_dimensions.items():
            if column.hidden:
                start = column.min or column_index_from_string(name)
                end = column.max or start
                hidden.update(get_column_letter(i) for i in range(start, min(end, 64) + 1))
        self.column_dimensions = Columns(hidden)
        self._images, self._charts = bool(source._images), bool(source._charts)

    def cell(self, row, column):
        if row < 1 or column < 1:
            raise ValueError("行列从1开始")
        value = self.rows[row - 1][column - 1] if row <= len(self.rows) and column <= len(self.rows[row - 1]) else None
        return value if value is not None else Cell(None, "n", row, column)

    def __getitem__(self, key):
        if type(key) is int:
            return tuple(self.cell(key, column) for column in range(1, self.max_column + 1))
        column, row = coordinate_from_string(key)
        return self.cell(row, column_index_from_string(column))

    def iter_rows(self, min_row=1, max_row=None, min_col=1, max_col=None, values_only=False):
        for row in range(min_row, (self.max_row if max_row is None else max_row) + 1):
            cells = tuple(self.cell(row, column) for column in range(min_col, (self.max_column if max_col is None else max_col) + 1))
            yield tuple(cell.value for cell in cells) if values_only else cells


class Book:
    def __init__(self, source):
        self.worksheets = tuple(Sheet(sheet) for sheet in source.worksheets)
        self.sheetnames = tuple(sheet.title for sheet in self.worksheets)
        self._sheets = dict(zip(self.sheetnames, self.worksheets))
        self.weight = sum(sheet.weight for sheet in self.worksheets)

    def __getitem__(self, name):
        return self._sheets[name]


@dataclass(frozen=True)
class Document:
    book: object
    result: object
    identities: dict
    orders: dict
    weight: int
    workbook_opens: int = 0
    formula_cache_opens: int = 0
    projection: str = "full"

    def ledger(self):
        # Callers may keep/alter their UI result; the cached model is private.
        return replace(self.result, group_keys=dict(self.result.group_keys), group_totals=dict(self.result.group_totals))

    def summary(self):
        result = replace(self.result, rows=(), group_keys={}, group_totals={}, _detail_count=self.result.detail_count)
        return Document(None, result, {}, {}, 2048 + sum(128 + len(i.message) * 4 for i in result.issues), projection="summary")


def failed_document(error, raw=None, unsupported=False):
    from .reading import LedgerRead, ReadIssue
    result = LedgerRead("unsupported" if unsupported else "error",
        issues=(ReadIssue("UNSUPPORTED_SCHEMA" if unsupported else "READ_ERROR", f"{type(error).__name__}: {error}"),),
        source_sha256=digest(raw) if raw is not None else None)
    return Document(None, result, {}, {}, 1024)


def parse_document(raw, *, opener=None, analyzer=None):
    from .reading import ReadIssue, analyze, open_book
    from .workbooks import view_identity
    from .orders import order_views
    opener, analyzer = opener or open_book, analyzer or analyze
    primary = cached = None
    opens = formula_opens = 0
    try:
        primary = opener(raw, for_write=False);opens += 1
        book = Book(primary)
        cached_rows = None
        if any(sheet.has_formulas for sheet in book.worksheets):
            cached = opener(raw, data_only=True, for_write=False);opens += 1;formula_opens += 1
            cached_rows = Book(cached)
        result = analyzer(book, cached_rows)
        with ZipFile(io.BytesIO(raw)) as package:
            extensions = any(any(token in name.lower() for token in ("externallinks", "vbaproject", "customxml", "embeddings", "drawings", "vml")) for name in package.namelist())
        if extensions and result.schema is not None:
            reason = "含未支持扩展/外链/宏/图片，原件只读保留"
            result = replace(result, schema=replace(result.schema, writable=False,
                readonly_reasons=(*result.schema.readonly_reasons, reason)), issues=(*result.issues, ReadIssue("READ_ONLY", reason)))
        result = replace(result, source_sha256=digest(raw))
        identities = view_identity(book, "利润明细") if result.status == "ok" and result.schema.writable else {}
        orders = order_views(book, result.schema) if result.status == "ok" and result.schema.identified else {}
        return Document(book, result, identities, orders, book.weight + len(result.rows) * 512 + len(raw), opens, formula_opens)
    except Exception as exc:
        return replace(failed_document(exc, raw, isinstance(exc, UnsupportedSchema)), workbook_opens=opens, formula_cache_opens=formula_opens)
    finally:
        try:
            if cached is not None:cached.close()
        finally:
            if primary is not None:primary.close()
