"""Closed workbook -> plain view data. No Tk, retained handles or file writes."""
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from quote_core.numbers import money_text, settle
from .reading import read_document
from .scheduling import checkpoint


@dataclass(frozen=True)
class View:
    path: str
    sheet: str
    sheets: tuple
    headings: tuple
    rows: tuple
    ledger: object
    identities: dict
    orders: dict


def cell_text(value, money=False):
    if value is None:
        return ""
    if money and type(value) in (int, float, Decimal, str):
        try:
            value = Decimal(str(value))
            if value.is_finite():
                return money_text(settle(value))
        except Exception:
            pass
    return str(value)


def read_view(path, sheet, role, context, document=None):
    context.check()
    document = document or read_document(path)
    result = document.ledger()
    context.check()
    book = document.book
    if book is None:
        raise ValueError("工作簿读取失败：" + "; ".join(i.message for i in result.issues))
    sheets = tuple(ws.title for ws in book.worksheets if ws.sheet_state == "visible")
    sheet = sheet or ("利润明细" if "利润明细" in sheets else sheets[0])
    ws = book[sheet]
    identities = dict(document.identities) if sheet == "利润明细" else {}
    orders = dict(document.orders)
    rows = []
    if role == "manager":
        from openpyxl.utils import get_column_letter
        columns = [i for i in range(1, ws.max_column + 1) if not ws.column_dimensions[get_column_letter(i)].hidden]
        headings = ("行", *(get_column_letter(i) for i in columns))
        for i, data in enumerate(ws.iter_rows(max_row=ws.max_row, max_col=ws.max_column, values_only=True), 1):
            checkpoint()
            context.check()
            rows.append((i, (str(i), *(cell_text(data[c - 1]) for c in columns))))
    else:
        if result.status != "ok" or sheet != "利润明细":
            raise ValueError("该工作表不能作为销售明细：" + "; ".join(i.message for i in result.issues))
        headings = ("淘宝ID", "订单号", "微信号", "记录时间", "项目名称", "基础类型", "高", "宽", "数量", "销售单价", "折扣", "销售额", "折前销售额", "总销售额")
        h = result.schema.columns
        mapping = {"高": "高度(m)", "宽": "宽度(m)"}
        for i, data in enumerate(ws.iter_rows(min_row=result.schema.header_row + 1, values_only=True), result.schema.header_row + 1):
            checkpoint()
            context.check()
            if i not in result.visible_rows:
                continue
            values = [cell_text(data[h[mapping.get(name, name)] - 1], name in {"销售单价", "销售额", "折前销售额"}) for name in headings[:-1]]
            total = result.group_totals.get(result.group_keys[i])
            values.append(format(total, ".2f") if total is not None else "—（历史归属待核对）")
            rows.append((i, tuple(values)))
    return View(str(Path(path).absolute()), sheet, sheets, headings, tuple(rows), result, identities, orders)
