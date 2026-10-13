"""777-output-v1: reproducible ledger snapshots, never executable Excel cells."""
from datetime import datetime
import hashlib
import io
import json
import os
from pathlib import Path

from quote_core.numbers import money_text
from .protocol import APIError

FORMAT = '777-output-v1'
MAX_BYTES = 4 * 1024 * 1024
KINDS = {'excel_quote': '.xlsx', 'excel_admin': '.xlsx', 'quote_png': '.png'}


def payload_hash(payload):
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True,
                                    separators=(',', ':'), allow_nan=False).encode()).hexdigest()


def cell_text(value):
    # XML cannot represent these controls. Display their exact code points;
    # original text remains in the immutable ledger snapshot and its hash.
    return ''.join('\\u%04x'%ord(ch) if ord(ch)<32 and ch not in '\t\n\r' else ch for ch in str(value))


def public_order(order):
    value = {k: order[k] for k in ('order_id', 'revision', 'state', 'price_version_id',
                                 'created_at', 'details', 'income_cents', 'notes',
                                 'rule_version', 'numeric_profile')}
    value['lines'] = [{k: line[k] for k in ('sequence', 'raw', 'subtotal_before_cents', 'income_cents')} |
                      {'normalized': {k: v for k, v in line['normalized'].items()
                                      if k in {'q', 'h', 'w', 'sale', 'k', 'percent', 'single'}}}
                      for line in order['lines']]
    return value


def render(payload):
    if payload.get('format') != FORMAT or payload.get('kind') not in KINDS: raise APIError(409,'OUTPUT_VERSION_CONFLICT')
    kind = payload['kind']; order = payload['order']; output = io.BytesIO()
    if kind.startswith('excel_'):
        from openpyxl import Workbook
        from openpyxl.styles import Font
        workbook = Workbook(); workbook.remove(workbook.active)
        workbook.properties.created = workbook.properties.modified = datetime(2000, 1, 1)
        def sheet(name, rows):
            ws = workbook.create_sheet(name)
            for row in rows:
                ws.append([cell_text(cell) for cell in row])
                for cell in ws[ws.max_row]:
                    # '='/'+'/'@' and customer text remain literal strings.
                    cell.data_type = 's'; cell.number_format = '@'
            for cell in ws[1]: cell.font = Font(bold=True)
            ws.freeze_panes = 'A2'
            for column in ws.columns: ws.column_dimensions[column[0].column_letter].width = 24
            return ws
        sheet('Protocol', [('format', FORMAT), ('job_id', payload['job_id']),
                           ('payload_sha256', payload_hash(payload)), ('kind', kind)])
        sheet('Order', [('field', 'value'), *[(k, order[k]) for k in
              ('order_id', 'revision', 'state', 'price_version_id', 'rule_version', 'numeric_profile', 'created_at')],
              *order['details'].items(), ('income_cents', order['income_cents']),
              ('income_yuan', money_text(int(order['income_cents'])))])
        if kind == 'excel_admin':
            for k in ('cost_cents', 'profit_cents'): workbook['Order'].append([k, order[k]])
            for row in workbook['Order']:
                for cell in row: cell.data_type = 's'; cell.number_format = '@'
        headers = ['sequence', 'raw', 'normalized', 'subtotal_before_cents', 'income_cents']
        if kind == 'excel_admin': headers += ['cost_cents', 'profit_cents']
        sheet('Lines', [headers, *[[json.dumps(line[k], ensure_ascii=False, sort_keys=True)
                                  if isinstance(line[k], dict) else str(line[k]) for k in headers]
                                  for line in order['lines']]])
        sheet('Notes', [('note_id', 'author', 'body', 'created_at'),
                        *[[note[k] for k in ('note_id', 'author', 'body', 'created_at')] for note in order['notes']]])
        workbook.save(output); workbook.close()
    else:
        from PIL import Image, ImageDraw, ImageFont, PngImagePlugin
        # Windows system CJK fonts; do not silently emit unreadable Chinese.
        fonts = [Path(os.environ.get('WINDIR', 'C:/Windows'))/'Fonts'/name
                 for name in ('msyh.ttc', 'simhei.ttf')]
        font = next((ImageFont.truetype(str(p), 24) for p in fonts if p.is_file()), None)
        if font is None and os.name == 'nt': raise APIError(503, 'OUTPUT_CJK_FONT_UNAVAILABLE')
        if font is None: font = ImageFont.load_default(size=22)  # Auxiliary only.
        lines = ['Quotation / 报价', order['order_id'], 'Revision / 版本: '+str(order['revision']),
                 'State / 状态: '+order['state'], '客户: '+order['details']['customer'],
                 '订单号: '+order['details']['order_ref'], '微信: '+order['details']['wechat'],
                 '收入: '+money_text(int(order['income_cents']))+' 元']
        for row in order['lines']:
            label = row['raw'].get('name') or row['raw'].get('product_id', '折扣')
            lines.append(str(row['sequence'])+'. '+label+' '+money_text(int(row['income_cents']))+' 元')
        for note in order['notes']: lines.append('备注: '+note['body'])
        wrapped = [line[i:i+40] for line in lines for i in range(0, max(1, len(line)), 40)]
        if len(wrapped) > 400: raise APIError(422, 'OUTPUT_IMAGE_TOO_TALL')
        image = Image.new('RGB', (1200, 40+len(wrapped)*36), 'white'); draw = ImageDraw.Draw(image)
        for index, line in enumerate(wrapped): draw.text((24, 20+index*36), line, fill='black', font=font)
        meta = PngImagePlugin.PngInfo()
        meta.add_text('777_format', FORMAT); meta.add_text('777_job_id', payload['job_id'])
        meta.add_text('777_payload_sha256', payload_hash(payload))
        image.save(output, format='PNG', pnginfo=meta); image.close()
    raw = output.getvalue()
    if len(raw) > MAX_BYTES: raise APIError(413, 'OUTPUT_TOO_LARGE')
    validate(raw, payload)
    return raw


def validate(raw, payload):
    if not 0 < len(raw) <= MAX_BYTES: raise APIError(422, 'OUTPUT_BLOB_INVALID')
    expected = payload_hash(payload)
    if payload['kind'].startswith('excel_'):
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            if workbook.sheetnames != ['Protocol', 'Order', 'Lines', 'Notes']: raise ValueError()
            values = dict(workbook['Protocol'].values)
            if values != {'format': FORMAT, 'job_id': payload['job_id'], 'payload_sha256': expected, 'kind': payload['kind']}: raise ValueError()
            if any(cell.data_type == 'f' for ws in workbook for row in ws for cell in row): raise ValueError()
        except (ValueError, KeyError): raise APIError(422, 'OUTPUT_BLOB_INVALID') from None
        finally: workbook.close()
    else:
        from PIL import Image
        with Image.open(io.BytesIO(raw)) as image:
            if image.format != 'PNG' or image.info.get('777_format') != FORMAT or image.info.get('777_job_id') != payload['job_id'] or image.info.get('777_payload_sha256') != expected:
                raise APIError(422, 'OUTPUT_BLOB_INVALID')
            image.verify()
